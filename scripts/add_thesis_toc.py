#!/usr/bin/env python3
"""在英文摘要 Keywords 后插入学位论文目录（点线页码 + PAGEREF）。"""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "华东师范大学硕士论文.docx"

CHAPTER_RE = re.compile(r"^第[一二三四五六七八九十百]+章")
SECTION_RE = re.compile(r"^\d+\.\d+(?!\.\d)")
SUBSECTION_RE = re.compile(r"^\d+\.\d+\.\d+")
END_MATTER = {"参考文献", "致谢", "附录"}


def set_run_font(run, *, east_asia="宋体", western="Times New Roman", size_pt=12, bold=False):
    run.font.name = western
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:ascii"), western)
    rFonts.set(qn("w:hAnsi"), western)


def insert_paragraph_after(para: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    para._p.addnext(new_p)
    new_para = Paragraph(new_p, para._parent)
    if text:
        new_para.add_run(text)
    return new_para


def add_page_break(para: Paragraph) -> None:
    run = para.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def heading_level(text: str) -> int | None:
    t = text.strip()
    if not t:
        return None
    if CHAPTER_RE.match(t) or t in END_MATTER:
        return 0
    if SUBSECTION_RE.match(t):
        return 2
    if SECTION_RE.match(t):
        return 1
    return None


def add_bookmark(paragraph: Paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p = paragraph._p
    if len(p):
        p.insert(0 if p[0].tag != qn("w:pPr") else 1, start)
    else:
        p.append(start)
    p.append(end)


def clear_toc_bookmarks(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag == qn("w:bookmarkStart") and (child.get(qn("w:name")) or "").startswith("_TocThesis"):
            bid = child.get(qn("w:id"))
            paragraph._p.remove(child)
            for child2 in list(paragraph._p):
                if child2.tag == qn("w:bookmarkEnd") and child2.get(qn("w:id")) == bid:
                    paragraph._p.remove(child2)


def add_pageref_field(paragraph: Paragraph, bookmark_name: str) -> None:
    run = paragraph.add_run()
    r = run._r
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" PAGEREF {bookmark_name} \\h "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "—"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(text)
    r.append(fld_end)
    set_run_font(run, size_pt=12, bold=False)


def style_toc_title(para: Paragraph) -> None:
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(12)
    pf.space_after = Pt(12)
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    for run in para.runs:
        set_run_font(run, east_asia="黑体", western="Times New Roman", size_pt=16, bold=True)


def style_toc_entry(para: Paragraph, level: int, tab_pos_cm: float) -> None:
    indents = {0: Cm(0), 1: Cm(0.74), 2: Cm(1.48)}
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.left_indent = indents.get(level, Cm(0))
    pf.tab_stops.add_tab_stop(Cm(tab_pos_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def normalize_title(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip())


def remove_existing_toc(doc: Document) -> None:
    paras = list(doc.paragraphs)
    toc_start = None
    for i, p in enumerate(paras):
        if (p.text or "").strip() == "目录":
            toc_start = i
            break
    if toc_start is None:
        return
    j = toc_start
    while j < len(paras):
        t = (paras[j].text or "").strip()
        if CHAPTER_RE.match(t) and j > toc_start + 1:
            # first 第一章 after 致谢 is body; stop before body chapter
            # If we are still in TOC, 第一章 appears as first entry — skip until after 致谢
            pass
        j += 1

    # Safer: delete from 目录 through the page-break paragraph before body 第一章
    seen_toc = False
    seen_ack = False
    to_remove = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "目录":
            seen_toc = True
            to_remove.append(p._p)
            continue
        if not seen_toc:
            continue
        if t.startswith("致谢"):
            seen_ack = True
            to_remove.append(p._p)
            continue
        if seen_ack and CHAPTER_RE.match(t):
            break
        to_remove.append(p._p)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def main() -> None:
    backup = ROOT / f"华东师范大学硕士论文.backup.{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    shutil.copy2(SRC, backup)
    doc = Document(str(SRC))
    remove_existing_toc(doc)

    keywords_para = None
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("Keywords:") or t.startswith("KEYWORDS:"):
            keywords_para = p
            break
    if keywords_para is None:
        raise SystemExit("未找到英文 Keywords 段落")

    chapter1 = None
    for p in doc.paragraphs:
        if CHAPTER_RE.match((p.text or "").strip()):
            chapter1 = p
            break
    if chapter1 is None:
        raise SystemExit("未找到第一章")

    to_remove = []
    crossed = False
    for p in doc.paragraphs:
        if p._p is keywords_para._p:
            crossed = True
            continue
        if p._p is chapter1._p:
            break
        if crossed:
            to_remove.append(p._p)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)

    sec = doc.sections[0]
    usable_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000
    tab_pos = max(usable_cm - 0.1, 12.0)

    cursor = keywords_para
    pb1 = insert_paragraph_after(cursor, "")
    add_page_break(pb1)
    cursor = pb1

    title = insert_paragraph_after(cursor, "目录")
    style_toc_title(title)
    cursor = title

    headings: list[tuple[Paragraph, str, int, str]] = []
    seen = False
    bi = 0
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("Keywords:") or t.startswith("KEYWORDS:"):
            seen = True
            continue
        if not seen or t == "目录":
            continue
        lvl = heading_level(t)
        if lvl is None:
            continue
        name = f"_TocThesis{bi:03d}"
        clear_toc_bookmarks(p)
        add_bookmark(p, name, 2000 + bi)
        headings.append((p, t, lvl, name))
        bi += 1

    for _, t, lvl, name in headings:
        entry = insert_paragraph_after(cursor, "")
        style_toc_entry(entry, lvl, tab_pos)
        run = entry.add_run(normalize_title(t))
        set_run_font(run, east_asia="宋体", western="Times New Roman", size_pt=12, bold=False)
        entry.add_run("\t")
        add_pageref_field(entry, name)
        cursor = entry

    pb2 = insert_paragraph_after(cursor, "")
    add_page_break(pb2)

    doc.save(str(SRC))
    print(f"backup: {backup.name}")
    print(f"toc entries: {len(headings)}")
    print(f"saved: {SRC.name}")
    print("请用 Word/WPS 打开后：全选 → 右键「更新域」→「更新整个目录」，以刷新页码。")


if __name__ == "__main__":
    main()
