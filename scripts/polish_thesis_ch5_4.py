#!/usr/bin/env python3
"""定点理顺论文 5.4：真表格 + 重生成并插入图5-11/5-12 + 表5-4。

用法:
  .venv/bin/python scripts/polish_thesis_ch5_4.py
  .venv/bin/python scripts/polish_thesis_ch5_4.py --no-backup
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from copy import deepcopy
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from matplotlib import rcParams

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"
SUMMARY = ROOT / "data" / "eval" / "table_6_3_summary_60.json"
BACKUP = ROOT / "华东师范大学硕士论文.docx.bak-before-ch5-4-polish"
FIG_DIR = ROOT / "docs" / "superpowers" / "architecture"
FIG_FACT = FIG_DIR / "fig-5-11-factual-60.png"
FIG_COMP = FIG_DIR / "fig-5-12-completeness-format-60.png"

BASELINES = ("B0", "B1", "B2", "B3")
COLORS = ["#9ca3af", "#3b82f6", "#86efac", "#f59e0b"]
PHASES = ("震前", "震中", "震后")


def _setup_matplotlib_font() -> str:
    rcParams["axes.unicode_minus"] = False
    for name in ("STHeiti", "Songti SC", "Arial Unicode MS", "PingFang SC", "Heiti SC"):
        rcParams["font.sans-serif"] = [name]
        return name
    rcParams["font.sans-serif"] = ["DejaVu Sans"]
    return "DejaVu Sans"


def load_summary() -> dict:
    return json.loads(SUMMARY.read_text(encoding="utf-8"))


def generate_figures(data: dict) -> None:
    _setup_matplotlib_font()
    overall = data["overall"]
    labels = list(BASELINES)
    factual = [float(overall[b]["factual_accuracy_pct"]) for b in BASELINES]
    completeness = [float(overall[b]["completeness_mean"]) for b in BASELINES]
    fmt = [float(overall[b]["format_compliance_pct"]) for b in BASELINES]

    FIG_DIR.mkdir(parents=True, exist_ok=True)

    fig, ax = plt.subplots(figsize=(7.2, 4.2), dpi=150)
    bars = ax.bar(labels, factual, color=COLORS, width=0.62)
    ax.set_ylabel("事实一致性（%）")
    ax.set_title("图5-11  四基线事实一致性对比（60题离线消融）")
    ax.set_ylim(0, 100)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    for bar, val in zip(bars, factual):
        ax.text(bar.get_x() + bar.get_width() / 2, val + 1.5, f"{val:g}", ha="center", va="bottom", fontsize=10)
    fig.tight_layout()
    fig.savefig(FIG_FACT, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(9.5, 4.2), dpi=150)
    bars0 = axes[0].bar(labels, completeness, color=COLORS, width=0.62)
    axes[0].set_ylabel("要点完整性（1–5）")
    axes[0].set_ylim(0, 5)
    axes[0].spines["top"].set_visible(False)
    axes[0].spines["right"].set_visible(False)
    for bar, val in zip(bars0, completeness):
        axes[0].text(bar.get_x() + bar.get_width() / 2, val + 0.08, f"{val:g}", ha="center", va="bottom", fontsize=9)

    bars1 = axes[1].bar(labels, fmt, color=COLORS, width=0.62)
    axes[1].set_ylabel("格式合规（%）")
    axes[1].set_ylim(0, 100)
    axes[1].spines["top"].set_visible(False)
    axes[1].spines["right"].set_visible(False)
    for bar, val in zip(bars1, fmt):
        axes[1].text(bar.get_x() + bar.get_width() / 2, val + 1.5, f"{val:g}", ha="center", va="bottom", fontsize=9)

    fig.suptitle("图5-12  四基线要点完整性与格式合规对比（60题离线消融）", y=1.02)
    fig.tight_layout()
    fig.savefig(FIG_COMP, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def _set_run_font(
    run,
    *,
    east_asia: str = "宋体",
    western: str = "Times New Roman",
    size_pt: float = 12,
    bold: bool = False,
) -> None:
    run.font.name = western
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), east_asia)


def _clear_runs(para: Paragraph) -> None:
    for run in para.runs:
        run.text = ""


def set_para_text(para: Paragraph, text: str) -> None:
    _clear_runs(para)
    if para.runs:
        para.runs[0].text = text
        _set_run_font(para.runs[0])
    else:
        run = para.add_run(text)
        _set_run_font(run)


def _format_para(para: Paragraph, level: str) -> None:
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if level == "h2":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=14, bold=True)
    elif level == "h3":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(3)
        pf.space_after = Pt(3)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=12, bold=True)
    elif level in ("figure_caption", "table_title"):
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
        size = 12 if level == "table_title" else 10.5
        ea = "黑体" if level == "table_title" else "宋体"
        for run in para.runs:
            _set_run_font(run, east_asia=ea, size_pt=size, bold=(level == "table_title"))
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.74)
        for run in para.runs:
            _set_run_font(run, east_asia="宋体", size_pt=12, bold=False)


def _insert_para_before(anchor: Paragraph, text: str, level: str) -> Paragraph:
    new_p = deepcopy(anchor._p)
    for child in list(new_p):
        local = child.tag.split("}")[-1]
        if local in ("drawing", "pict", "r", "hyperlink"):
            new_p.remove(child)
    anchor._p.addprevious(new_p)
    para = Paragraph(new_p, anchor._parent)
    set_para_text(para, text)
    _format_para(para, level)
    return para


def _delete_paragraph(para: Paragraph) -> None:
    el = para._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def _find_body_idx(doc: Document, startswith: str, contain: str) -> int:
    idxs = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if "\t" in p.text:
            continue
        if t.startswith(startswith) and contain in t:
            idxs.append(i)
    if not idxs:
        raise KeyError(f"找不到: {startswith} {contain}")
    return idxs[-1]


def _add_picture(para: Paragraph, image_path: Path, *, width_in: float) -> None:
    run = para.add_run()
    run.add_picture(str(image_path), width=Inches(width_in))
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)


def _set_cell_border(cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        element = tcBorders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tcBorders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def _style_table(table: Table) -> None:
    for style_name in ("Table Grid", "表格网格", "Normal Table", "Table Normal"):
        try:
            table.style = style_name
            break
        except KeyError:
            continue
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(cell)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.space_before = Pt(2)
                pf.space_after = Pt(2)
                for run in p.runs:
                    _set_run_font(run, east_asia="宋体", size_pt=10.5, bold=False)
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                _set_run_font(run, east_asia="黑体", size_pt=10.5, bold=True)


def _insert_table_before(doc: Document, anchor: Paragraph, headers: list[str], rows: list[list[str]]) -> Table:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            table.rows[i + 1].cells[j].text = str(val)
    _style_table(table)
    anchor._p.addprevious(table._tbl)
    return table


def _fmt(v) -> str:
    if isinstance(v, float):
        if abs(v - round(v)) < 1e-9:
            return str(int(round(v))) if v >= 10 else f"{v:g}"
        return f"{v:g}"
    return str(v)


def rewrite_section(doc: Document, data: dict) -> None:
    i54 = _find_body_idx(doc, "5.4", "系统实验")
    i55 = _find_body_idx(doc, "5.5", "小结")
    if i55 <= i54:
        raise RuntimeError(f"区间异常: 5.4={i54}, 5.5={i55}")

    # delete [i54, i55)
    for p in list(doc.paragraphs[i54:i55]):
        _delete_paragraph(p)

    anchor = doc.paragraphs[_find_body_idx(doc, "5.5", "小结")]
    overall = data["overall"]
    by_phase = data["by_phase"]

    f = {b: overall[b]["factual_accuracy_pct"] for b in BASELINES}
    c = {b: overall[b]["completeness_mean"] for b in BASELINES}
    m = {b: overall[b]["format_compliance_pct"] for b in BASELINES}

    blocks: list[tuple[str, str]] = [
        ("h2", "5.4  系统实验与分析"),
        ("h3", "5.4.1  实验设置"),
        (
            "body",
            "实验环境为作者本地工作站（macOS，Python 3.10+，PyTorch 自动选择设备），"
            "图数据库与向量索引按项目默认端口部署，微调权重为地震应急 LoRA（秩8，α16）。"
            "评测问集按震前、震中、震后各20题共60题构造，版本2026-06；"
            "开发期快速子集（每阶段4题共12题）仅作调试参考，正文结论以全量60题为准。",
        ),
        (
            "body",
            "四组基线定义：B0 仅语言模型，不注入图谱与检索；B1 仅图谱上下文；B2 仅向量检索；"
            "B3 同时启用图谱与检索。离线消融统一关闭动态检索与第三层输出增强，以隔离静态知识贡献；"
            "解码参数固定（温度约0.55、top-p约0.88、最大新生成约384 token），history 为空。"
            "评分采用自动 grounding 重叠与长度/关键词启发式 proxy（事实一致性百分比、要点完整性1—5、格式合规率），"
            "非人工双盲、非专家问卷；自动分用于批量对比，不能替代人工专家评测。配置快照见表5-1。",
        ),
        ("table_title", "表5-1  离线消融实验配置快照"),
        ("table1", ""),
        ("h3", "5.4.2  对比实验结果与分析"),
        (
            "body",
            "表5-2 汇总四基线在60题全量离线消融上的自动评分均值。"
            f"事实一致性：B0 {_fmt(f['B0'])}%，B1 {_fmt(f['B1'])}%，B2 {_fmt(f['B2'])}%，B3 {_fmt(f['B3'])}%；"
            f"要点完整性：B0 {_fmt(c['B0'])}，B1 {_fmt(c['B1'])}，B2 {_fmt(c['B2'])}，B3 {_fmt(c['B3'])}；"
            f"格式合规：B0 {_fmt(m['B0'])}%，B1 {_fmt(m['B1'])}%，B2 {_fmt(m['B2'])}%，B3 {_fmt(m['B3'])}%。"
            "图5-11 给出四基线事实一致性柱状对比；图5-12 给出要点完整性与格式合规对照。",
        ),
        ("table_title", "表5-2  四基线60题自动评分汇总"),
        ("table2", ""),
        ("image_fact", ""),
        ("figure_caption", "图5-11  四基线事实一致性对比（60题离线消融）"),
        ("image_comp", ""),
        ("figure_caption", "图5-12  四基线要点完整性与格式合规对比（60题离线消融）"),
        (
            "body",
            "从通路增益看，B3 事实一致性高于 B0、B1 与 B2，说明图谱可核查字段与检索步骤细节在 grounding proxy 上具有叠加项；"
            "相对 B0，B1 单独引入图谱亦提升事实分，而 B2 单独检索在本评分规则下未形成正向事实增益（见5.4.3）。"
            "要点完整性未必单调最优：更长证据上下文可使模型回答更短或更聚焦，长度分档下降——"
            f"B3 完整性（{_fmt(c['B3'])}）低于 B0（{_fmt(c['B0'])}）即属此类，不宜表述为“全面显著优于”。",
        ),
        (
            "body",
            "与 GraphRAG、KnowledGPT 的对比仅做设计维度对照，不报告二者官方流水线在本问集上的分数。"
            "表5-3 从构图粒度、阶段感知、动态源与显式调度四方面对照。",
        ),
        ("table_title", "表5-3  与 GraphRAG、KnowledGPT 的设计对比"),
        ("table3", ""),
        (
            "body",
            "本文未在同一60题问集上复现 GraphRAG 与 KnowledGPT 的官方流水线，故不做同设定数值对比实验；"
            "对比结论限于方法假设与工程边界，避免虚构跨系统分数高低。",
        ),
        ("h3", "5.4.3  消融实验"),
        (
            "body",
            "以 B0 为参照观察模块贡献：引入图谱（B1）主要抬升事实一致性 proxy；"
            "单独检索（B2）在本规则下事实分下降，但可为完整性提供步骤性文本；"
            f"图谱与检索联用（B3）事实 proxy 最高（{_fmt(f['B3'])}%），说明两静态源互补——"
            "字段以图谱为准、动作步骤以检索主题 steps 为准——是本场景合理配置。"
            f"格式合规在各组均维持较高水平（约 {_fmt(m['B0'])}%—{_fmt(m['B3'])}%）。"
            "上述贡献均来自关闭动态的离线消融，不可外推为在线全功能路径的动态源收益。"
            "分阶段（各20题）明细见表5-4。",
        ),
        ("table_title", "表5-4  分阶段自动评分汇总（各阶段20题）"),
        ("table4", ""),
        (
            "body",
            "由表5-4可见：震前子集少涉精确震例参数，B0 与 B3 差距相对可控；"
            "震中子集图谱区域震例与检索避险主题贡献增大；震后子集政策与心理主题检索命中更突出，B3 事实一致性最高。"
            "解读时应结合阶段关键词命中，而非仅看总分。",
        ),
        (
            "body",
            f"全量60题上 B2（仅检索）事实一致性为 {_fmt(f['B2'])}%，低于 B0（{_fmt(f['B0'])}%）与 B1（{_fmt(f['B1'])}%），"
            "与“RAG 必然提升事实性”的直觉相悖。根因分析如下，而非归因于“样本小”。"
            "第一，主题级分块将 title、category、phase_tag 与 steps 合并为一段，步骤性自然语言与评分抽取的“可核对事实 token”重叠度往往低于图谱结构化字段。"
            "第二，嵌入与 Top-K 召回在口语问句上可能命中邻近但非精确主题，注入噪声后模型复述标题而省略数值型事实。"
            "第三，提示中可核对片段以【参考资料】步骤为主；当 B2 关闭图谱段时，图谱重叠项权重为零，检索段在 grounding proxy 下相对吃亏。"
            f"第四，B3（{_fmt(f['B3'])}%）事实分最高，表明检索应与结构化图谱联用，而非否定检索路径本身。",
        ),
        ("h3", "5.4.4  案例分析"),
        (
            "body",
            "案例均为定性演示，与60题消融问集口径分离，不把案例分数写入表5-2。",
        ),
        (
            "body",
            "案例一（在线演示路径，开启动态）：问句“刚才地震多大？震中在哪？”——"
            "分类为震中，动态段写入最新速报摘要，静态段补充实时信息获取渠道；"
            "回答应优先引用动态参数并提示权威来源。该案例不参与 B0—B3 数值表。",
        ),
        (
            "body",
            "案例二（离线 B3）：问句“室内应该躲哪里？”——"
            "检索命中室内避险主题，图谱关联步骤节点，回答应含“伏地、遮挡、手抓牢”等要点。",
        ),
        (
            "body",
            "案例三（离线 B1）：问句“四川近年来发生过哪些典型地震？”——"
            "图谱返回汶川、芦山、九寨沟、泸定等节点摘要，震级与时间字段可 ground。"
            "上述案例用于说明路径边界与证据形态，不作统计推断。",
        ),
        ("h3", "5.4.5  系统可用性评测"),
        (
            "body",
            "可用性评测（任务完成率、满意度 Likert、主观可信度等）本次未开展正式用户问卷，列为后续工作（待开展）。"
            "当前论文不以虚构问卷或小样本访谈替代系统评价；主体结论建立在60题离线消融与自动 proxy 之上。"
            "后续计划：分层抽取问句开展双人专家打分并报告一致性，再辅以有限用户可用性测试。"
            "在正式问卷完成前，界面截图与在线案例仅作功能展示，不构成可用性主证据。",
        ),
    ]

    # Same-anchor addprevious: later inserts land closer to 5.5 → use forward order.
    for level, text in blocks:
        if level == "table1":
            _insert_table_before(
                doc,
                anchor,
                ["配置项", "取值"],
                [
                    ["图谱注入", "B1/B3 开启，B0/B2 关闭"],
                    ["向量检索", "B2/B3 开启，B0/B1 关闭"],
                    ["动态检索", "全部关闭"],
                    ["输出增强（第三层）", "全部关闭"],
                    ["评测问量", "60（震前/震中/震后各20）"],
                    ["评分指标", "事实一致性、要点完整性、格式合规（自动 proxy）"],
                ],
            )
            continue
        if level == "table2":
            _insert_table_before(
                doc,
                anchor,
                ["指标", "B0", "B1", "B2", "B3"],
                [
                    ["事实一致性(%)", _fmt(f["B0"]), _fmt(f["B1"]), _fmt(f["B2"]), _fmt(f["B3"])],
                    ["要点完整性(1-5)", _fmt(c["B0"]), _fmt(c["B1"]), _fmt(c["B2"]), _fmt(c["B3"])],
                    ["格式合规(%)", _fmt(m["B0"]), _fmt(m["B1"]), _fmt(m["B2"]), _fmt(m["B3"])],
                ],
            )
            continue
        if level == "table3":
            _insert_table_before(
                doc,
                anchor,
                ["维度", "GraphRAG", "KnowledGPT（或同类）", "本文"],
                [
                    [
                        "构图粒度",
                        "开放语料实体图与社区摘要",
                        "依赖可查询大规模 KG 与实体链接",
                        "curated 震例+主题库轻量领域图",
                    ],
                    [
                        "阶段感知",
                        "一般不硬编码震前/震中/震后",
                        "侧重图谱接口与提示一体化",
                        "三阶段先验写入调度第一输入",
                    ],
                    [
                        "动态源",
                        "增量机制需另行设计",
                        "通常非实时震情双源",
                        "CEIC 轻量抓取 + USGS FDSN，可关闭",
                    ],
                    [
                        "显式调度",
                        "侧重图检索/摘要策略",
                        "侧重程序化访问 KG",
                        "规则阈值调度，开关可审计",
                    ],
                ],
            )
            continue
        if level == "table4":
            rows: list[list[str]] = []
            for phase in PHASES:
                for metric_key, metric_name in (
                    ("factual_accuracy_pct", "事实一致性(%)"),
                    ("completeness_mean", "要点完整性(1-5)"),
                    ("format_compliance_pct", "格式合规(%)"),
                ):
                    row = [phase, metric_name]
                    for b in BASELINES:
                        row.append(_fmt(by_phase[b][phase][metric_key]))
                    rows.append(row)
            _insert_table_before(
                doc,
                anchor,
                ["阶段", "指标", "B0", "B1", "B2", "B3"],
                rows,
            )
            continue
        if level == "image_fact":
            para = _insert_para_before(anchor, "", "body")
            set_para_text(para, "")
            # clear indent for image
            para.paragraph_format.first_line_indent = Cm(0)
            _add_picture(para, FIG_FACT, width_in=5.2)
            continue
        if level == "image_comp":
            para = _insert_para_before(anchor, "", "body")
            para.paragraph_format.first_line_indent = Cm(0)
            _add_picture(para, FIG_COMP, width_in=5.6)
            continue
        _insert_para_before(anchor, text, level)


def verify(doc: Document, shapes_before_ch54: int) -> dict[str, bool]:
    texts = [p.text.strip() for p in doc.paragraphs]
    i54 = _find_body_idx(doc, "5.4", "系统实验")
    i55 = _find_body_idx(doc, "5.5", "小结")
    section = "\n".join(texts[i54:i55])
    # headings must appear in ascending index order
    order_keys = ["5.4  ", "5.4.1", "5.4.2", "5.4.3", "5.4.4", "5.4.5"]
    order_idxs = []
    for key in order_keys:
        for i, t in enumerate(texts[i54:i55], start=i54):
            if t.startswith(key.strip()) or t.startswith(key):
                order_idxs.append(i)
                break
        else:
            order_idxs.append(-1)
    checks = {
        "has_541": "5.4.1" in section,
        "has_542": "5.4.2" in section,
        "has_543": "5.4.3" in section,
        "has_544": "5.4.4" in section,
        "has_545": "5.4.5" in section,
        "order_ok": order_idxs == sorted(order_idxs) and all(i >= 0 for i in order_idxs),
        "has_table1_title": "表5-1" in section,
        "has_table2_title": "表5-2" in section,
        "has_table3_title": "表5-3" in section,
        "has_table4_title": "表5-4" in section,
        "has_fig11": "图5-11" in section,
        "has_fig12": "图5-12" in section,
        "no_repo_path": "docs/superpowers/" not in section,
        "has_825": "82.5" in section,
        "no_pipe_table": "指标 | B0 | B1" not in section,
        "tables_ge_4": len(doc.tables) >= 4,
        "shapes_ok": len(doc.inline_shapes) >= max(12, shapes_before_ch54),
        "images_in_54": section.count("图5-11") >= 1 and section.count("图5-12") >= 1,
        "figs_exist": FIG_FACT.exists() and FIG_COMP.exists(),
    }
    return checks


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-backup", action="store_true")
    parser.add_argument("--figures-only", action="store_true")
    args = parser.parse_args()

    if not THESIS.exists():
        raise SystemExit(f"找不到论文: {THESIS}")
    if not SUMMARY.exists():
        raise SystemExit(f"找不到汇总: {SUMMARY}")

    data = load_summary()
    generate_figures(data)
    print(f"wrote {FIG_FACT.name}, {FIG_COMP.name}")
    if args.figures_only:
        return

    if not args.no_backup:
        shutil.copy2(THESIS, BACKUP)
        print(f"backup -> {BACKUP.name}")

    doc0 = Document(str(THESIS))
    shapes_before = len(doc0.inline_shapes)

    doc = Document(str(THESIS))
    rewrite_section(doc, data)
    doc.save(str(THESIS))

    doc2 = Document(str(THESIS))
    checks = verify(doc2, shapes_before)
    for k, v in checks.items():
        print(f"  [{'OK' if v else 'FAIL'}] {k}")
    if not all(checks.values()):
        raise SystemExit("验收未全部通过")
    print(f"done. tables={len(doc2.tables)} inline_shapes={len(doc2.inline_shapes)}")


if __name__ == "__main__":
    main()
