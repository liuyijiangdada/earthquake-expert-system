#!/usr/bin/env python3
"""按篇幅对比计划加厚论文：扩写第2–4章，并在5.4插入分阶段事实一致性图5-13。

用法:
  python3 scripts/thicken_thesis_from_gap_plan.py
"""
from __future__ import annotations

import json
import re
import shutil
from copy import deepcopy
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from matplotlib import rcParams

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"
BACKUP = ROOT / "华东师范大学硕士论文.docx.bak-before-thicken"
SUMMARY = ROOT / "data" / "eval" / "table_6_3_summary_60.json"
FIG_DIR = ROOT / "docs" / "superpowers" / "architecture"
FIG_PHASE = FIG_DIR / "fig-5-13-phase-factual-60.png"

BASELINES = ("B0", "B1", "B2", "B3")
PHASES = ("震前", "震中", "震后")
COLORS = ["#9ca3af", "#3b82f6", "#86efac", "#f59e0b"]


def _chinese_font_prop():
    from matplotlib.font_manager import FontProperties

    rcParams["axes.unicode_minus"] = False
    candidates = [
        Path("/System/Library/Fonts/STHeiti Medium.ttc"),
        Path("/System/Library/Fonts/STHeiti Light.ttc"),
        Path("/System/Library/Fonts/Hiragino Sans GB.ttc"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return FontProperties(fname=str(path))
    return FontProperties()


def generate_fig_5_13(data: dict) -> None:
    fp = _chinese_font_prop()
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    by_phase = data["by_phase"]
    x = list(range(len(PHASES)))
    width = 0.18
    fig, ax = plt.subplots(figsize=(8.2, 4.4), dpi=150)
    for i, b in enumerate(BASELINES):
        vals = [float(by_phase[b][p]["factual_accuracy_pct"]) for p in PHASES]
        offs = [xi + (i - 1.5) * width for xi in x]
        bars = ax.bar(offs, vals, width=width, color=COLORS[i], label=b)
        for bar, val in zip(bars, vals):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                val + 1.2,
                f"{val:g}",
                ha="center",
                va="bottom",
                fontsize=8,
                fontproperties=fp,
            )
    ax.set_xticks(x)
    ax.set_xticklabels(PHASES, fontproperties=fp)
    ax.set_ylabel("事实一致性（%）", fontproperties=fp)
    ax.set_ylim(0, 100)
    ax.set_title("图5-13  分阶段事实一致性对比（各阶段20题）", fontproperties=fp)
    leg = ax.legend(frameon=False, ncol=4, loc="upper center", bbox_to_anchor=(0.5, 1.02))
    for text in leg.get_texts():
        text.set_fontproperties(fp)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(FIG_PHASE, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def find_body(doc: Document, pred, *, start_after: int = 140) -> int:
    for i, p in enumerate(doc.paragraphs):
        if i <= start_after:
            continue
        if pred(p.text.strip()):
            return i
    raise RuntimeError(f"anchor not found: {pred}")


def ref_style(doc: Document):
    for p in doc.paragraphs:
        if "事实一致性：衡量生成内容" in p.text and p.runs:
            return deepcopy(p._element.find(qn("w:pPr"))), deepcopy(p.runs[0]._element)
    for p in doc.paragraphs:
        if p.runs and len(p.text) > 40:
            return deepcopy(p._element.find(qn("w:pPr"))), deepcopy(p.runs[0]._element)
    raise RuntimeError("no style ref")


def make_body_para(text: str, pPr, ref_run):
    new_p = OxmlElement("w:p")
    if pPr is not None:
        pPr2 = deepcopy(pPr)
        jc = pPr2.find(qn("w:jc"))
        if jc is not None:
            pPr2.remove(jc)
        new_p.insert(0, pPr2)
    new_r = deepcopy(ref_run)
    t_nodes = list(new_r.iter(qn("w:t")))
    if not t_nodes:
        t = OxmlElement("w:t")
        new_r.append(t)
        t_nodes = [t]
    for k, t in enumerate(t_nodes):
        t.text = text if k == 0 else ""
        if k == 0:
            t.set(qn("xml:space"), "preserve")
    new_p.append(new_r)
    return new_p


def make_center_para(text: str, pPr, ref_run):
    new_p = make_body_para(text, pPr, ref_run)
    pPr_el = new_p.find(qn("w:pPr"))
    if pPr_el is None:
        pPr_el = OxmlElement("w:pPr")
        new_p.insert(0, pPr_el)
    jc = pPr_el.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr_el.append(jc)
    jc.set(qn("w:val"), "center")
    return new_p


def insert_after(paragraph, element) -> None:
    paragraph._element.addnext(element)


def insert_texts_after(doc: Document, anchor_pred, texts: list[str], *, center_first: bool = False) -> None:
    pPr, ref_run = ref_style(doc)
    idx = find_body(doc, anchor_pred)
    anchor = doc.paragraphs[idx]
    # insert in reverse so first text ends up immediately after anchor
    for j, text in enumerate(reversed(texts)):
        use_center = center_first and j == len(texts) - 1
        el = make_center_para(text, pPr, ref_run) if use_center else make_body_para(text, pPr, ref_run)
        insert_after(anchor, el)


def insert_figure_block(doc: Document, after_pred, image: Path, caption: str, note: str) -> None:
    # Skip if body caption already exists (ignore TOC lines with tab+page).
    if any(p.text.strip().startswith(caption.split("（")[0].strip()) and "\t" not in p.text for p in doc.paragraphs):
        # Still refresh embedded image if caption exists in body.
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip().startswith("图5-13") and "\t" not in p.text:
                img_p = doc.paragraphs[i - 1]
                if "w:drawing" in img_p._element.xml:
                    for child in list(img_p._element):
                        if child.tag != qn("w:pPr"):
                            img_p._element.remove(child)
                    img_p.add_run().add_picture(str(image), width=Inches(5.6))
                    print(f"refreshed image for {caption}")
                    return
        print(f"skip existing {caption}")
        return
    pPr, ref_run = ref_style(doc)
    idx = find_body(doc, after_pred)
    # order after insert_after chain: note, caption, image (reverse insert)
    anchor = doc.paragraphs[idx]
    note_el = make_center_para(note, pPr, ref_run)
    cap_el = make_center_para(caption, pPr, ref_run)
    img_p = OxmlElement("w:p")
    img_pPr = OxmlElement("w:pPr")
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    img_pPr.append(jc)
    img_p.append(img_pPr)
    insert_after(anchor, note_el)
    insert_after(anchor, cap_el)
    insert_after(anchor, img_p)
    # bind picture to img_p
    for p in doc.paragraphs:
        if p._element is img_p:
            p.add_run().add_picture(str(image), width=Inches(5.6))
            break


def patch_summary_25(doc: Document) -> None:
    idx = find_body(doc, lambda t: t.startswith("2.5") and "本章小结" in t)
    # next para is summary body
    p = doc.paragraphs[idx + 1]
    if "国内外对照" in p.text:
        return
    new = (
        "本章从四方面梳理了与本文方法相关的技术基础，并补充了领域对照要点："
        "知识图谱构建的模式驱动范式、开放抽取噪声与中文实践；"
        "向量检索与检索增强生成（RAG）的分块、召回、稠密检索思路及与图谱互补原则；"
        "动态震情数据接入中的多源优先级、去重、缓存时效与失败降级；"
        "以及事实一致性、要点完整性与格式合规三类自动评价指标及其局限说明。"
        "系统总体逻辑架构、请求数据流与技术栈映射属于实现层内容，详见第五章 5.2；"
        "本章不再展开系统说明书式架构叙述，以免与实现章重复。"
        "下一章将在此基础上给出静态知识构建与检索的具体方法。"
    )
    if p.runs:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = new


def patch_toc_figure(doc: Document) -> None:
    # insert into 图目录 after 图5-12 line if missing
    for i, p in enumerate(doc.paragraphs[:160]):
        t = p.text.strip()
        if t.startswith("图5-12") and "\t" in t:
            if any(x.text.strip().startswith("图5-13") for x in doc.paragraphs[:160]):
                return
            pPr, ref_run = ref_style(doc)
            # keep TOC-like text
            line = "图5-13  分阶段事实一致性对比（各阶段20题）\t—"
            el = make_body_para(line, deepcopy(p._element.find(qn("w:pPr"))), deepcopy(p.runs[0]._element) if p.runs else ref_run)
            insert_after(p, el)
            return


def main() -> None:
    shutil.copy2(THESIS, BACKUP)
    data = json.loads(SUMMARY.read_text(encoding="utf-8"))
    generate_fig_5_13(data)
    print("generated", FIG_PHASE)

    doc = Document(str(THESIS))

    # ----- Chapter 2 expansions -----
    insert_texts_after(
        doc,
        lambda t: t.startswith("在应急垂直领域，常见做法是预定义节点类型"),
        [
            "从国内外对照看，徐增林等[27]、李涓子等[28]分别从技术体系与学科综述角度强调知识图谱的表示学习、融合与应用链条；肖仰华[38]、王昊奋等[37]则更侧重工程化构建与中文场景落地。"
            "对地震应急而言，完全依赖开放域抽取易把非正式帖文中的震级传言写入库，因此本文采取“小规模权威整理 + 模式约束导入”的策略：先固定节点/关系类型，再批量写入震例与主题，最后用抽样图查询自检。"
            "该策略牺牲了自动覆盖广度，换取字段可核验与消融实验可复现，与第一章对静态知识的操作性定义一致。",
        ],
    )
    insert_texts_after(
        doc,
        lambda t: "本文第三章将据此设计主题级向量索引及其与图谱的联合注入顺序" in t,
        [
            "在检索器形态上，Karpukhin 等[29]提出的稠密段落检索（DPR）表明：以双塔编码计算问句—段落相似度，可在开放域问答中显著优于传统稀疏检索；Reimers 等[32]的 Sentence-BERT 则为中文句向量提供了常用编码器训练范式。"
            "本文实验侧采用轻量中文嵌入与主题级索引，目标不是追赶开放域排行榜，而是保证应急主题步骤链完整、并与图谱字段分区注入。"
            "需要指出：Edge 等[2]的 GraphRAG 通过社区摘要提升长文档全局问答，但离线构图与增量更新成本较高；在已有人工主题库与固定模式图谱的前提下，本文不重建社区图，而把“可关闭的动态源 + 阶段先验”作为差异化重点，相关设计对照见第五章表5-3。",
        ],
    )
    insert_texts_after(
        doc,
        lambda t: "具体双源合并、调度开关与消融关闭策略将在第四章展开" in t,
        [
            "工程上，动态接入还需区分“目录查询”与“速报展示”两类语义：前者回答近时段区域是否有事件，后者强调最新一条参数的可读性。"
            "合并时建议记录来源标签与拉取时间戳，以便界面侧展示可靠性提示；当国内源与国际源对同一事件给出接近但不等的震级时，应以预先声明的优先级与去重半径裁决，而不是让生成模型自行折中。"
            "此外，动态段写入提示前应做最小字段校验（时间、震级、地名非空），避免把半解析失败的页面噪声注入模型。上述约束为第四章规则阈值调度提供可探测信号。",
        ],
    )
    insert_texts_after(
        doc,
        lambda t: "论文表述因此坚持描述性对比，避免在缺乏严格假设检验时宣称“全面显著优于”" in t,
        [
            "除自动代理指标外，应急问答评价还可引入人工维度：步骤可执行性、口径稳健性（是否给出过度绝对化承诺）、来源可追溯性等。"
            "受时间与标注成本限制，本文未开展正式双人专家打分，第五章将诚实标注“可用性问卷待开展”，并把结论边界限制在60题离线消融与自动评分口径之内。"
            "该安排与“先把可复现基线跑通，再补强人工评测”的研究节奏一致，也避免用小样本访谈替代系统评价。",
        ],
    )
    patch_summary_25(doc)

    # ----- Chapter 3 expansions -----
    insert_texts_after(
        doc,
        lambda t: "但不改变本章“轻量权威库 + 主题级检索”的核心设定" in t,
        [
            "为说明静态双源如何落到一次问答，给出两类典型查询路径。"
            "路径甲（结构化为主）：问句含明确区域实体（如“四川近年来有哪些典型地震”）→ 图谱按区域返回震例摘要（震级、时间、深度）→ 检索段仅补充“建议关注主题”中的避险/政策要点，避免重复罗列震级。"
            "路径乙（语义为主）：问句口语化且缺区域（如“刚才晃了一下室内怎么办”）→ 图谱可能无区域命中而写“（无）”→ 向量检索召回室内避险主题步骤并保持顺序 → 第四章若判定为震中且动态可用，再追加【动态信息】。"
            "失败情形同样需要显式处理：主题库外的冷门问法可能导致检索得分偏低，此时应缩短参考资料或追加“静态依据不足”；震例库未覆盖的新发地震则依赖动态源，而不能伪造历史节点。"
            "把成功路径与失败降级写清楚，是为了让第五章消融在关闭动态后仍能单独检验静态贡献，并避免把“库外事件”误判为模型能力不足。",
        ],
    )

    # ----- Chapter 4 expansions -----
    insert_texts_after(
        doc,
        lambda t: t.startswith("该伪代码透明可审计"),
        [
            "下面用三个阶段各举一例，说明规则如何改变提示组成（均为方法示意，数值以探测结果为准）。"
            "震前问句“家庭应急包要准备什么”：阶段=震前，动态强制关闭；图谱可返回准备类主题关联，检索注入应急包步骤；阶段指令强调勿编造即时震情。"
            "震中问句“刚才地震多大，震中在哪”：阶段=震中且需要动态；若动态可用则【动态信息】优先，静态侧补充“权威信息获取渠道”主题；若动态不可用则写明暂不可用并禁止编造震级。"
            "震后问句“房子裂了还能住吗”：阶段=震后；图谱/检索常开，动态仅当问句含“最新余震”等时效线索时开启；政策类主题附加“以省级最新方案为准”的约束，降低过时条文被当成现行规定的风险。"
            "上述例子表明：调度器改变的是“写哪些段、加哪些约束”，而不是改写知识库内容本身，从而保持证据来源可审计。",
        ],
    )
    insert_texts_after(
        doc,
        lambda t: "空段显式化——无命中写“（无）”或“本路径已关闭/暂不可用”，避免模型误以为有证据。" in t,
        [
            "与 GraphRAG、KnowledGPT 等通用方案对照时，应区分“能否同设定复现分数”与“设计假设是否匹配”。"
            "GraphRAG 擅长从大规模无结构语料归纳社区摘要，但本文静态侧已是人工主题库与固定模式图谱，再造社区图收益有限且增加离线成本；KnowledGPT 强调程序化访问知识库，与本文“分区提示注入”部分同向，但其开放域实体链接全栈并非震灾速报场景的必需品。"
            "本文差异化在于：把震前/震中/震后先验写成可关闭规则，并把 CEIC/USGS 动态源纳入同一消融协议。第五章表5-3仅做设计维度对照，不报告未复现流水线的虚构分数；若未来具备同问集复现条件，再补充数值对比更为妥当。",
        ],
    )

    # ----- Chapter 5.4: fig 5-13 + error typology -----
    insert_figure_block(
        doc,
        lambda t: t.startswith("由表5-4可见"),
        FIG_PHASE,
        "图5-13  分阶段事实一致性对比（各阶段20题）",
        "该图按震前/震中/震后对比 B0–B3 的事实一致性自动评分。",
    )
    # Insert error analysis after the B2 root-cause paragraph (or after 表5-4 discussion chain)
    insert_texts_after(
        doc,
        lambda t: "表明检索应与结构化图谱联用，而非否定检索路径本身。" in t,
        [
            "在自动评分口径下，还可将失误粗分为三类，便于后续人工抽检。"
            "类型A（字段漂移）：回答中的震级/时间/地点与图谱段不一致，或在无动态段时编造即时参数——B0 更常见，B1/B3 因图谱约束而减少。"
            "类型B（主题错配）：检索命中邻近主题但关键步骤缺失或顺序打乱——B2 相对高发，也是其事实分偏低的机制来源之一。"
            "类型C（阶段串话）：震中问句展开大段震后重建政策，或震前问句夹带未确认的速报口吻——规则阈值与阶段指令主要用于抑制此类串话。"
            "需要强调：上述分型是对自动评分与案例观察的归纳，尚未经双人标注的一致性检验；其作用是指导错误分析与系统迭代，而不是替代表5-2的汇总结论。"
            "图5-13 从分阶段视角补充说明：B3 在三个阶段均保持相对最高的事实一致性，其中震后子集（88.0%）提升更明显，与政策/心理类主题检索及图谱字段叠加有关；震中子集仍受动态关闭（离线消融）约束，分数解释必须回到实验协议边界。",
        ],
    )

    # Update inline mention in 5.4.2 if useful
    for p in doc.paragraphs:
        if "图5-11 给出四基线事实一致性柱状对比；图5-12 给出要点完整性与格式合规对照。" in p.text:
            if "图5-13" not in p.text:
                if p.runs:
                    p.runs[0].text = p.text.replace(
                        "图5-12 给出要点完整性与格式合规对照。",
                        "图5-12 给出要点完整性与格式合规对照；分阶段事实一致性见图5-13。",
                    )
                    for r in p.runs[1:]:
                        r.text = ""
            break

    patch_toc_figure(doc)

    # Refresh 5.5 summary to mention fig5-13 / error types lightly
    for p in doc.paragraphs:
        if p.text.strip().startswith("本章完成了系统需求分析、总体设计"):
            if "图5-13" in p.text:
                break
            new = (
                p.text.strip().rstrip("。")
                + "；并以图5-13展示分阶段事实一致性，归纳字段漂移、主题错配与阶段串话三类失误形态（自动口径，待人工复核）。"
            )
            if p.runs:
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
            break

    doc.save(str(THESIS))

    # verify
    doc2 = Document(str(THESIS))
    blob = "\n".join(p.text for p in doc2.paragraphs)
    checks = {
        "has_fig513_caption": "图5-13  分阶段事实一致性对比" in blob,
        "has_fig513_file": FIG_PHASE.exists(),
        "has_error_types": "类型A（字段漂移）" in blob,
        "has_ch2_contrast": "小规模权威整理 + 模式约束导入" in blob,
        "has_ch3_paths": "路径甲（结构化为主）" in blob,
        "has_ch4_examples": "震前问句“家庭应急包要准备什么”" in blob,
        "has_graphrag_discuss": "不报告未复现流水线的虚构分数" in blob,
        "inline_shapes_ge_14": len(doc2.inline_shapes) >= 14,
    }
    print("checks:", checks)
    if not all(checks.values()):
        missing = [k for k, v in checks.items() if not v]
        raise SystemExit(f"verification failed: {missing}")
    print("OK; backup:", BACKUP)


if __name__ == "__main__":
    main()
