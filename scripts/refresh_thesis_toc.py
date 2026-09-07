#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按正文标题重写「目录 / 图目录 / 表目录」，并为条目建立书签超链接跳转。

用法:
  .venv/bin/python scripts/refresh_thesis_toc.py
  .venv/bin/python scripts/refresh_thesis_toc.py --no-backup
"""
from __future__ import annotations

import argparse
import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"
BACKUP = ROOT / "华东师范大学硕士论文.docx.bak-before-toc-links"

CHAPTER_RE = re.compile(r"^第[一二三四五六七八九十]+章(\s|$)")
SECTION_RE = re.compile(r"^\d+\.\d+(?!\.\d)")
SUBSECTION_RE = re.compile(r"^\d+\.\d+\.\d+")
FIG_CAP_RE = re.compile(r"^图\s*(\d+)\s*[-–—]\s*(\d+)\s+(.+)$")
TAB_CAP_RE = re.compile(r"^表\s*(\d+)\s*[-–—]\s*(\d+)\s+(.+)$")
END_MATTER = ("参考文献", "致谢", "附录")
BOOKMARK_PREFIX = "_TocThesis"


def _set_run_font(run, *, east_asia="宋体", western="Times New Roman", size_pt=12, bold=False) -> None:
    run.font.name = western
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:ascii"), western)
    rFonts.set(qn("w:hAnsi"), western)


def set_para_text(para: Paragraph, text: str) -> None:
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def _insert_para_before(anchor: Paragraph, text: str = "") -> Paragraph:
    new_p = deepcopy(anchor._p)
    for child in list(new_p):
        local = child.tag.split("}")[-1]
        if local in ("drawing", "pict", "r", "hyperlink", "bookmarkStart", "bookmarkEnd"):
            new_p.remove(child)
    anchor._p.addprevious(new_p)
    para = Paragraph(new_p, anchor._parent)
    set_para_text(para, text)
    return para


def _delete_paragraph(para: Paragraph) -> None:
    el = para._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def _find_exact(doc: Document, title: str) -> int:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == title:
            return i
    raise KeyError(title)


def _find_body_chapter1(doc: Document) -> int:
    idxs = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if "\t" in p.text:
            continue
        if t.startswith("第一章") and "绪论" in t:
            idxs.append(i)
    if not idxs:
        raise KeyError("正文第一章")
    return idxs[-1]


def fix_chapter2_title(doc: Document) -> None:
    for p in doc.paragraphs:
        t = p.text.strip()
        if "\t" in p.text:
            continue
        if t == "相关技术与理论基础" or (t.startswith("第二章") and "相关技术" in t):
            set_para_text(p, "第二章  相关技术与理论基础")
            _format_heading(p, 0)
            return


def _format_heading(para: Paragraph, level: int) -> None:
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(0)
    if level == 0:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=16, bold=True)
    elif level == 1:
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=14, bold=True)
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=12, bold=True)


def _heading_level(text: str) -> int | None:
    t = text.strip()
    if not t or len(t) > 80:
        return None
    if CHAPTER_RE.match(t) or t in END_MATTER:
        return 0
    if SUBSECTION_RE.match(t):
        return 2
    if SECTION_RE.match(t):
        return 1
    return None


def clear_toc_bookmarks(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag == qn("w:bookmarkStart") and (child.get(qn("w:name")) or "").startswith(BOOKMARK_PREFIX):
            bid = child.get(qn("w:id"))
            paragraph._p.remove(child)
            for child2 in list(paragraph._p):
                if child2.tag == qn("w:bookmarkEnd") and child2.get(qn("w:id")) == bid:
                    paragraph._p.remove(child2)


def add_bookmark(paragraph: Paragraph, name: str, bookmark_id: int) -> None:
    clear_toc_bookmarks(paragraph)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p = paragraph._p
    if len(p) and p[0].tag == qn("w:pPr"):
        p.insert(1, start)
    else:
        p.insert(0, start)
    p.append(end)


def collect_heading_targets(doc: Document, body_start: int) -> list[tuple[int, str, Paragraph]]:
    out: list[tuple[int, str, Paragraph]] = []
    for p in doc.paragraphs[body_start:]:
        if "\t" in p.text:
            continue
        t = p.text.strip()
        lvl = _heading_level(t)
        if lvl is None:
            continue
        out.append((lvl, t, p))
    return out


def collect_caption_targets(doc: Document, body_start: int) -> tuple[list[tuple[str, Paragraph]], list[tuple[str, Paragraph]]]:
    figs: list[tuple[str, Paragraph]] = []
    tabs: list[tuple[str, Paragraph]] = []
    seen_f: set[str] = set()
    seen_t: set[str] = set()
    for p in doc.paragraphs[body_start:]:
        t = p.text.strip()
        m = FIG_CAP_RE.match(t)
        if m:
            rest = m.group(3).strip()
            if rest.startswith("展示") or len(rest) > 60:
                continue
            key = f"{m.group(1)}-{m.group(2)}"
            if key in seen_f:
                continue
            seen_f.add(key)
            figs.append((f"图{m.group(1)}-{m.group(2)}  {rest}", p))
            continue
        m = TAB_CAP_RE.match(t)
        if m:
            rest = m.group(3).strip()
            if rest.startswith("汇总") or len(rest) > 55:
                continue
            key = f"{m.group(1)}-{m.group(2)}"
            if key in seen_t:
                continue
            seen_t.add(key)
            tabs.append((f"表{m.group(1)}-{m.group(2)}  {rest}", p))
    return figs, tabs


def _style_toc_title(para: Paragraph) -> None:
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(0)
    pf.left_indent = Cm(0)
    pf.space_before = Pt(12)
    pf.space_after = Pt(12)
    for run in para.runs:
        _set_run_font(run, east_asia="黑体", size_pt=16, bold=True)


def _style_toc_entry(para: Paragraph, level: int, tab_pos_cm: float, doc: Document | None = None) -> None:
    """套用 Word 内置 toc 1/2/3，对齐学校模板目录层级与点线页码。"""
    if doc is not None:
        style_name = {0: "toc 1", 1: "toc 2", 2: "toc 3"}.get(level, "toc 1")
        try:
            para.style = doc.styles[style_name]
        except KeyError:
            pass
    indents = {0: Cm(0), 1: Cm(0.74), 2: Cm(1.48)}
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(0)
    # toc 样式通常自带层级缩进；无样式时用缩进保底
    st_name = str(getattr(para.style, "name", "") or "")
    if not st_name.startswith("toc"):
        pf.left_indent = indents.get(level, Cm(0))
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    try:
        pf.tab_stops.clear_all()
    except Exception:
        pass
    pf.tab_stops.add_tab_stop(Cm(tab_pos_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def ensure_black_hyperlink_styles(doc: Document) -> None:
    """对齐戚媛媛答辩版目录：黑字、无下划线；避免访问后变绿。"""
    for style_name in ("Hyperlink", "FollowedHyperlink"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        try:
            style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        except Exception:
            pass
        try:
            style.font.underline = False
        except Exception:
            pass

    # 直接改 styles.xml，覆盖主题色/已访问链接色
    styles_el = doc.styles.element
    for style_id, display in (("Hyperlink", "Hyperlink"), ("FollowedHyperlink", "FollowedHyperlink")):
        st = None
        for candidate in styles_el.findall(qn("w:style")):
            if candidate.get(qn("w:styleId")) == style_id:
                st = candidate
                break
        if st is None:
            st = OxmlElement("w:style")
            st.set(qn("w:type"), "character")
            st.set(qn("w:styleId"), style_id)
            name_el = OxmlElement("w:name")
            name_el.set(qn("w:val"), display)
            st.append(name_el)
            styles_el.append(st)
        rPr = st.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            st.append(rPr)
        for tag in ("w:color", "w:u", "w:sz", "w:szCs"):
            node = rPr.find(qn(tag))
            if node is not None:
                rPr.remove(node)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        rPr.append(color)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "none")
        rPr.append(u)


def _add_hyperlink_run(paragraph: Paragraph, text: str, bookmark: str) -> None:
    """Insert a Word hyperlink anchored to an internal bookmark (black, no underline)."""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark)
    hyperlink.set(qn("w:history"), "1")

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    # 显式覆盖 Hyperlink 字符样式，防止蓝/绿
    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), "Hyperlink")
    rPr.append(rStyle)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "宋体")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")  # 12pt
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "24")
    rPr.append(szCs)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "none")
    rPr.append(u)
    run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_tab_run(paragraph: Paragraph) -> None:
    run = paragraph.add_run("\t")
    _set_run_font(run, size_pt=12, bold=False)


def _add_pageref_field(paragraph: Paragraph, bookmark: str) -> None:
    run = paragraph.add_run()
    r = run._r
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" PAGEREF {bookmark} \\h "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "—"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(text)
    r.append(fld_end)
    _set_run_font(run, size_pt=12, bold=False)


def _clear_paragraph_content(para: Paragraph) -> None:
    for child in list(para._p):
        local = child.tag.split("}")[-1]
        if local != "pPr":
            para._p.remove(child)


def fill_toc_entry(
    para: Paragraph, title: str, bookmark: str, level: int, tab_pos: float, doc: Document | None = None
) -> None:
    _clear_paragraph_content(para)
    _style_toc_entry(para, level, tab_pos, doc=doc)
    _add_hyperlink_run(para, title, bookmark)
    _add_tab_run(para)
    _add_pageref_field(para, bookmark)


def rebuild_front_lists(
    doc: Document,
    headings: list[tuple[int, str, Paragraph]],
    figs: list[tuple[str, Paragraph]],
    tabs: list[tuple[str, Paragraph]],
) -> int:
    ensure_black_hyperlink_styles(doc)
    # assign bookmarks on body targets first
    bid = 3000
    targets: list[tuple[int, str, str]] = []  # level, title, bookmark
    for lvl, title, para in headings:
        name = f"{BOOKMARK_PREFIX}{bid - 3000:03d}"
        add_bookmark(para, name, bid)
        targets.append((lvl, title, name))
        bid += 1

    fig_targets: list[tuple[str, str]] = []
    for title, para in figs:
        name = f"{BOOKMARK_PREFIX}{bid - 3000:03d}"
        add_bookmark(para, name, bid)
        fig_targets.append((title, name))
        bid += 1

    tab_targets: list[tuple[str, str]] = []
    for title, para in tabs:
        name = f"{BOOKMARK_PREFIX}{bid - 3000:03d}"
        add_bookmark(para, name, bid)
        tab_targets.append((title, name))
        bid += 1

    toc_i = _find_exact(doc, "目录")
    body_i = _find_body_chapter1(doc)
    for p in list(doc.paragraphs[toc_i + 1 : body_i]):
        _delete_paragraph(p)

    toc_title = doc.paragraphs[_find_exact(doc, "目录")]
    body_ch1 = doc.paragraphs[_find_body_chapter1(doc)]
    _style_toc_title(toc_title)

    sec = doc.sections[0]
    usable_cm = float((sec.page_width - sec.left_margin - sec.right_margin) / 360000)
    tab_pos = max(usable_cm - 0.2, 12.0)

    # Build entry list then insert forward before body
    # level=-1: 图目录/表目录分节标题；0/1/2: 正文标题（套 toc 样式）；10: 图/表条目（Normal+点线）
    entries = []
    for lvl, title, bm in targets:
        entries.append((lvl, title, bm))
    entries.append((-1, "图目录", None))
    for title, bm in fig_targets:
        entries.append((10, title, bm))
    entries.append((-1, "表目录", None))
    for title, bm in tab_targets:
        entries.append((10, title, bm))

    for level, title, bm in entries:
        para = _insert_para_before(body_ch1, "")
        if level == -1:
            set_para_text(para, title or "")
            _style_toc_title(para)
        else:
            assert title and bm
            if level == 10:
                fill_toc_entry(para, title, bm, 0, tab_pos, doc=None)
            else:
                fill_toc_entry(para, title, bm, level, tab_pos, doc=doc)

    return len(targets) + len(fig_targets) + len(tab_targets)


def verify(doc: Document) -> dict[str, bool]:
    texts = [p.text.strip() for p in doc.paragraphs]
    toc_i = texts.index("目录")
    body_i = _find_body_chapter1(doc)
    zone_paras = doc.paragraphs[toc_i:body_i]
    zone = [p.text.strip() for p in zone_paras]
    joined = "\n".join(zone)
    hyperlink_n = 0
    pageref_n = 0
    for p in zone_paras:
        xml = p._p.xml
        hyperlink_n += xml.count("w:hyperlink")
        pageref_n += xml.count("PAGEREF")
    body_bm = 0
    for p in doc.paragraphs[body_i:]:
        for child in p._p:
            if child.tag == qn("w:bookmarkStart") and (child.get(qn("w:name")) or "").startswith(BOOKMARK_PREFIX):
                body_bm += 1
    return {
        "has_toc": zone and zone[0] == "目录",
        "has_ch1": any(t.startswith("第一章") for t in zone),
        "has_ch2": any(t.startswith("第二章") for t in zone),
        "has_ch5": any(t.startswith("第五章") for t in zone),
        "has_ch6": any(t.startswith("第六章") for t in zone),
        "has_54": any(t.startswith("5.4") for t in zone),
        "has_fig_toc": "图目录" in zone,
        "has_tab_toc": "表目录" in zone,
        "has_fig511": "图5-11" in joined,
        "has_tab54": "表5-4" in joined,
        "hyperlinks_ge_20": hyperlink_n >= 20,
        "pageref_ge_20": pageref_n >= 20,
        "body_bookmarks_ge_20": body_bm >= 20,
    }


def update_fields_with_word(docx_path: Path) -> bool:
    """用 Microsoft Word 更新 PAGEREF 域，写入真实页码。"""
    import subprocess

    posix = str(docx_path.resolve())
    script = f'''
set thePath to POSIX file "{posix}"
tell application "Microsoft Word"
    activate
    open thePath
    delay 2
    set theDoc to active document
    try
        update fields theDoc
    end try
    try
        tell theDoc
            update
        end tell
    end try
    -- 再扫一遍 story，尽量刷新目录 PAGEREF
    try
        set storyCount to count of stories of theDoc
        repeat with i from 1 to storyCount
            try
                update fields (get story i of theDoc)
            end try
        end repeat
    end try
    save theDoc
    delay 1
    close theDoc saving yes
end tell
'''
    try:
        r = subprocess.run(
            ["osascript", "-e", script],
            capture_output=True,
            text=True,
            timeout=180,
        )
        if r.returncode != 0:
            print("Word 更新域失败:", (r.stderr or r.stdout)[:500])
            return False
        print("已用 Microsoft Word 更新域并保存页码")
        return True
    except Exception as e:
        print(f"调用 Word 失败: {e}")
        return False


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-backup", action="store_true")
    ap.add_argument("--skip-word-update", action="store_true", help="不调用 Word 刷新页码")
    args = ap.parse_args()

    if not THESIS.exists():
        raise SystemExit(f"找不到 {THESIS}")

    if not args.no_backup:
        shutil.copy2(THESIS, BACKUP)
        print(f"backup -> {BACKUP.name}")

    doc = Document(str(THESIS))
    fix_chapter2_title(doc)
    body_i = _find_body_chapter1(doc)
    headings = collect_heading_targets(doc, body_i)
    figs, tabs = collect_caption_targets(doc, body_i)

    print(f"headings={len(headings)} figs={len(figs)} tables={len(tabs)}")
    n = rebuild_front_lists(doc, headings, figs, tabs)
    doc.save(str(THESIS))

    doc2 = Document(str(THESIS))
    checks = verify(doc2)
    for k, v in checks.items():
        print(f"  [{'OK' if v else 'FAIL'}] {k}")
    if not all(checks.values()):
        raise SystemExit("目录验收未全部通过")

    print(f"linked entries: {n}")
    if not args.skip_word_update:
        ok = update_fields_with_word(THESIS)
        if ok:
            # 抽查页码是否已非 —
            doc3 = Document(str(THESIS))
            toc_i = [p.text.strip() for p in doc3.paragraphs].index("目录")
            sample = [doc3.paragraphs[toc_i + k].text for k in range(1, 6)]
            print("sample TOC after Word:", sample)
        else:
            print("页码未自动写入。请用 Word 打开后全选 → 右键「更新域」。")
    print("完成。目录为黑字可跳转（对齐戚媛媛风格），页码由 Word 域刷新。")


if __name__ == "__main__":
    main()

