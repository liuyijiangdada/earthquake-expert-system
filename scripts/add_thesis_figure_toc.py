#!/usr/bin/env python3
"""在目录后插入图目录（点线页码 + PAGEREF）。"""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "华东师范大学硕士论文.docx"

FIG_RE = re.compile(r"^图\s*(\d+)[-.．](\d+)\s*(.*)$")
CHAPTER_RE = re.compile(r"^第[一二三四五六七八九十百]+章")


def set_run_font(run, *, east_asia="宋体", western="Times New Roman", size_pt=12, bold=False):
    run.font.name = western
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:ascii"), western)
    rFonts.set(qn("w:hAnsi"), western)


def insert_paragraph_after(para: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    para._p.addnext(new_p)
    new_para = Paragraph(new_p, para._parent)
    if text:
        new_para.add_run(text)
    return new_para


def add_page_break(para: Paragraph) -> None:
    run = para.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def clear_bookmarks_with_prefix(paragraph: Paragraph, prefix: str) -> None:
    for child in list(paragraph._p):
        if child.tag == qn("w:bookmarkStart") and (child.get(qn("w:name")) or "").startswith(prefix):
            bid = child.get(qn("w:id"))
            paragraph._p.remove(child)
            for child2 in list(paragraph._p):
                if child2.tag == qn("w:bookmarkEnd") and child2.get(qn("w:id")) == bid:
                    paragraph._p.remove(child2)


def add_bookmark(paragraph: Paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p = paragraph._p
    if len(p):
        p.insert(0 if p[0].tag != qn("w:pPr") else 1, start)
    else:
        p.append(start)
    p.append(end)


def add_pageref_field(paragraph: Paragraph, bookmark_name: str) -> None:
    run = paragraph.add_run()
    r = run._r
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" PAGEREF {bookmark_name} \\h "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "—"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(text)
    r.append(fld_end)
    set_run_font(run, size_pt=12, bold=False)


def style_title(para: Paragraph) -> None:
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(12)
    pf.space_after = Pt(12)
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    for run in para.runs:
        set_run_font(run, east_asia="黑体", western="Times New Roman", size_pt=16, bold=True)


def style_entry(para: Paragraph, tab_pos_cm: float) -> None:
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.left_indent = Cm(0)
    pf.tab_stops.add_tab_stop(Cm(tab_pos_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def format_fig_label(ch: str, seq: str, title: str) -> str:
    title = re.sub(r"\s+", " ", title.strip())
    return f"图 {ch}.{seq}  {title}"


def remove_existing_fig_toc(doc: Document) -> None:
    seen = False
    to_remove = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "图目录":
            seen = True
            to_remove.append(p._p)
            continue
        if not seen:
            continue
        if CHAPTER_RE.match(t):
            break
        to_remove.append(p._p)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def find_toc_ack_para(doc: Document) -> Paragraph | None:
    in_toc = False
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "目录":
            in_toc = True
            continue
        if in_toc and t.startswith("致谢"):
            return p
    return None


def find_body_chapter1(doc: Document) -> Paragraph | None:
    seen_toc_ack = False
    in_toc = False
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "目录":
            in_toc = True
            continue
        if in_toc and t.startswith("致谢"):
            seen_toc_ack = True
            continue
        if seen_toc_ack and CHAPTER_RE.match(t):
            return p
    return None


def main() -> None:
    backup = ROOT / f"华东师范大学硕士论文.backup.{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    shutil.copy2(SRC, backup)
    doc = Document(str(SRC))
    remove_existing_fig_toc(doc)

    figures: list[tuple[Paragraph, str, str, str]] = []
    passed_body = False
    seen_toc_ack = False
    in_toc = False
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "目录":
            in_toc = True
            continue
        if in_toc and t.startswith("致谢"):
            seen_toc_ack = True
            continue
        if seen_toc_ack and CHAPTER_RE.match(t):
            passed_body = True
        if not passed_body:
            continue
        m = FIG_RE.match(t)
        if m:
            figures.append((p, m.group(1), m.group(2), m.group(3)))

    if not figures:
        raise SystemExit("未找到图题")

    for i, (p, _, _, _) in enumerate(figures):
        clear_bookmarks_with_prefix(p, "_TocFig")
        add_bookmark(p, f"_TocFig{i:03d}", 3000 + i)

    toc_ack = find_toc_ack_para(doc)
    body_ch1 = find_body_chapter1(doc)
    if toc_ack is None or body_ch1 is None:
        raise SystemExit("无法定位目录致谢或正文第一章")

    to_remove = []
    crossed = False
    for p in doc.paragraphs:
        if p._p is toc_ack._p:
            crossed = True
            continue
        if p._p is body_ch1._p:
            break
        if crossed:
            to_remove.append(p._p)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)

    sec = doc.sections[0]
    usable_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000
    tab_pos = max(usable_cm - 0.1, 12.0)

    cursor = toc_ack
    pb1 = insert_paragraph_after(cursor, "")
    add_page_break(pb1)
    cursor = pb1

    title_p = insert_paragraph_after(cursor, "图目录")
    style_title(title_p)
    cursor = title_p

    prev_ch = None
    for i, (_, ch, seq, title) in enumerate(figures):
        if prev_ch is not None and ch != prev_ch:
            blank = insert_paragraph_after(cursor, "")
            blank.paragraph_format.space_before = Pt(0)
            blank.paragraph_format.space_after = Pt(0)
            cursor = blank
        entry = insert_paragraph_after(cursor, "")
        style_entry(entry, tab_pos)
        run = entry.add_run(format_fig_label(ch, seq, title))
        set_run_font(run, east_asia="宋体", western="Times New Roman", size_pt=12, bold=False)
        entry.add_run("\t")
        add_pageref_field(entry, f"_TocFig{i:03d}")
        cursor = entry
        prev_ch = ch

    pb2 = insert_paragraph_after(cursor, "")
    add_page_break(pb2)
    doc.save(str(SRC))
    print(f"backup: {backup.name}")
    print(f"figures: {len(figures)}")
    print("请用 Word/WPS 打开后全选并更新域，以刷新页码。")


if __name__ == "__main__":
    main()
