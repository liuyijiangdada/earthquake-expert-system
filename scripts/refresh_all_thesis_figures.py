#!/usr/bin/env python3
"""重新生成论文中非「系统页面截图」的全部插图，并写回 docx。

包含：
  图3-1 总体方法框架图
  图4-1 端到端推理流程
  图5-1 系统总体逻辑架构图
  图5-2 问答请求数据流图
  图5-3 技术栈与模块映射关系
  图5-11 / 图5-12 实验结果对比图

排除（不改动）：
  图5-4 ～ 图5-10 系统/移动端界面示例截图
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"
BACKUP = ROOT / "华东师范大学硕士论文.docx.bak-before-refresh-all-figures"
ARCH = ROOT / "docs" / "superpowers" / "architecture"

sys.path.insert(0, str(ROOT / "scripts"))

import generate_e2e_inference_flow as e2e  # noqa: E402
import generate_fig_3_1_method_framework as fig31  # noqa: E402
import generate_query_flow as qf  # noqa: E402
import generate_system_architecture as sa  # noqa: E402
import generate_tech_stack_table as ts  # noqa: E402
import polish_thesis_ch5_4 as ch54  # noqa: E402


# (题注前缀匹配关键词, 图片路径, 宽度英寸)
REPLACEMENTS = [
    ("图3-1", ARCH / "fig-3-1-method-framework.png", 5.8),
    ("图4-1", ARCH / "fig-6-1-e2e-inference-flow.png", 5.4),
    ("图5-1", ARCH / "fig-5-1-system-architecture.png", 5.8),
    ("图5-2", ARCH / "fig-5-2-query-flow.png", 4.6),
    ("图5-3", ARCH / "fig-5-3-tech-stack.png", 5.8),
    ("图5-11", ARCH / "fig-5-11-factual-60.png", 5.2),
    ("图5-12", ARCH / "fig-5-12-completeness-format-60.png", 5.6),
]


def has_drawing(p: Paragraph) -> bool:
    return bool(p._element.xpath('.//*[local-name()="drawing"]'))


def clear_para_media(para: Paragraph) -> None:
    for child in list(para._element):
        local = child.tag.split("}")[-1]
        if local in ("drawing", "pict"):
            para._element.remove(child)
        elif local == "r":
            for sub in list(child):
                sl = sub.tag.split("}")[-1]
                if sl in ("drawing", "pict"):
                    child.remove(sub)
            # drop empty runs after clearing drawings
            if len(list(child)) == 0 or (
                all(
                    (s.tag.split("}")[-1] in ("rPr",) or not (s.text or "").strip())
                    for s in child
                )
                and not any(s.tag.split("}")[-1] in ("drawing", "pict", "t") for s in child)
            ):
                # keep structure simple: remove whole run if no text left
                texts = [s for s in child if s.tag.split("}")[-1] == "t"]
                if not texts or all(not (t.text or "").strip() for t in texts):
                    # only remove if no meaningful text
                    has_text = any(
                        (s.text or "").strip()
                        for s in child
                        if s.tag.split("}")[-1] == "t"
                    )
                    if not has_text:
                        para._element.remove(child)


def find_caption_index(paras, prefix: str) -> int:
    # Prefer body captions (no TOC tab), later in document
    soft = None
    for i, p in enumerate(paras):
        t = p.text.strip()
        if not t.startswith(prefix):
            continue
        if "\t" in t:
            continue  # skip TOC
        soft = i
        # prefer exact body caption paragraphs
        if any(k in t for k in ("框架", "流程", "架构", "数据流", "技术栈", "事实", "完整性", "映射")):
            return i
    if soft is not None:
        return soft
    raise KeyError(f"找不到题注: {prefix}")


def replace_image_before_caption(doc: Document, caption_prefix: str, image: Path, width_in: float) -> None:
    paras = list(doc.paragraphs)
    cap_i = find_caption_index(paras, caption_prefix)

    target = None
    for j in range(cap_i - 1, max(-1, cap_i - 6), -1):
        p = paras[j]
        if has_drawing(p) or not p.text.strip():
            target = p
            if has_drawing(p):
                break
    if target is None:
        target = paras[cap_i].insert_paragraph_before("")

    clear_para_media(target)
    # also clear leftover text in image para
    for run in target.runs:
        run.text = ""

    run = target.add_run()
    run.add_picture(str(image), width=Inches(width_in))
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target.paragraph_format.first_line_indent = Cm(0)
    print(f"  已替换 {caption_prefix} <- {image.name} ({image.stat().st_size // 1024} KB)")


def generate_all() -> list[Path]:
    print("1/7 图3-1 总体方法框架…")
    p31 = fig31.generate_png()
    fig31.generate_svg()

    print("2/7 图4-1 端到端推理流程…")
    p41 = e2e.generate_png()
    e2e.generate_svg()

    print("3/7 图5-1 系统架构…")
    p51 = sa.generate_png()
    sa.generate_svg()

    print("4/7 图5-2 问答数据流…")
    p52 = qf.generate_png()
    qf.generate_svg()

    print("5/7 图5-3 技术栈…")
    p53 = ts.generate_png()
    ts.generate_svg()

    print("6–7/7 图5-11 / 图5-12 实验结果…")
    data = ch54.load_summary()
    ch54.generate_figures(data)

    paths = [p31, p41, p51, p52, p53, ch54.FIG_FACT, ch54.FIG_COMP]
    for p in paths:
        im = Image.open(p)
        print(f"  {p.name}: {im.size[0]}×{im.size[1]}px")
    return paths


def main() -> None:
    if not THESIS.exists():
        raise FileNotFoundError(THESIS)

    shutil.copy2(THESIS, BACKUP)
    print(f"备份 -> {BACKUP.name}")

    generate_all()

    print("写入 docx（跳过图5-4～5-10 系统页面）…")
    doc = Document(str(THESIS))
    for prefix, path, width in REPLACEMENTS:
        if not path.exists():
            raise FileNotFoundError(path)
        replace_image_before_caption(doc, prefix, path, width)
    doc.save(str(THESIS))
    print(f"已保存 -> {THESIS.name}")
    print("完成。未改动：图5-4～图5-10 界面截图。")


if __name__ == "__main__":
    main()
