#!/usr/bin/env python3
"""生成图3-1「总体方法框架图」（PNG + SVG），并可选写入论文 docx。"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "docs" / "superpowers" / "architecture"
PNG_PATH = OUT_DIR / "fig-3-1-method-framework.png"
SVG_PATH = OUT_DIR / "fig-3-1-method-framework.svg"
THESIS = ROOT / "华东师范大学硕士论文.docx"

W, H = 1200, 980
SCALE = 3

COLORS = {
    "bg": (248, 250, 252),
    "card": (255, 255, 255),
    "border": (203, 213, 225),
    "title": (15, 23, 42),
    "sub": (100, 116, 139),
    "sys_fill": (37, 99, 235),
    "sys_bg": (239, 246, 255),
    "sch_fill": (124, 58, 237),
    "sch_bg": (245, 243, 255),
    "dyn_fill": (220, 38, 38),
    "dyn_bg": (254, 242, 242),
    "sta_fill": (8, 145, 178),
    "sta_bg": (236, 254, 255),
    "arrow": (100, 116, 139),
    "white": (255, 255, 255),
}


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size=size, index=1 if bold else 0)
            except OSError:
                try:
                    return ImageFont.truetype(path, size=size)
                except OSError:
                    continue
    return ImageFont.load_default()


def rr(draw: ImageDraw.ImageDraw, box, radius, fill, outline=None, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow_down(draw, x, y1, y2, scale, color):
    draw.line([(x, y1), (x, y2 - 8 * scale)], fill=color, width=max(2, 2 * scale))
    tip = [
        (x, y2),
        (x - 7 * scale, y2 - 12 * scale),
        (x + 7 * scale, y2 - 12 * scale),
    ]
    draw.polygon(tip, fill=color)


def arrow_up(draw, x, y1, y2, scale, color):
    """y1 bottom -> y2 top"""
    draw.line([(x, y1), (x, y2 + 8 * scale)], fill=color, width=max(2, 2 * scale))
    tip = [
        (x, y2),
        (x - 7 * scale, y2 + 12 * scale),
        (x + 7 * scale, y2 + 12 * scale),
    ]
    draw.polygon(tip, fill=color)


def text_center(draw, xy, text, font, fill):
    draw.text(xy, text, font=font, fill=fill, anchor="mm")


def generate_png(path: Path = PNG_PATH) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    scale = SCALE
    img = Image.new("RGB", (W * scale, H * scale), COLORS["bg"])
    draw = ImageDraw.Draw(img)

    f_title = load_font(28 * scale, bold=True)
    f_sub = load_font(15 * scale)
    f_layer = load_font(20 * scale, bold=True)
    f_body = load_font(15 * scale)
    f_small = load_font(13 * scale)
    f_tag = load_font(12 * scale)

    # outer card
    margin = 28 * scale
    rr(
        draw,
        [margin, margin, W * scale - margin, H * scale - margin],
        18 * scale,
        COLORS["card"],
        COLORS["border"],
        max(2, scale),
    )

    text_center(draw, (W * scale // 2, 70 * scale), "动态—静态知识协同问答 · 总体方法框架", f_title, COLORS["title"])
    text_center(
        draw,
        (W * scale // 2, 105 * scale),
        "自下而上：静态知识层 → 动态知识层 → 协同调度层 → 系统实现层",
        f_sub,
        COLORS["sub"],
    )

    left = 90 * scale
    right = (W - 90) * scale
    mid = W * scale // 2

    # Layer geometry (top to bottom visually: system, schedule, then static|dynamic)
    # System
    sys_top, sys_bot = 150 * scale, 280 * scale
    # Schedule
    sch_top, sch_bot = 330 * scale, 470 * scale
    # Base pair
    base_top, base_bot = 530 * scale, 740 * scale
    gap = 18 * scale
    half_w = (right - left - gap) // 2

    # --- System layer ---
    rr(draw, [left, sys_top, right, sys_bot], 14 * scale, COLORS["sys_bg"], COLORS["sys_fill"], max(2, scale))
    # accent bar
    draw.rectangle([left, sys_top, left + 10 * scale, sys_bot], fill=COLORS["sys_fill"])
    text_center(draw, (mid, sys_top + 32 * scale), "系统实现层", f_layer, COLORS["sys_fill"])
    text_center(
        draw,
        (mid, sys_top + 62 * scale),
        "端到端服务 · 提示组装 · 本地 LLM 生成 · 守卫与多模态回显 · 实验评测",
        f_body,
        COLORS["title"],
    )
    text_center(draw, (mid, sys_top + 92 * scale), "对应第五章", f_tag, COLORS["sub"])
    # chips
    chips = ["提问", "检索", "生成", "资源回显"]
    chip_y = sys_bot - 28 * scale
    chip_w = 110 * scale
    start_x = mid - (len(chips) * chip_w + (len(chips) - 1) * 12 * scale) / 2
    for i, lab in enumerate(chips):
        x0 = start_x + i * (chip_w + 12 * scale)
        rr(
            draw,
            [x0, chip_y - 16 * scale, x0 + chip_w, chip_y + 16 * scale],
            10 * scale,
            COLORS["white"],
            COLORS["sys_fill"],
            max(1, scale),
        )
        text_center(draw, (x0 + chip_w / 2, chip_y), lab, f_small, COLORS["sys_fill"])

    # --- Schedule layer ---
    rr(draw, [left, sch_top, right, sch_bot], 14 * scale, COLORS["sch_bg"], COLORS["sch_fill"], max(2, scale))
    draw.rectangle([left, sch_top, left + 10 * scale, sch_bot], fill=COLORS["sch_fill"])
    text_center(draw, (mid, sch_top + 32 * scale), "协同调度层", f_layer, COLORS["sch_fill"])
    text_center(
        draw,
        (mid, sch_top + 65 * scale),
        "三阶段分类（震前 / 震中 / 震后） · 阶段先验 + 规则阈值 · 启闭与优先级决策",
        f_body,
        COLORS["title"],
    )
    text_center(
        draw,
        (mid, sch_top + 98 * scale),
        "只决定「给什么」　｜　对应第四章 4.5—4.8 节",
        f_tag,
        COLORS["sub"],
    )

    # --- Static / Dynamic base ---
    sta_box = [left, base_top, left + half_w, base_bot]
    dyn_box = [left + half_w + gap, base_top, right, base_bot]
    rr(draw, sta_box, 14 * scale, COLORS["sta_bg"], COLORS["sta_fill"], max(2, scale))
    rr(draw, dyn_box, 14 * scale, COLORS["dyn_bg"], COLORS["dyn_fill"], max(2, scale))
    draw.rectangle([sta_box[0], sta_box[1], sta_box[0] + 10 * scale, sta_box[3]], fill=COLORS["sta_fill"])
    draw.rectangle([dyn_box[0], dyn_box[1], dyn_box[0] + 10 * scale, dyn_box[3]], fill=COLORS["dyn_fill"])

    sx = (sta_box[0] + sta_box[2]) / 2
    dx = (dyn_box[0] + dyn_box[2]) / 2
    text_center(draw, (sx, base_top + 36 * scale), "静态知识层", f_layer, COLORS["sta_fill"])
    text_center(draw, (dx, base_top + 36 * scale), "动态知识层", f_layer, COLORS["dyn_fill"])

    for cx, lines, color in (
        (
            sx,
            [
                "知识图谱（震例—区域—主题—步骤）",
                "主题级向量检索（RAG）",
                "可离线固化 · 可核对证据底座",
                "对应第四章 4.3—4.4 节",
            ],
            COLORS["title"],
        ),
        (
            dx,
            [
                "CEIC 公开速报（轻量接入）",
                "USGS FDSN 地震目录",
                "强时效 · 按需写入提示",
                "对应第四章 4.9 节",
            ],
            COLORS["title"],
        ),
    ):
        y = base_top + 80 * scale
        for line in lines:
            text_center(draw, (cx, y), line, f_body if "对应" not in line else f_tag, color if "对应" not in line else COLORS["sub"])
            y += 32 * scale

    # arrows between layers
    arrow_down(draw, mid, sys_bot + 4 * scale, sch_top - 4 * scale, scale, COLORS["arrow"])
    text_center(draw, (mid + 70 * scale, (sys_bot + sch_top) / 2), "调度结果", f_tag, COLORS["sub"])

    # from schedule down to both bases
    arrow_down(draw, sx, sch_bot + 4 * scale, base_top - 4 * scale, scale, COLORS["arrow"])
    arrow_down(draw, dx, sch_bot + 4 * scale, base_top - 4 * scale, scale, COLORS["arrow"])
    text_center(draw, (sx + 55 * scale, (sch_bot + base_top) / 2), "取证据", f_tag, COLORS["sub"])
    text_center(draw, (dx + 55 * scale, (sch_bot + base_top) / 2), "取震情", f_tag, COLORS["sub"])

    # upward evidence flow annotation on right
    rx = right + 8 * scale
    # keep inside card - put annotation at bottom
    text_center(
        draw,
        (mid, base_bot + 40 * scale),
        "统一上下文组装接口解耦：调度层决定「给什么」· 检索层负责「取什么」· 生成层负责「怎么说」",
        f_small,
        COLORS["sub"],
    )
    text_center(
        draw,
        (mid, base_bot + 68 * scale),
        "可记录 · 可消融 · 可复现",
        f_tag,
        COLORS["sub"],
    )

    # user / answer badges at very top of system area
    text_center(draw, (left + 70 * scale, sys_top - 18 * scale), "用户问句 ↓", f_tag, COLORS["sys_fill"])
    text_center(draw, (right - 70 * scale, sys_top - 18 * scale), "↓ 回答与资源", f_tag, COLORS["sys_fill"])

    img = img.resize((W, H), Image.Resampling.LANCZOS)
    img.save(path, format="PNG", optimize=True)
    print(f"Wrote {path} ({path.stat().st_size // 1024} KB)")
    return path


def generate_svg(path: Path = SVG_PATH) -> Path:
    """轻量 SVG 占位说明；论文插图以 PNG 为准。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">
  <rect width="{W}" height="{H}" fill="#f8fafc"/>
  <text x="{W//2}" y="48" text-anchor="middle" font-family="PingFang SC,Microsoft YaHei,sans-serif"
        font-size="22" font-weight="700" fill="#0f172a">总体方法框架图（见图 PNG）</text>
  <text x="{W//2}" y="80" text-anchor="middle" font-family="PingFang SC,Microsoft YaHei,sans-serif"
        font-size="14" fill="#64748b">静态知识层 / 动态知识层 / 协同调度层 / 系统实现层</text>
</svg>
""",
        encoding="utf-8",
    )
    print(f"Wrote {path}")
    return path


def insert_into_thesis(image: Path, thesis: Path = THESIS, width_in: float = 5.8) -> None:
    backup = thesis.with_suffix(thesis.suffix + ".bak-before-fig3-1")
    shutil.copy2(thesis, backup)
    doc = Document(str(thesis))
    cap_i = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith("图3-1") and "总体方法" in t:
            cap_i = i
            break
    if cap_i is None:
        raise KeyError("找不到题注：图3-1 总体方法框架图")

    target = None
    for j in range(cap_i - 1, max(-1, cap_i - 4), -1):
        p = doc.paragraphs[j]
        if not p.text.strip():
            target = p
            break
    if target is None:
        # insert empty para before caption
        cap = doc.paragraphs[cap_i]
        new_p = cap.insert_paragraph_before("")
        target = new_p

    # clear existing content/drawings
    for child in list(target._element):
        local = child.tag.split("}")[-1]
        if local in ("drawing", "pict", "r"):
            target._element.remove(child)

    run = target.add_run()
    run.add_picture(str(image), width=Inches(width_in))
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target.paragraph_format.first_line_indent = Cm(0)
    doc.save(str(thesis))
    print(f"已插入论文：{thesis.name}（备份 {backup.name}）")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-docx", action="store_true", help="只生成图片，不写入 docx")
    args = parser.parse_args()
    png = generate_png()
    generate_svg()
    if not args.no_docx:
        insert_into_thesis(png)


if __name__ == "__main__":
    main()
