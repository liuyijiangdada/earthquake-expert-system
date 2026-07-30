#!/usr/bin/env python3
"""按六章制方案重写《华东师范大学硕士论文.docx》正文。

保留封面、声明、摘要标题结构；替换摘要正文与关键词；
从「第一章」起至「致谢」前重建第1–6章、参考文献与附录。
实验数字占位符可由 --fill-eval 用 60 题归档结果替换。
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"
OUT_DEFAULT = ROOT / "华东师范大学硕士论文_6章修改版.docx"
RED = RGBColor(0xFF, 0x00, 0x00)

import sys

sys.path.insert(0, str(ROOT))
from scripts.thesis_rewrite_content_6ch import (  # noqa: E402
    ABSTRACT_CN,
    ABSTRACT_EN,
    APPENDIX_BLOCKS,
    BODY_BLOCKS,
    KEYWORDS_CN,
    KEYWORDS_EN,
    REFERENCES,
    TOC_LINES,
)


def _set_run_font(
    run,
    *,
    east_asia: str = "宋体",
    western: str = "Times New Roman",
    size_pt: float = 12,
    bold: bool = False,
    mark_red: bool = False,
) -> None:
    run.font.name = western
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if mark_red:
        run.font.color.rgb = RED
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), east_asia)


def _clear_runs(para: Paragraph) -> None:
    for run in para.runs:
        run.text = ""


def set_para_text(para: Paragraph, text: str) -> None:
    _clear_runs(para)
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def _format_para(para: Paragraph, level: str, *, mark_red: bool = True) -> None:
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if level == "h1":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(12)
        pf.space_after = Pt(12)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=16, bold=True, mark_red=mark_red)
    elif level == "h2":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=14, bold=True, mark_red=mark_red)
    elif level == "h3":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(3)
        pf.space_after = Pt(3)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=12, bold=True, mark_red=mark_red)
    elif level in ("caption", "figure_caption", "table_title"):
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
        size = 12 if level == "table_title" else 10.5
        ea = "黑体" if level == "table_title" else "宋体"
        for run in para.runs:
            _set_run_font(
                run,
                east_asia=ea,
                size_pt=size,
                bold=(level == "table_title"),
                mark_red=mark_red,
            )
    elif level == "ref":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        for run in para.runs:
            _set_run_font(run, east_asia="宋体", size_pt=10.5, bold=False, mark_red=mark_red)
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.74)
        for run in para.runs:
            _set_run_font(run, east_asia="宋体", size_pt=12, bold=False, mark_red=mark_red)


def _insert_para_after(
    anchor: Paragraph, text: str, level: str, *, mark_red: bool = True
) -> Paragraph:
    new_p = deepcopy(anchor._p)
    for child in list(new_p):
        if child.tag == qn("w:r") or child.tag.endswith("}r"):
            new_p.remove(child)
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    set_para_text(para, text)
    _format_para(para, level, mark_red=mark_red)
    return para


def _delete_paragraph(para: Paragraph) -> None:
    el = para._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def _find_para_index(doc: Document, predicate) -> int:
    for i, p in enumerate(doc.paragraphs):
        if predicate(p.text.strip()):
            return i
    raise KeyError("paragraph not found")


def _expand_body_blocks(blocks: list[tuple[str, str]]) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    for level, text in blocks:
        if level == "body" and "\n\n" in text:
            for part in text.split("\n\n"):
                part = part.strip()
                if part:
                    out.append(("body", part))
        elif level == "body" and "\n" in text and text.startswith("指标"):
            for line in text.split("\n"):
                line = line.strip()
                if line:
                    out.append(("body", line))
        else:
            out.append((level, text))
    return out


def fill_placeholders(text: str, mapping: dict[str, str]) -> str:
    def repl(m: re.Match) -> str:
        key = m.group(1)
        return mapping.get(key, m.group(0))

    return re.sub(r"\{\{([A-Z0-9_]+)\}\}", repl, text)


def build_eval_mapping(summary_path: Path) -> dict[str, str]:
    data = json.loads(summary_path.read_text(encoding="utf-8"))
    overall = data.get("overall") or data
    by_phase = data.get("by_phase") or {}
    m: dict[str, str] = {}
    for bid in ("B0", "B1", "B2", "B3"):
        s = overall.get(bid, {})
        m[f"FACT_{bid}"] = str(s.get("factual_accuracy_pct", "—"))
        m[f"COMP_{bid}"] = str(s.get("completeness_mean", "—"))
        m[f"FMT_{bid}"] = str(s.get("format_compliance_pct", "—"))
        phase_map = by_phase.get(bid, {})
        for phase, key in (
            ("震前", "PRE"),
            ("震中", "DUR"),
            ("震后", "POST"),
        ):
            ps = phase_map.get(phase, {})
            m[f"FACT_{bid}_{key}"] = str(ps.get("factual_accuracy_pct", "—"))
            m[f"COMP_{bid}_{key}"] = str(ps.get("completeness_mean", "—"))
            m[f"FMT_{bid}_{key}"] = str(ps.get("format_compliance_pct", "—"))
    # 六章制正文图号为图5-1/5-2（对应仓库 fig-6-1/6-2 资源）
    m["FIG_6_1"] = "（见图5-1）"
    m["FIG_6_2"] = "（见图5-2）"
    m["N_QUESTIONS"] = str(data.get("question_count", 60))
    return m


def apply_content(
    doc: Document,
    *,
    mapping: dict[str, str] | None = None,
    mark_red: bool = True,
) -> None:
    mapping = mapping or {}

    # --- abstract / keywords ---
    abs_i = _find_para_index(doc, lambda t: t == "摘要")
    for j in range(abs_i + 1, abs_i + 5):
        t = doc.paragraphs[j].text.strip()
        if t and not t.startswith("关键词"):
            set_para_text(doc.paragraphs[j], fill_placeholders(ABSTRACT_CN, mapping))
            _format_para(doc.paragraphs[j], "body", mark_red=mark_red)
            break
    for j in range(abs_i + 1, abs_i + 8):
        t = doc.paragraphs[j].text.strip()
        if t.startswith("关键词"):
            set_para_text(doc.paragraphs[j], f"关键词：{KEYWORDS_CN}")
            _format_para(doc.paragraphs[j], "body", mark_red=mark_red)
            break

    en_i = _find_para_index(doc, lambda t: t.upper() == "ABSTRACT")
    for j in range(en_i + 1, en_i + 5):
        t = doc.paragraphs[j].text.strip()
        if t and not t.lower().startswith("keywords"):
            set_para_text(doc.paragraphs[j], fill_placeholders(ABSTRACT_EN, mapping))
            _format_para(doc.paragraphs[j], "body", mark_red=mark_red)
            break
    for j in range(en_i + 1, en_i + 8):
        t = doc.paragraphs[j].text.strip()
        if t.lower().startswith("keywords"):
            set_para_text(doc.paragraphs[j], f"Keywords: {KEYWORDS_EN}")
            _format_para(doc.paragraphs[j], "body", mark_red=mark_red)
            break

    # --- locate body chapter start and 致谢 ---
    body_start = _find_para_index(
        doc, lambda t: t.startswith("第一章") and "绪论" in t and "\t" not in t
    )
    ack_i = _find_para_index(doc, lambda t: t == "致谢" or t.startswith("致谢"))
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() in ("致谢",) and i > body_start:
            ack_i = i

    paras = list(doc.paragraphs)
    to_delete = [paras[i] for i in range(body_start, ack_i)]
    for p in to_delete:
        _delete_paragraph(p)

    ack = None
    for p in doc.paragraphs:
        if p.text.strip() == "致谢":
            ack = p
    if ack is None:
        raise RuntimeError("致谢段落丢失")

    from docx.oxml import OxmlElement

    new_el = OxmlElement("w:p")
    ack._p.addprevious(new_el)
    cursor = Paragraph(new_el, ack._parent)
    set_para_text(cursor, "")

    blocks = _expand_body_blocks(BODY_BLOCKS)
    stream: list[tuple[str, str]] = list(blocks)
    stream.append(("h1", "参考文献"))
    for ref in REFERENCES:
        stream.append(("ref", fill_placeholders(ref, mapping)))
    stream.extend(_expand_body_blocks(APPENDIX_BLOCKS))

    for level, text in stream:
        text = fill_placeholders(text, mapping)
        cursor = _insert_para_after(cursor, text, level, mark_red=mark_red)

    for p in list(doc.paragraphs):
        if p.text.strip() != "":
            continue
        nxt = p._element.getnext()
        if nxt is None:
            continue
        try:
            np = Paragraph(nxt, p._parent)
            if (np.text or "").strip().startswith("第一章"):
                _delete_paragraph(p)
                break
        except Exception:
            continue

    # update TOC using TOC_LINES (6 chapters)
    try:
        toc_i = _find_para_index(doc, lambda t: t == "目录")
    except KeyError:
        return

    stop_titles = {"图目录", "表目录", "第一章  绪论"}
    toc_paras = []
    for i in range(toc_i + 1, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip()
        if t in stop_titles and i > toc_i + 1:
            break
        if t.startswith("第") or re.match(r"^\d+\.\d+", t) or t.startswith("参考") or t.startswith("致"):
            toc_paras.append(doc.paragraphs[i])
        elif not t:
            continue
        else:
            if toc_paras:
                break

    new_toc = list(TOC_LINES)
    for i, line in enumerate(new_toc):
        if i < len(toc_paras):
            set_para_text(toc_paras[i], line)
            # 正文致谢不改；目录「致谢」条目也不标红，避免与「致谢非红」口径冲突
            line_red = mark_red and line.strip() != "致谢"
            _format_para(toc_paras[i], "body", mark_red=line_red)
        else:
            # 底稿目录行不足时，在末条后插入，保证六章+参考文献/附录/致谢齐全
            anchor = toc_paras[-1] if toc_paras else doc.paragraphs[toc_i]
            line_red = mark_red and line.strip() != "致谢"
            new_p = _insert_para_after(anchor, line, "body", mark_red=line_red)
            toc_paras.append(new_p)
    for j in range(len(new_toc), len(toc_paras)):
        set_para_text(toc_paras[j], "")

    _refresh_figure_table_toc(doc, mark_red=mark_red)


def _refresh_figure_table_toc(doc: Document, *, mark_red: bool = True) -> None:
    """清空旧八章遗留的图/表目录条目，写入六章制图题表题清单。"""
    fig_lines = [
        "图3-1  地震应急知识图谱模式示意（正文预留）",
        "图4-1  端到端推理流程",
        "图5-1  四基线事实一致性对比（60题离线消融）",
        "图5-2  四基线要点完整性与格式合规对比（60题离线消融）",
    ]
    table_lines = [
        "表5-1  离线消融实验配置快照",
        "表5-2  四基线60题自动评分汇总",
        "表5-3  与 GraphRAG / KnowledGPT 的设计对比（文字表）",
    ]

    def _rewrite_list_section(title: str, lines: list[str]) -> None:
        try:
            start = _find_para_index(doc, lambda t: t == title)
        except KeyError:
            return
        # Collect until next major title; do not stop on blank lines mid-list
        stop_exact = {"目录", "图目录", "表目录", "摘要", "ABSTRACT"}
        entries: list = []
        end_idx = start + 1
        for i in range(start + 1, len(doc.paragraphs)):
            t = doc.paragraphs[i].text.strip()
            if t in stop_exact:
                break
            if t.startswith("第") and "章" in t[:8]:
                break
            if t.startswith(("1.", "1．")) and "研究" in t:
                break
            end_idx = i + 1
            if t:
                entries.append(doc.paragraphs[i])
        # Also collect any leftover 图/表 lines that look like old TOC numbering
        # (already included above if contiguous)

        for i, line in enumerate(lines):
            if i < len(entries):
                set_para_text(entries[i], line)
                _format_para(entries[i], "body", mark_red=mark_red)
            else:
                anchor = entries[-1] if entries else doc.paragraphs[start]
                new_p = _insert_para_after(anchor, line, "body", mark_red=mark_red)
                entries.append(new_p)
        for j in range(len(lines), len(entries)):
            set_para_text(entries[j], "")
            # strip runs so empty paras don't resurface as ghost TOC
            for r in entries[j].runs:
                r.text = ""

    _rewrite_list_section("图目录", fig_lines)
    _rewrite_list_section("表目录", table_lines)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--thesis", type=Path, default=THESIS, help="输入底稿 docx")
    ap.add_argument(
        "--output",
        type=Path,
        default=OUT_DEFAULT,
        help="输出修改版 docx（默认不覆盖底稿）",
    )
    ap.add_argument(
        "--fill-eval",
        type=Path,
        default=ROOT / "data/eval/table_6_3_summary_60.json",
        help="60题汇总 JSON（table_6_3_summary_60.json）",
    )
    ap.add_argument("--no-backup", action="store_true")
    ap.add_argument(
        "--no-red",
        action="store_true",
        help="不将新增正文标为红色",
    )
    args = ap.parse_args()

    thesis = args.thesis
    output = args.output
    if not args.no_backup:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = thesis.with_name(f"{thesis.stem}.backup.{stamp}{thesis.suffix}")
        shutil.copy2(thesis, backup)
        print(f"底稿备份: {backup}")

    # 始终从底稿复制到输出再改写，避免误改原件
    if output.resolve() == thesis.resolve():
        raise SystemExit(
            "错误：--output 不能与 --thesis 相同；请指定独立输出路径以避免覆盖底稿。"
        )
    shutil.copy2(thesis, output)
    work = output

    mapping: dict[str, str] = {}
    if args.fill_eval and args.fill_eval.exists():
        mapping = build_eval_mapping(args.fill_eval)
        print(f"已加载评测汇总: {args.fill_eval}")
        print("映射键:", ", ".join(sorted(mapping)[:12]), "...")

    doc = Document(str(work))
    apply_content(doc, mapping=mapping, mark_red=not args.no_red)
    doc.save(str(work))
    print(f"已写入: {work}")
    chars = sum(len(p.text) for p in doc.paragraphs)
    red_n = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.color and r.font.color.rgb == RED:
                red_n += 1
                break
    print(f"段落数={len(doc.paragraphs)}, 字符数≈{chars}, 含红字段落≈{red_n}")


if __name__ == "__main__":
    main()
