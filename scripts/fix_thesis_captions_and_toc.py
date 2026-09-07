#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""修复论文图题/表题编号与正文引用，并重建目录/图目录/表目录（戚媛媛格式）。"""
from __future__ import annotations

import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from scripts.refresh_thesis_toc import (  # noqa: E402
    collect_caption_targets,
    collect_heading_targets,
    rebuild_front_lists,
    update_fields_with_word,
    verify,
    _find_body_chapter1,
    fix_chapter2_title,
)
from scripts.restructure_ch345_docx import T, set_text, body_start  # noqa: E402

DOC = ROOT / "华东师范大学硕士论文.docx"
CAP_RE = re.compile(r"^(图|表)\s*(\d+)\s*[-–—]\s*(\d+)\s+(.*)$")


def chapter_at(doc: Document, bi: int, idx: int) -> int | None:
    ch = None
    for i in range(bi, idx + 1):
        t = T(doc.paragraphs[i])
        if t.startswith("第一章"):
            ch = 1
        elif t.startswith("第二章") or t == "相关技术与理论基础":
            ch = 2
        elif t.startswith("第三章") or ("多源静态知识构建" in t and len(t) < 40):
            ch = 3
        elif t.startswith("第四章"):
            ch = 4
        elif t.startswith("第五章"):
            ch = 5
        elif t.startswith("第六章"):
            ch = 6
    return ch


def is_real_caption(text: str) -> bool:
    m = CAP_RE.match(text.strip())
    if not m:
        return False
    rest = m.group(4).strip()
    if "\t" in rest:
        rest = rest.split("\t")[0].strip()
    # 叙述句不是题注
    if "。" in rest or rest.endswith("：") or rest.endswith(":"):
        return False
    if len(rest) > 55:
        return False
    if rest.count("，") >= 2:
        return False
    return bool(rest)


def main() -> None:
    bak = DOC.with_name(f"{DOC.stem}.backup_fix_toc_{datetime.now():%Y%m%d_%H%M%S}{DOC.suffix}")
    shutil.copy2(DOC, bak)
    print("备份:", bak.name)

    doc = Document(str(DOC))
    bi = body_start(doc)

    # 1) 章标题
    for p in doc.paragraphs[bi:]:
        if T(p) == "面向地震应急的多源静态知识构建":
            set_text(p, "第三章  面向地震应急的多源静态知识构建")
            print("补全第三章标题")
            break

    for p in doc.paragraphs[bi:]:
        t = T(p)
        if "静态知识层至此构建完成" in t and "下节起" in t:
            set_text(
                p,
                "静态知识层至此构建完成。下一章将在此基础上给出动态—静态协同问答方法，并完成实验评估。",
            )
            break

    # 2) 收集题注并分配新编号
    caps: list[tuple[int, str, str, str, str]] = []
    # idx, kind, old_id, title, old_full
    for i, p in enumerate(doc.paragraphs):
        if i < bi:
            continue
        t = T(p)
        if not is_real_caption(t):
            continue
        m = CAP_RE.match(t)
        assert m
        title = m.group(4).strip().split("\t")[0].strip()
        old_id = f"{m.group(1)}{m.group(2)}-{m.group(3)}"
        caps.append((i, m.group(1), old_id, title, t))

    counters: dict[tuple[str, int], int] = {}
    id_map: dict[str, str] = {}
    para_new: dict[int, str] = {}

    for idx, kind, old_id, title, old_full in caps:
        ch = chapter_at(doc, bi, idx)
        if ch is None:
            print("WARN skip", old_full)
            continue
        counters[(kind, ch)] = counters.get((kind, ch), 0) + 1
        seq = counters[(kind, ch)]
        new_id = f"{kind}{ch}-{seq}"
        new_full = f"{new_id}  {title}"
        para_new[idx] = new_full
        # 同一 old_id 多次出现时以后写为准（正文与目录残留）
        id_map[old_id] = new_id
        print(f"  {old_id:8s} → {new_id:8s}  {title}")

    # 写题注
    for idx, new_full in para_new.items():
        set_text(doc.paragraphs[idx], new_full)

    # 3) 文内引用：用临时占位避免连锁
    # 只替换明确的 图a-b / 表a-b
    tokens = {}
    for i, (oid, nid) in enumerate(id_map.items()):
        tokens[oid] = f"⟦ID{i}⟧"

    for p in doc.paragraphs:
        raw = p.text
        if not raw or not raw.strip():
            continue
        # 跳过已是最终题注的短行
        if CAP_RE.match(raw.strip()) and is_real_caption(raw.strip()):
            # 若已是 new 格式则不动
            continue
        nt = raw
        for oid, tok in sorted(tokens.items(), key=lambda x: -len(x[0])):
            # 常见前缀组合
            for pre in ("如图", "见图", "参见图", "如表", "见表", "参见表", "图", "表"):
                # 避免 图表 误伤：pre 末字须与 oid 首字一致
                if pre[-1] != oid[0]:
                    continue
                nt = nt.replace(pre + oid, pre + tok)
                nt = nt.replace(pre + " " + oid, pre + tok)
        for oid, tok in tokens.items():
            nt = nt.replace(tok, id_map[oid])
        if nt != raw:
            set_text(p, nt)

    # 框架图引用校正
    fw = None
    for p in doc.paragraphs[bi:]:
        t = T(p)
        if t.startswith("图") and "总体方法框架" in t:
            m = CAP_RE.match(t)
            if m:
                fw = f"图{m.group(2)}-{m.group(3)}"
            break
    if fw:
        for p in doc.paragraphs[bi:]:
            raw = p.text or ""
            if "自下而上分为四层" in raw and "如图" in raw:
                nt = re.sub(r"如图\s*\d+\s*[-–—]\s*\d+\s*所示", f"如{fw}所示", raw)
                if nt != raw:
                    set_text(p, nt)
                    print("框架图引用改为", fw)

    doc.save(str(DOC))
    print("题注修复完成，重建目录…")

    # 4) 重建目录
    doc = Document(str(DOC))
    fix_chapter2_title(doc)
    body_i = _find_body_chapter1(doc)
    headings = collect_heading_targets(doc, body_i)
    figs, tabs = collect_caption_targets(doc, body_i)
    print(f"headings={len(headings)} figs={len(figs)} tabs={len(tabs)}")
    for title, _ in figs:
        print("  FIG", title)
    for title, _ in tabs:
        print("  TAB", title)

    rebuild_front_lists(doc, headings, figs, tabs)
    doc.save(str(DOC))

    d2 = Document(str(DOC))
    checks = verify(d2)
    for k, v in checks.items():
        print(f"  [{'OK' if v else 'FAIL'}] {k}")

    texts = [p.text.strip() for p in d2.paragraphs]
    ti = texts.index("目录")
    bi2 = _find_body_chapter1(d2)
    zone = texts[ti:bi2]
    for key in ("3.5", "4.6.1", "4.8", "图目录", "表目录", "第三章"):
        print(f"TOC有[{key}]:", any(key in t for t in zone))

    # 正文确认 3.5
    print("正文有3.5:", any(T(p).startswith("3.5") for p in d2.paragraphs[bi2:]))

    ok = update_fields_with_word(DOC)
    if not ok:
        print("提示：请用 WPS 打开 → 引用/目录处右键「更新整个目录」，或 Ctrl+A 后更新域。")
    print("完成。备份:", bak.name)


if __name__ == "__main__":
    main()
