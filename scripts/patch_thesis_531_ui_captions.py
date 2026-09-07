#!/usr/bin/env python3
"""按设计规格补全 5.3.1 各图操作说明并调整插图尺寸。"""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"

EMU_PER_CM = 360000
WEB_W = int(12.5 * EMU_PER_CM)
MOBILE_W = int(6.5 * EMU_PER_CM)

INTRO = (
    "本系统前端包含 Web 端与 Android 移动端两部分，二者复用同一套后端问答与图谱接口。"
    "Web 端面向桌面浏览器场景，采用“对话区 + 震情洞察区”的双栏布局，贯通“提问—检索—生成—资源回显”的交互闭环；"
    "Android 端面向移动应急场景，底部导航栏提供问答、指引、更多三个入口，其中求助、家人、上报与志愿等辅助功能收纳于“更多”。"
    "下面结合关键页面与用户操作路径分别说明。"
)

CAPTIONS = {
    "5-4": "图5-4  Web端登录界面",
    "5-5": "图5-5  Web端问答主界面",
    "5-6": "图5-6  Web端刷新数据",
    "5-7": "图5-7  移动端登录界面",
    "5-8": "图5-8  移动端问答界面",
    "5-9": "图5-9  移动端应急安全指引界面",
    "5-10": "图5-10  移动端辅助页面",
}

DESCS = {
    "5-4": (
        "Web 端登录界面如图5-4所示。用户在浏览器打开系统地址后，于表单中输入用户名与密码，点击“登录”进入主界面；"
        "校验失败时页面提示错误信息，不进入问答能力。该界面作为统一身份入口，保护问答、震情刷新与图谱相关能力，"
        "避免未授权访问演示环境中的业务接口。"
    ),
    "5-5": (
        "Web 端问答主界面如图5-5所示，由顶部导航栏、场景快捷提问区、对话交互区、输入控制条与右侧震情洞察区组成。"
        "页面加载后自动推送欢迎语；用户可单击阶段快捷按钮将预设问题填入输入框，或直接输入问句后点击“发送”（Enter 提交，Shift+Enter 换行）。"
        "回答以气泡返回并标注阶段标签；右侧震情洞察区以统计卡片与图表呈现知识库震情摘要，页面加载时自动获取并定时刷新。"
        "该布局将提问、作答与态势浏览集于一屏，降低检索与阅读门槛。"
    ),
    "5-6": (
        "Web 端动态数据刷新如图5-6所示。用户在问答主界面点击“刷新地震数据”后，系统向后端请求同步公开震情源，"
        "并在对话区返回本次更新条数或失败提示。该操作使动态源可按需手动同步，保证演示与联调场景下的震情时效，"
        "而不必依赖仅定时拉取。"
    ),
    "5-7": (
        "移动端登录界面如图5-7所示。用户启动应用后填写服务器地址、用户名与密码，点击“登录”完成鉴权并进入主界面；"
        "服务器地址支持切换不同部署环境，便于联调与现场演示。该设计在窄屏上保留必要配置项，使同一客户端可对接开发、测试或演示后端。"
    ),
    "5-8": (
        "移动端问答界面如图5-8所示。用户可通过底部快捷按钮一键填入典型问题，或在输入框输入文字并发送；支持上传图片作为问句附件。"
        "回答区展示阶段标签与要点化结果，便于在窄屏下快速阅读避险与处置信息。该界面将移动场景下的高频提问路径缩短为“点选或输入—发送—阅读”。"
    ),
    "5-9": (
        "移动端应急安全指引界面如图5-9所示。用户在底栏进入“指引”后，可按震前、震中、震后浏览要点卡片，"
        "并点选安全、注意或危险等状态以标记个人处境。该页面将静态安全指引与轻量状态上报结合，补充问答之外的结构化自助查阅。"
    ),
    "5-10": (
        "移动端辅助功能由底栏“更多”进入，如图5-10所示。用户可在此访问求助、家人、上报与志愿等入口："
        "求助页发起救援/医疗/物资/避难等请求；家人页维护成员与安全状态；上报页提交灾情类型、位置与描述；志愿页完成技能与地区报名。"
        "上述能力与问答互补，覆盖行动类应急需求，而不占用主问答路径的屏幕空间。"
    ),
}


def set_run_font(run, size_pt: float = 12) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "宋体")
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")


def set_para_text(p, text: str, size_pt: float = 12) -> None:
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size_pt)
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run(text)
        set_run_font(run, size_pt)


def insert_paragraph_after(para, text: str):
    new_p = OxmlElement("w:p")
    para._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    new_para = Paragraph(new_p, para._parent)
    run = new_para.add_run(text)
    set_run_font(run, 12)
    # first line indent like body
    pPr = new_para._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLine"), "480")  # ~0.85cm
    return new_para


def has_image(p) -> bool:
    return bool(
        p._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    )


def resize_images_in_para(p, target_cx: int) -> None:
    for ext in p._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
        cx = ext.get("cx")
        cy = ext.get("cy")
        if not cx or not cy:
            continue
        cx_i, cy_i = int(cx), int(cy)
        if cx_i <= 0:
            continue
        new_cy = int(cy_i * (target_cx / cx_i))
        ext.set("cx", str(target_cx))
        ext.set("cy", str(new_cy))
    # also wp:extent if present
    for ext in p._p.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
    ):
        cx = ext.get("cx")
        cy = ext.get("cy")
        if not cx or not cy:
            continue
        cx_i, cy_i = int(cx), int(cy)
        if cx_i <= 0:
            continue
        new_cy = int(cy_i * (target_cx / cx_i))
        ext.set("cx", str(target_cx))
        ext.set("cy", str(new_cy))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def fig_key(text: str) -> str | None:
    m = re.match(r"^图\s*5-(\d+)", text.strip())
    if not m:
        return None
    return f"5-{int(m.group(1))}"


def main() -> int:
    if not THESIS.exists():
        print("missing", THESIS)
        return 1
    bak = THESIS.with_suffix(THESIS.suffix + f".bak-531-{datetime.now():%Y%m%d-%H%M%S}")
    shutil.copy2(THESIS, bak)
    print("backup", bak.name)

    doc = Document(str(THESIS))
    start = end = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if i > 200 and t.startswith("5.3.1") and "前端" in t:
            start = i
        if start is not None and i > start and t.startswith("5.3.2"):
            end = i
            break
    if start is None or end is None:
        print("cannot locate 5.3.1 body range")
        return 1
    print(f"range {start}-{end}")

    # 1) intro
    set_para_text(doc.paragraphs[start + 1], INTRO)

    # 2) walk and collect caption indices (refresh after inserts)
    # First pass: fix captions, resize images, delete old descriptive paras that are not captions
    # Strategy: remove all non-empty non-caption non-heading paras between start+2 and end,
    # then insert fresh desc after each caption.

    # Identify caption paras and image paras
    caption_idxs = {}
    image_before_caption = {}
    for i in range(start + 1, end):
        t = doc.paragraphs[i].text.strip()
        k = fig_key(t)
        if k and k in CAPTIONS:
            caption_idxs[k] = i
            if i > 0 and has_image(doc.paragraphs[i - 1]):
                image_before_caption[k] = i - 1

    print("captions", caption_idxs)

    # Delete explanatory paragraphs in range (keep title, intro, images, captions)
    # Collect elements to remove from high index to low
    keep_idxs = {start, start + 1}
    keep_idxs.update(caption_idxs.values())
    keep_idxs.update(image_before_caption.values())
    # also keep empty spacers? remove non-caption text paras
    to_delete = []
    for i in range(start + 2, end):
        if i in keep_idxs:
            continue
        p = doc.paragraphs[i]
        t = p.text.strip()
        if has_image(p):
            # orphan image without caption mapping - keep if any
            continue
        if t:
            to_delete.append(p._p)
    for elem in to_delete:
        parent = elem.getparent()
        if parent is not None:
            parent.remove(elem)
    print(f"removed {len(to_delete)} old description paras")

    # Refresh paragraph list after deletion
    doc = Document(str(THESIS)) if False else doc  # same doc object, paragraphs property refreshes
    # Re-find range
    start = end = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if i > 200 and t.startswith("5.3.1") and "前端" in t:
            start = i
        if start is not None and i > start and t.startswith("5.3.2"):
            end = i
            break

    # Fix captions + resize + insert descriptions (from last caption to first to keep indices stable... 
    # inserting after caption shifts later indices; process from bottom to top)
    items = []
    for i, p in enumerate(doc.paragraphs):
        if i <= start or i >= end:
            continue
        t = p.text.strip()
        k = fig_key(t)
        if k and k in CAPTIONS:
            items.append((i, k))
    items.sort(reverse=True)

    for i, k in items:
        p = doc.paragraphs[i]
        set_para_text(p, CAPTIONS[k])
        # resize preceding image
        if i > 0 and has_image(doc.paragraphs[i - 1]):
            w = WEB_W if k in ("5-4", "5-5", "5-6") else MOBILE_W
            resize_images_in_para(doc.paragraphs[i - 1], w)
        # insert description after caption if next is not already our desc
        nxt = doc.paragraphs[i + 1] if i + 1 < len(doc.paragraphs) else None
        desc = DESCS[k]
        if nxt is not None and nxt.text.strip() == desc:
            continue
        # if next is another caption/image/section, insert; if next is old text, replace
        if nxt is not None:
            nt = nxt.text.strip()
            if nt and not fig_key(nt) and not nt.startswith("5.3.2") and not has_image(nxt):
                set_para_text(nxt, desc)
                continue
        insert_paragraph_after(p, desc)
        print(f"  inserted/updated desc for 图{k}")

    # Re-set intro in case
    for i, p in enumerate(doc.paragraphs):
        if i > 200 and p.text.strip().startswith("5.3.1") and "前端" in p.text:
            set_para_text(doc.paragraphs[i + 1], INTRO)
            start = i
            break

    doc.save(str(THESIS))
    print("saved", THESIS)

    # Verify
    doc = Document(str(THESIS))
    start = end = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if i > 200 and t.startswith("5.3.1") and "前端" in t:
            start = i
        if start is not None and i > start and t.startswith("5.3.2"):
            end = i
            break
    print("\n=== 5.3.1 after ===")
    for i in range(start, end):
        p = doc.paragraphs[i]
        t = p.text.strip()
        img = has_image(p)
        if img:
            for ext in p._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
                cx = ext.get("cx")
                if cx:
                    print(f"{i}: IMG width_cm={int(cx)/EMU_PER_CM:.2f}")
                    break
        elif t:
            print(f"{i}: {t[:80]}")
    # checks
    blob = "\n".join(p.text for p in doc.paragraphs[start:end])
    assert "登陆" not in blob
    assert "三个入口" in blob or "问答、指引、更多" in blob
    for k, cap in CAPTIONS.items():
        assert cap in blob or cap.replace("  ", " ") in blob.replace("  ", " "), cap
        assert DESCS[k][:20] in blob, k
    assert doc.paragraphs[end].text.strip().startswith("5.3.2")
    print("\nOK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
