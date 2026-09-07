#!/usr/bin/env python3
"""Extract and review thesis docx for honesty/consistency issues."""
import re
import json
from collections import defaultdict
from docx import Document
from docx.oxml.ns import qn

DOCX_PATH = "/Users/xiaoxiaoqingnian/Desktop/biyelunwen/华东师范大学硕士论文.docx"

# Heading style name patterns
HEADING_STYLES = {"Heading 1", "Heading 2", "Heading 3", "Heading 4",
                  "标题 1", "标题 2", "标题 3", "标题 4",
                  "1", "2", "3", "4", "5", "6", "7", "8", "9"}

SEARCH_TERMS = [
    "显著", "验证了有效", "不确定性感知", "隔离静态与动态", "开启动态提升",
    "4B", "1.8B", "7B", "LangChain", "并行探测", "16题", "91.2%",
    "人工评测已完成", "Kappa", "P50", "方法研究",
    "快准全", "事实可核验", "全量", "协同调度", "动态—静态", "动态-静态",
    "B0", "B1", "B2", "B3", "B4", "B5", "82.5", "52.4", "66.8", "41.2",
    "Qwen", "LoRA", "LightRAG", "GB/T 7714", "显著性",
    "calibrat", "校准", "weak-base", "弱基线",
    "不纳入结论", "placeholder", "占位",
    "graph+RAG", "DYNAMIC", "use_kg", "use_rag",
    "78", "80", "0.64", "人工评测", "人类评估", "人评",
]

def is_heading(para):
    style = para.style.name if para.style else ""
    if any(h in style for h in ["Heading", "标题", "heading"]):
        return True
    if style and re.match(r"^[1-9]$", style.strip()):
        return True
    text = para.text.strip()
    # Chinese chapter headings
    if re.match(r"^第[一二三四五六七八九十\d]+章", text):
        return True
    if re.match(r"^\d+(\.\d+)*\s+[\u4e00-\u9fff]", text) and len(text) < 80:
        return True
    if re.match(r"^[一二三四五六七八九十]+、", text) and len(text) < 60:
        return True
    return False

def heading_level(para):
    style = para.style.name if para.style else ""
    m = re.search(r"(\d+)", style)
    if m:
        return int(m.group(1))
    text = para.text.strip()
    if re.match(r"^第[一二三四五六七八九十\d]+章", text):
        return 1
    m = re.match(r"^(\d+(?:\.\d+)*)", text)
    if m:
        return m.group(1).count(".") + 1
    return 2

def get_para_index(doc):
    """Return list of (idx, para, text)."""
    return [(i, p, p.text) for i, p in enumerate(doc.paragraphs)]

def extract_tables(doc):
    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            rows.append(cells)
        tables.append({"index": ti, "rows": rows})
    return tables

def main():
    doc = Document(DOCX_PATH)
    paras = get_para_index(doc)
    tables = extract_tables(doc)

    # 1. Outline
    outline = []
    for idx, para, text in paras:
        t = text.strip()
        if not t:
            continue
        if is_heading(para):
            outline.append({
                "idx": idx,
                "level": heading_level(para),
                "style": para.style.name if para.style else "",
                "text": t[:120],
            })

    # 2. Character count - exclude refs section
    ref_start = None
    for idx, para, text in paras:
        t = text.strip()
        if re.match(r"^(参考文献|References|REFERENCES)$", t):
            ref_start = idx
            break
    body_chars = 0
    cn_chars = 0
    for idx, para, text in paras:
        if ref_start is not None and idx >= ref_start:
            break
        for ch in text:
            body_chars += 1
            if "\u4e00" <= ch <= "\u9fff":
                cn_chars += 1

    # 3. Search terms with context
    hits = defaultdict(list)
    for idx, para, text in paras:
        if ref_start is not None and idx >= ref_start:
            # still search refs for some terms but mark them
            in_refs = True
        else:
            in_refs = False
        for term in SEARCH_TERMS:
            if term.lower() in text.lower() or term in text:
                snippet = text.strip()[:200]
                hits[term].append({
                    "idx": idx,
                    "in_refs": in_refs,
                    "snippet": snippet,
                    "style": para.style.name if para.style else "",
                })

    # Title / abstract extraction
    title_cn = ""
    title_en = ""
    abstract_cn = []
    abstract_en = []
    keywords = []
    innovation = []
    toc_entries = []

    in_abstract_cn = False
    in_abstract_en = False
    for idx, para, text in paras[:200]:
        t = text.strip()
        if not title_cn and len(t) > 10 and "地震" in t and idx < 30:
            if "基于" in t or "设计" in t:
                title_cn = t
        if "Abstract" in t or "ABSTRACT" in t:
            in_abstract_en = True
            in_abstract_cn = False
        if t == "摘要" or t.startswith("摘要"):
            in_abstract_cn = True
            in_abstract_en = False
        if in_abstract_cn and t and t != "摘要":
            abstract_cn.append({"idx": idx, "text": t})
        if in_abstract_en and t and t not in ("Abstract", "ABSTRACT"):
            abstract_en.append({"idx": idx, "text": t})
        if "创新" in t and len(t) < 50:
            innovation.append({"idx": idx, "text": t})
        if "目录" in t or para.style.name.startswith("toc") if para.style else False:
            toc_entries.append({"idx": idx, "text": t})

    # Chapter count
    chapters = [o for o in outline if re.match(r"^第[一二三四五六七八九十\d]+章", o["text"])]

    # Table 5-1 search
    table_51 = None
    for ti, table in enumerate(tables):
        flat = " ".join(" ".join(r) for r in table["rows"])
        if "Kappa" in flat or "kappa" in flat or ("78" in flat and "80" in flat):
            table_51 = {"table_index": ti, "rows": table["rows"], "flat": flat[:500]}
        if "表5-1" in flat or "表 5-1" in flat:
            table_51 = {"table_index": ti, "rows": table["rows"], "flat": flat[:500]}

    # Env table search
    env_tables = []
    for ti, table in enumerate(tables):
        flat = " ".join(" ".join(r) for r in table["rows"])
        if any(k in flat for k in ["Qwen", "LangChain", "Python", "GPU", "实验环境", "硬件", "软件"]):
            env_tables.append({"table_index": ti, "rows": table["rows"]})

    # Page count approximation (word count / 800 chars per page for Chinese thesis)
    approx_pages = cn_chars / 800

    result = {
        "file": DOCX_PATH,
        "title_cn_candidates": title_cn,
        "chapter_count": len(chapters),
        "chapters": chapters,
        "outline_count": len(outline),
        "outline_sample": outline[:80],
        "outline_full": outline,
        "cn_char_count_body": cn_chars,
        "total_char_count_body": body_chars,
        "approx_pages": round(approx_pages, 1),
        "ref_start_idx": ref_start,
        "search_hits": {k: v[:15] for k, v in hits.items() if v},
        "abstract_cn": abstract_cn[:20],
        "abstract_en": abstract_en[:20],
        "innovation_headings": innovation,
        "table_51": table_51,
        "env_tables": env_tables,
        "total_tables": len(tables),
        "total_paragraphs": len(paras),
    }

    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
