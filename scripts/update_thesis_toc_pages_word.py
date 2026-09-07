#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""用 Microsoft Word 更新论文目录 PAGEREF 页码域。"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"


def main() -> None:
    doc = THESIS.resolve()
    if not doc.exists():
        raise SystemExit(f"找不到 {doc}")

    script = f'''
set docFile to POSIX file "{doc}"
tell application "Microsoft Word"
    activate
    open docFile
    delay 4
    set theDoc to active document
    try
        set n to count of fields of theDoc
        repeat with i from 1 to n
            try
                update field (field i of theDoc)
            end try
        end repeat
    on error
        try
            update fields theDoc
        end try
    end try
    save theDoc
    delay 1
    close theDoc saving yes
end tell
return "OK"
'''
    r = subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=240)
    print("stdout:", (r.stdout or "").strip())
    print("stderr:", (r.stderr or "").strip()[:800])
    if r.returncode != 0:
        raise SystemExit(r.returncode)

    from docx import Document
    import re

    d = Document(str(doc))
    print("--- TOC sample ---")
    for i, p in enumerate(d.paragraphs):
        if p.text.strip() == "目录":
            for j in range(1, 15):
                print(repr(d.paragraphs[i + j].text[:90]))
            break
    nums = dashes = 0
    in_toc = False
    for p in d.paragraphs:
        t = p.text.strip()
        if t == "目录":
            in_toc = True
            continue
        if not in_toc:
            continue
        if t.startswith("第一章") and "绪论" in t and "\t" not in p.text:
            # reached body if no tab - but body also 第一章; TOC entries have tab/pageref
            # stop when we hit body: no hyperlink
            if "w:hyperlink" not in p._p.xml and "PAGEREF" not in p._p.xml:
                break
        if not t or t in ("图目录", "表目录"):
            continue
        if t.endswith("—"):
            dashes += 1
        elif re.search(r"\d+\s*$", t):
            nums += 1
    print(f"toc numbered≈{nums}, still dash≈{dashes}")


if __name__ == "__main__":
    main()
