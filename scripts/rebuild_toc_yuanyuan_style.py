#!/usr/bin/env python3
"""按媛媛论文样式重建主目录、图目录、表目录。"""
from __future__ import annotations

import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DOCX = ROOT / "华东师范大学硕士论文.docx"

CAPTION_VERBS = re.compile(r"给出|展示|说明|如图|如表|所示|描述|为图|为表|是图|是表")

FRONT_BOOKMARKS = (
    ("摘要", "_TocFrontAbstract", 8001),
    ("ABSTRACT", "_TocFrontABSTRACT", 8002),
    ("Abstract", "_TocFrontABSTRACT", 8002),
    ("图目录", "_TocFrontFigList", 8003),
    ("表目录", "_TocFrontTblList", 8004),
)

# 主目录静态入口（摘要/ABSTRACT/图目录/表目录），插在 TOC 域前，不依赖更新域即可看见
FRONT_TOC_ENTRIES = (
    ("摘要", "_TocFrontAbstract"),
    ("ABSTRACT", "_TocFrontABSTRACT"),
    ("图目录", "_TocFrontFigList"),
    ("表目录", "_TocFrontTblList"),
)

def is_caption(text: str, kind: str) -> bool:
    t = text.strip()
    if not re.match(rf"^{kind}\s*\d+-\d+\s+\S", t):
        return False
    if kind == "图" and len(t) > 60:
        return False
    if kind == "表" and len(t) > 80:
        return False
    if CAPTION_VERBS.search(t[:20]):
        return False
    return True


def parse_label(text: str, kind: str):
    m = re.match(rf"^({kind}\s*\d+-\d+)\s+(.+)$", text.strip())
    if not m:
        return None
    label = re.sub(r"\s+", "", m.group(1))
    nums = re.search(r"(\d+)-(\d+)", label)
    if not nums:
        return None
    return int(nums.group(1)), int(nums.group(2)), label, m.group(2).strip()


def backup_docx(path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    bak = path.with_suffix(path.suffix + f".bak-toc-{stamp}")
    shutil.copy2(path, bak)
    return bak


def set_outline_levels(doc: Document) -> int:
    """为摘要/ABSTRACT/图目录/表目录/章/节/小节/参考文献/致谢/附录设大纲级别。"""
    count = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        level = None
        if t in ("摘要", "ABSTRACT", "Abstract", "图目录", "表目录", "参考文献", "致谢"):
            level = 0
        elif re.match(r"^附录\s*[A-Z]", t) or re.match(r"^附录[A-Z]", t):
            level = 0
        elif re.match(r"^第[一二三四五六七八九十]+章", t) and len(t) < 50:
            level = 0
        elif re.match(r"^[A-Z]\.\d+\s", t) and len(t) < 40:  # A.1
            level = 1
        elif re.match(r"^\d+\.\d+\.\d+\s", t) and len(t) < 60:
            level = 2
        elif re.match(r"^\d+\.\d+\s", t) and len(t) < 50:
            level = 1
        if level is None:
            continue
        pPr = p._p.find(qn("w:pPr"))
        if pPr is None:
            pPr = parse_xml(f"<w:pPr {nsdecls('w')}/>")
            p._p.insert(0, pPr)
        existing = pPr.find(qn("w:outlineLvl"))
        if existing is not None:
            pPr.remove(existing)
        pPr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{level}"/>'))
        count += 1
    return count


def clear_between(start_elem, stop_predicate) -> int:
    removed = 0
    cur = start_elem.getnext()
    while cur is not None:
        if cur.tag != qn("w:p"):
            cur = cur.getnext()
            continue
        texts = cur.findall(f".//{qn('w:t')}")
        full = "".join(t.text or "" for t in texts).strip()
        if stop_predicate(full):
            break
        nxt = cur.getnext()
        cur.getparent().remove(cur)
        removed += 1
        cur = nxt
    return removed


def insert_toc_field(after_elem) -> None:
    toc_xml = (
        f'<w:p {nsdecls("w")}>'
        f"<w:pPr/>"
        f'<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
        f'<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r><w:t>（请在 WPS/Word 中右键此处→更新域，生成目录）</w:t></w:r>'
        f'<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        f"</w:p>"
    )
    after_elem.addnext(parse_xml(toc_xml))


def ensure_bookmark(para, name: str, bm_id: int) -> str:
    """若段落尚无指定书签则插入，返回书签名。"""
    for child in para._p:
        if child.tag == qn("w:bookmarkStart") and child.get(qn("w:name")) == name:
            return name
    para._p.insert(
        0,
        parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bm_id}" w:name="{name}"/>'),
    )
    para._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bm_id}"/>'))
    return name


def get_bookmark(para) -> str | None:
    m = re.search(r'w:bookmarkStart[^>]*w:name="([^"]+)"', para._p.xml)
    return m.group(1) if m else None


def add_bookmark(para, name: str, bm_id: int) -> str:
    """兼容旧调用；幂等插入，委托 ensure_bookmark。"""
    return ensure_bookmark(para, name, bm_id)


def make_toc_entry(label: str, title: str, bookmark: str, tab_pos: str = "8295"):
    safe = (
        title.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )
    display = f"{label}  {safe}".strip() if safe else label
    ppr = (
        "<w:pPr>"
        f'<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{tab_pos}"/></w:tabs>'
        '<w:jc w:val="left"/>'
        "</w:pPr>"
    )
    rpr = (
        '<w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体" '
        'w:hAnsi="Times New Roman"/><w:sz w:val="21"/><w:szCs w:val="21"/></w:rPr>'
    )
    return parse_xml(
        f'<w:p {nsdecls("w")}>{ppr}'
        f'<w:hyperlink w:anchor="{bookmark}" w:history="1">'
        f'<w:r>{rpr}<w:t xml:space="preserve">{display}</w:t></w:r>'
        f"<w:r><w:tab/></w:r></w:hyperlink>"
        f'<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> PAGEREF {bookmark} \\h </w:instrText></w:r>'
        f'<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f"<w:r>{rpr}<w:t>1</w:t></w:r>"
        f'<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
    )


def make_front_toc_entry(title: str, bookmark: str, tab_pos: str = "8295"):
    """主目录前端静态入口：标题 + 点引导线 + PAGEREF。"""
    return make_toc_entry(title, "", bookmark, tab_pos)


def insert_front_toc_entries(after_elem) -> object:
    """在 after_elem 后插入 4 条前端入口，返回最后插入的元素。"""
    prev = after_elem
    for title, bookmark in FRONT_TOC_ENTRIES:
        elem = make_front_toc_entry(title, bookmark)
        prev.addnext(elem)
        prev = elem
    return prev


def collect_captions(paras, kind: str, min_idx: int):
    items = []
    for i, p in enumerate(paras):
        if i < min_idx:
            continue
        if not is_caption(p.text, kind):
            continue
        parsed = parse_label(p.text, kind)
        if not parsed:
            continue
        chapter, seq, label, title = parsed
        items.append((chapter, seq, i, p, label, title, get_bookmark(p)))
    items.sort(key=lambda x: (x[0], x[1], x[2]))
    return [(p, label, title, bm) for _, _, _, p, label, title, bm in items]


def center_heading(para) -> None:
    pPr = para._p.find(qn("w:pPr"))
    if pPr is None:
        pPr = parse_xml(f"<w:pPr {nsdecls('w')}/>")
        para._p.insert(0, pPr)
    existing = pPr.find(qn("w:jc"))
    if existing is not None:
        pPr.remove(existing)
    pPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))


def _first_para(doc: Document, text: str):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    return None


def _ensure_front_bookmarks(doc: Document) -> list[str]:
    ensured: list[str] = []
    seen_names: set[str] = set()
    for heading, bm_name, bm_id in FRONT_BOOKMARKS:
        if bm_name in seen_names:
            continue
        para = _first_para(doc, heading)
        if para is None and heading == "ABSTRACT":
            para = _first_para(doc, "Abstract")
        if para is None:
            continue
        ensure_bookmark(para, bm_name, bm_id)
        ensured.append(bm_name)
        seen_names.add(bm_name)
    return ensured


def _body_start_idx(doc: Document, tbl_heading) -> int | None:
    after_tbl = False
    for i, p in enumerate(doc.paragraphs):
        if tbl_heading is not None and p._p is tbl_heading._p:
            after_tbl = True
            continue
        if after_tbl and re.match(r"^第[一二三四五六七八九十]+章", p.text.strip()) and len(p.text.strip()) < 50:
            return i
    return None


def _ensure_caption_bookmarks(items: list, bm_id_start: int = 9100) -> tuple[list, int]:
    bm_id = bm_id_start
    out = []
    for p, label, title, bm in items:
        if bm is None:
            bm = add_bookmark(p, f"_TocAuto{bm_id}", bm_id)
            bm_id += 1
        out.append((p, label, title, bm))
    return out, bm_id


def rebuild_fig_tbl_tocs(doc: Document) -> tuple[int, int]:
    """清空并重建图目录、表目录，返回 (图条数, 表条数)。"""
    fig_heading = _first_para(doc, "图目录")
    tbl_heading = _first_para(doc, "表目录")
    if fig_heading is None or tbl_heading is None:
        raise RuntimeError("未找到「图目录」或「表目录」标题")

    body_idx = _body_start_idx(doc, tbl_heading)
    if body_idx is None:
        raise RuntimeError("未找到正文第一章（表目录之后）")

    paras = doc.paragraphs
    figs, bm_id = _ensure_caption_bookmarks(collect_captions(paras, "图", body_idx))
    tbls, _ = _ensure_caption_bookmarks(collect_captions(paras, "表", body_idx), bm_id)
    print(f"图标题: {len(figs)} 个")
    print(f"表标题: {len(tbls)} 个")

    clear_between(fig_heading._p, lambda t: t == "表目录")
    prev = fig_heading._p
    for _, label, title, bm in figs:
        elem = make_toc_entry(label, title, bm)
        prev.addnext(elem)
        prev = elem

    def tbl_stop(t: str) -> bool:
        return bool(re.match(r"^第[一二三四五六七八九十]+章", t)) or t in ("参考文献", "致谢")

    clear_between(tbl_heading._p, tbl_stop)
    prev = tbl_heading._p
    for _, label, title, bm in tbls:
        elem = make_toc_entry(label, title, bm)
        prev.addnext(elem)
        prev = elem

    return len(figs), len(tbls)


def main(argv: list[str] | None = None) -> int:
    args = argv if argv is not None else sys.argv[1:]
    dry_run = "--dry-run" in args
    args = [a for a in args if a != "--dry-run"]
    docx_path = Path(args[0]) if args else DEFAULT_DOCX
    if not docx_path.exists():
        print(f"文件不存在: {docx_path}")
        return 1

    if dry_run:
        doc = Document(str(docx_path))
        toc = _first_para(doc, "目录")
        fig = _first_para(doc, "图目录")
        tbl = _first_para(doc, "表目录")
        body_idx = _body_start_idx(doc, tbl)
        print(
            f"dry-run: 目录={toc is not None} 图目录={fig is not None} "
            f"表目录={tbl is not None} 正文第一章={body_idx is not None}"
        )
        if toc is None or fig is None or tbl is None or body_idx is None:
            print("定位失败：缺少「目录」「图目录」「表目录」或正文第一章")
            return 1
        n = set_outline_levels(doc)
        print(f"大纲级别(未保存): {n}")
        bookmarks = _ensure_front_bookmarks(doc)
        print(f"前端书签(未保存): {bookmarks}")
        figs = collect_captions(doc.paragraphs, "图", body_idx)
        tbls = collect_captions(doc.paragraphs, "表", body_idx)
        print(f"图标题: {len(figs)} 个")
        print(f"表标题: {len(tbls)} 个")
        return 0

    bak = backup_docx(docx_path)
    print(f"备份: {bak}")
    doc = Document(str(docx_path))

    # 全部改动在内存完成后再保存，避免半清空状态落盘
    n = set_outline_levels(doc)
    print(f"大纲级别: {n}")

    toc_heading = _first_para(doc, "目录")
    fig_heading = _first_para(doc, "图目录")
    tbl_heading = _first_para(doc, "表目录")
    if toc_heading is None or fig_heading is None or tbl_heading is None:
        print("未找到「目录」「图目录」或「表目录」标题")
        return 1

    bookmarks = _ensure_front_bookmarks(doc)
    print(f"前端书签: {bookmarks}")

    removed = clear_between(toc_heading._p, lambda t: t == "图目录")
    print(f"清空主目录区: {removed}")
    # 静态 PAGEREF 入口（不依赖更新域即可看见）+ TOC 域（更新后填充章节）
    last_front = insert_front_toc_entries(toc_heading._p)
    print(f"主目录前端入口: {[t for t, _ in FRONT_TOC_ENTRIES]}")
    insert_toc_field(last_front)

    n_fig, n_tbl = rebuild_fig_tbl_tocs(doc)
    print(f"图目录 {n_fig} 条，表目录 {n_tbl} 条。")

    # 重新取标题段（图/表目录重建后对象仍有效）
    for title in ("目录", "图目录", "表目录"):
        para = _first_para(doc, title)
        if para is not None:
            center_heading(para)

    doc.save(str(docx_path))
    print(f"已保存：{docx_path}")
    print()
    print("=== 请在 WPS/Word 中完成以下操作 ===")
    print(f"1. 打开：{docx_path.name}")
    print("2. 在主目录灰色域上右键 → 更新域 → 更新整个目录")
    print("3. 全选（Ctrl/Cmd+A）→ 按 F9 或选「更新域」，刷新图/表目录页码")
    print("4. 目视对照媛媛论文：主目录应含摘要、ABSTRACT、图目录、表目录入口；")
    print("   其后为独立的图目录页与表目录页（主目录不含图/表题注条目）。")
    print("   （摘要等入口为静态 PAGEREF；章节列表由 TOC 域更新后生成。）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
