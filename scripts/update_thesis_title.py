#!/usr/bin/env python3
"""统一更新论文题名（中英文封面、声明、正文引用）。"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent

CN_TITLE = "基于动态-静态知识协同的地震灾害问答系统"
EN_L1 = "An Earthquake Disaster Question Answering System"
EN_L2 = "Based on Dynamic-Static Knowledge Collaboration"

OLD_CN_PATTERNS = [
    "基于知识图谱与向量检索协同的地震应急问答方法研究",
    "基于动态—静态知识协同的地震应急问答方法研究",
    "基于动态-静态知识协同的地震应急问答方法研究",
    "基于动态—静态知识协同的地震灾害问答系统",
]

OLD_EN_PATTERNS = [
    ("Research on Earthquake Emergency Question Answering", EN_L1),
    ("via Knowledge Graph and Vector Retrieval Collaboration", EN_L2),
    ("via Dynamic–Static Knowledge Collaboration", EN_L2),
    ("via Dynamic-Static Knowledge Collaboration", EN_L2),
]


def replace_in_text(text: str) -> str:
    out = text
    for old in OLD_CN_PATTERNS:
        out = out.replace(old, CN_TITLE)
        out = out.replace(f"《{old}》", f"《{CN_TITLE}》")
    for old, new in OLD_EN_PATTERNS:
        if old != new:
            out = out.replace(old, new)
    return out


def set_para_text(para, text: str) -> None:
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def apply(doc_path: Path) -> None:
    doc = Document(str(doc_path))
    set_para_text(doc.paragraphs[8], CN_TITLE)
    set_para_text(doc.paragraphs[23], EN_L1)
    set_para_text(doc.paragraphs[24], EN_L2)

    for para in doc.paragraphs:
        t = replace_in_text(para.text)
        if t != para.text:
            set_para_text(para, t)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    t = replace_in_text(para.text)
                    if t != para.text:
                        set_para_text(para, t)

    doc.save(str(doc_path))


def main():
    targets = [
        ROOT / "华东师范大学硕士论文.docx",
        ROOT / "华东师范大学硕士论文.backup.20260606_183810.docx",
    ]
    for path in targets:
        if not path.exists():
            continue
        backup = path.with_name(
            f"{path.stem}.title.{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        )
        shutil.copy2(path, backup)
        apply(path)
        print(f"已更新: {path}")
        print(f"  备份: {backup}")


if __name__ == "__main__":
    main()
