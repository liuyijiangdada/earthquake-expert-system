#!/usr/bin/env python3
"""将近期代码实现同步到《华东师范大学硕士论文1.docx》。"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文1.docx"


def set_para_text(para, text: str) -> None:
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


REPLACEMENTS: dict[int, str] = {
    49: (
        "地震灾害具有突发性和破坏性，震前防御、震中应急与震后恢复三阶段的信息需求差异显著，"
        "传统问答方法难以同时满足静态知识的系统性与动态情报的实时性。本研究在知识图谱与向量检索协同框架之上，"
        "进一步构建动态—静态知识协同的地震应急问答系统：融合 Neo4j 知识图谱、Milvus/内存向量检索、"
        "CEIC+USGS 统一动态震情（core/earthquake_feed.py）、三阶段问题分类与不确定性感知调度器，"
        "并采用 Qwen1.5-1.8B+LoRA 本地推理。"
        "具体工作包括：（1）设计地震事件、区域、应急主题与处置步骤的知识图谱模式，实现问题驱动的子图查询；"
        "（2）采用 BAAI/bge-small-zh-v1.5 对应急知识 topic 级分块并向量化，完成 Top-K 语义检索；"
        "（3）提出震前/震中/震后阶段分类与 static_confidence、dynamic_availability、urgency "
        "三信号调度策略，按阶段路由静态图谱、向量检索与动态数据源；"
        "（4）实现 QueryContextBuilder 统一组装上下文，配合 response_guard 质量兜底、"
        "多模态资源匹配、震中地图/避难所/余震图等第三层输出增强；"
        "（5）基于 Flask+Vue 3 与 Android 客户端完成 Web/移动端部署，提供分阶段评测问集接口。"
        "通过配置开关构造仅 LLM、仅图谱、仅 RAG 与全协同四组消融实验，并按震前/震中/震后分阶段统计指标，"
        "验证协同机制与阶段调度对事实一致性、要点完整性与格式合规性的提升（定量结果见第 7 章）。"
    ),
    54: (
        "Earthquake disasters are characterized by suddenness and destructiveness. Information needs differ "
        "markedly across the three phases of pre-earthquake preparedness, during-earthquake response, and "
        "post-earthquake recovery. Conventional question-answering approaches struggle to simultaneously "
        "satisfy the systematic nature of static knowledge and the timeliness of dynamic intelligence. "
        "Building upon a collaborative framework of knowledge graphs and vector retrieval, this study further "
        "develops a dynamic–static knowledge cooperative earthquake emergency QA system that integrates a "
        "Neo4j knowledge graph, Milvus/in-memory vector retrieval, unified CEIC+USGS dynamic earthquake "
        "information (core/earthquake_feed.py), three-phase question classification, and an uncertainty-aware "
        "scheduler, with local inference based on Qwen1.5-1.8B+LoRA. The main work includes: (1) designing a "
        "knowledge graph schema for earthquake events, regions, emergency topics, and procedural steps, and "
        "implementing query-driven subgraph retrieval; (2) applying BAAI/bge-small-zh-v1.5 to chunk emergency "
        "knowledge at the topic level, embedding the chunks, and performing Top-K semantic retrieval; "
        "(3) proposing pre-, during-, and post-earthquake phase classification and a three-signal scheduling "
        "strategy based on static_confidence, dynamic_availability, and urgency, routing static graph context, "
        "vector retrieval, and dynamic data sources by phase; (4) implementing QueryContextBuilder for unified "
        "context assembly, together with response_guard quality fallback, multimodal resource matching, and "
        "Layer-3 output enrichment such as epicenter maps, shelter information, and aftershock charts; "
        "(5) deploying Web and mobile clients based on Flask+Vue 3 and an Android app, and providing a phased "
        "evaluation query-set API. Four ablation settings—LLM-only, graph-only, RAG-only, and full "
        "collaboration—are constructed via configuration switches; metrics are stratified by pre-, during-, and "
        "post-earthquake phases to verify improvements in factual consistency, key-point completeness, and format "
        "compliance brought by the collaborative mechanism and phase-aware scheduling (quantitative results are "
        "presented in Chapter 7)."
    ),
    83: (
        "（3）动态—静态知识协同与不确定性感知调度。实现 PhaseClassifier 判定震前/震中/震后，"
        "DynamicRetriever 经 core/earthquake_feed.py 统一接入 CEIC 与中国区 USGS 实时目录"
        "（去重合并、TTL 缓存），Scheduler 综合三信号决定 KG/RAG/动态源的启用优先级，"
        "并由 services/context_builder.py 统一组装提示；services/response_guard.py 对乱码与低置信度回答做 RAG 降级。"
    ),
    97: (
        "展示层包括 Vue 3 SPA（frontend/，构建产物 static/spa/）与 Android 客户端（disaster_app/android/）。"
        "Web 端提供三阶段快捷提问胶囊、阶段标签、多模态资源卡片与用户可见的 reliability_hint 可靠性提示条，"
        "元信息字段（phase、static_confidence 等）默认不对普通用户展示；"
        "Android 端通过 Retrofit 调用同一 Flask API（模拟器默认 http://10.0.2.2:8000）。"
    ),
    99: (
        "协同调度层（core/）含 PhaseClassifier、Scheduler、DynamicRetriever、MultimodalOutput，"
        "以及 knowledge_signals 计算 static_confidence 与 dynamic_availability。"
        "知识数据层含 Neo4j（kg/neo4j_kg.py）、RAG（rag/emergency_rag.py，Milvus 或内存）、"
        "shelters.json、multimodal_resources.json 与 CEIC+USGS 统一动态源（core/earthquake_feed.py）。"
        "模型层为常驻 Qwen1.5-1.8B+LoRA 与懒加载 Qwen2-VL-2B-Instruct。"
    ),
    104: (
        "（1）文本问答流（query_type: llm）：前端提交 params.input 与可选 history（最近 3 轮）。"
        "服务端先经 PhaseClassifier 判定震前/震中/震后，Scheduler 结合三信号决定 KG/RAG/动态源优先级；"
        "QueryContextBuilder 并行探测知识信号并组装【知识图谱】【参考资料】【动态信息】【阶段提示】等块，"
        "再套入 Qwen 对话模板，左侧截断保尾后由 LoRA 模型生成；"
        "guard_response 检测乱码/低置信度并在必要时降级为 RAG 摘要；"
        "响应可含元信息 meta（phase、schedule_reasoning、reliability_hint、media_resources 等）"
        "及第三层地图/图表链接。"
    ),
    106: (
        "（3）数据更新流：POST /api/update-data 触发 kg.update_from_realtime_data()，"
        "经 core/earthquake_feed.py 拉取 CEIC+USGS 统一目录并增量写入 Neo4j；"
        "动态问答路径另由 DynamicRetriever 带缓存拉取同一 feed。"
    ),
    111: (
        "模块映射：前端（Vue/Android）→ app.py 路由 → services/context_builder.py → "
        "kg/neo4j_kg.py、rag/emergency_rag.py、core/earthquake_feed.py、core/dynamic_retriever.py → "
        "services/response_guard.py、services/output_enricher.py（第三层）→ transformers+peft / qwen_vl_handler。"
        "docker-compose.yml 提供 Neo4j 与 Milvus 栈；本地开发可设 RAG_USE_MEMORY_RAG=True 免 Milvus。"
    ),
    147: (
        "Scheduler 读取 static_confidence（RAG/图谱探测）、dynamic_availability（CEIC/USGS 统一探测）"
        "与 urgency，按阶段调用 _schedule_pre/_during/_post；"
        "动态源不可达时自动降级静态分支，并通过 reliability_hint 向前端提示。"
        "中等置信度场景下亦会附带可靠性说明，避免用户误判生成内容的时效性。"
    ),
    148: (
        "形式化描述：记知识图谱 G、应急分块集合 D、动态源 Δ（CEIC+USGS 统一 feed 等）。"
        "文本路径抽象为：a = guard(f_θ(Φ(q,G), Ψ(q,D), Δ(q), Σ(q), q))，"
        "其中 Φ 为图谱上下文，Ψ 为 RAG Top-K，Δ 为动态震情摘要，"
        "Σ 为阶段分类与调度产生的 prompt_suffix 与 validity_hint，guard 为 response_guard 质量兜底。"
        "纯图谱路径返回结构化列表，不经 f_θ。"
    ),
    160: (
        "动态分支 Δ 职责：DynamicRetriever 在震中或高紧急度问句下经 earthquake_feed 拉取 "
        "CEIC 与中国区 USGS 目录，去重合并、经 bbox 过滤与缓存后写入【动态信息】；"
        "单源失败时自动降级另一源。震后政策类问题可触发 policy_rss 摘要。"
        "MultimodalOutput 按阶段与关键词匹配 static/emergency 示意图资源。"
    ),
    163: (
        "端到端流程：（1）阶段分类 PhaseClassifier.classify(q)；"
        "（2）探测知识信号 compute_knowledge_signals；"
        "（3）Scheduler.decide() 生成 ScheduleDecision；"
        "（4）组装 Φ、Ψ、Δ 与阶段 prompt_suffix；"
        "（5）LoRA 生成；（6）guard_response 乱码检测与 RAG 降级；"
        "（7）OutputEnricher 附加地图/避难所/余震图等媒体；"
        "（8）_sanitize_response 过滤异常标记。"
    ),
    189: (
        "（3）EmergencyTopic（应急主题）：表示一类应急知识主题，属性包括 id、title、category、source、"
        "phase_tag（震前/震中/震后/通用）、temporal_validity（半衰期，供调度器判断静态置信度）等。"
    ),
    195: (
        "数据加载在 Neo4jKG.run() 中执行。优先路径（JSON）：先导入应急知识 emergency_knowledge.json，"
        "读取 topics 列表（当前 13 个主题），对每个主题 MERGE EmergencyTopic 并 SET phase_tag、"
        "temporal_validity 等属性，逐条 MERGE GuidanceStep，建立 HAS_STEP；"
        "再导入真实地震目录 real_earthquakes_catalog.json，遍历 earthquakes，"
        "对每条 MERGE Earthquake 并 SET 属性，若记录含 region 则 MERGE Region 并建立 OCCURRED_IN。"
        "该顺序保证应急主题节点已存在。兼容路径（CSV）按列映射导入。"
        "运行期更新通过 update_from_realtime_data() 经 earthquake_feed 实现增量合并。"
    ),
    237: (
        "Web 前端（Vue 3 + Vite）采用应急深色主题。AppHeader 展示系统能力标签；"
        "QuickPanel 按震前/震中/震后分组快捷问句；ChatMessages 渲染 Markdown、多模态卡片"
        "与 reliability_hint 可靠性提示条；Composer 支持文本发送。"
        "元信息字段（phase、static_confidence 等）仅在后端 meta 载荷中保留，默认不向用户展示。"
    ),
    258: (
        "图谱规模：Earthquake/Region/EmergencyTopic/GuidanceStep 节点及 OCCURRED_IN、HAS_STEP 等关系"
        "（以 Neo4j 实际统计为准，当前约 14 条地震事件、13 个应急主题）；"
        "RAG 为 topic 级 chunk，条数等于应急主题数；"
        "shelters.json 提供避难所 POI；multimodal_resources.json 提供分阶段示意图元数据。"
    ),
    271: (
        "分阶段观察：震前问句应主要命中 RAG 与静态科普，dynamic_availability 低；"
        "震中问句应触发 CEIC+USGS 动态段与短句式回答；震后问句侧重政策 RSS 与房屋评估流程图。"
        "meta 元信息可用于核对调度 reasoning、reliability_hint 与阶段标签是否一致。"
    ),
    276: (
        "案例1（震中+动态）：问「刚才地震多大？震中在哪？」B0 易幻觉；B3 应触发震中阶段、"
        "拉取 CEIC+USGS 统一动态摘要并附震中地图（高德或 OSM 降级），事实一致性优于 B0/B2。"
    ),
    289: (
        "（1）深化 CEIC 正式目录 API 对接与容错（当前已实现 CEIC+USGS 统一 feed 框架），"
        "扩充避难所与救援资源实体；完善 response_guard 与自动化质量评测。"
    ),
    305: (
        "（4）基于 Flask 后端与 Vue 3 前端实现了完整的 Web 问答系统，支持自然语言对话、"
        "图谱直连查询、ECharts 数据可视化、CEIC+USGS 统一震情更新与 reliability_hint 可靠性提示。"
    ),
    308: (
        "（1）知识图谱扩展：当前图谱规模有限（14 条地震事件、13 个应急主题），"
        "未来可接入更多权威数据源（如中国地震台网正式目录 API），"
        "并增加震害评估、救援资源等实体类型。"
    ),
}


def apply_sync() -> None:
    if not THESIS.exists():
        raise FileNotFoundError(THESIS)

    backup = THESIS.with_name(
        f"华东师范大学硕士论文1.backup.{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    )
    shutil.copy2(THESIS, backup)
    print(f"已备份: {backup}")

    doc = Document(str(THESIS))
    updated = 0
    for idx, text in REPLACEMENTS.items():
        if idx < len(doc.paragraphs):
            set_para_text(doc.paragraphs[idx], text)
            updated += 1
        else:
            print(f"警告: 段落索引 {idx} 超出范围（共 {len(doc.paragraphs)} 段）")

    doc.save(str(THESIS))
    print(f"已更新 {updated} 处段落: {THESIS}")


if __name__ == "__main__":
    apply_sync()
