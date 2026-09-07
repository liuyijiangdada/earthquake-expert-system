#!/usr/bin/env python3
"""在图目录后插入表目录（点线页码 + 超链接书签）。"""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "华东师范大学硕士论文.docx"

CHAPTER_RE = re.compile(r"^第[一二三四五六七八九十百]+章")
TABLE_CAPTION_SHORT = re.compile(r"^表\s*(\d+)[-.．](\d+)\s*(.*)$")
WS_RE = re.compile(r"\s+")


def set_run_font(run, *, east_asia="宋体", western="Times New Roman", size_pt=12, bold=False):
    run.font.name = western
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:ascii"), western)
    rFonts.set(qn("w:hAnsi"), western)


def set_run_font_elem(run_elem, *, east_asia="宋体", western="Times New Roman", size_pt=12):
    rPr = run_elem.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        run_elem.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), western)
    rFonts.set(qn("w:hAnsi"), western)
    rFonts.set(qn("w:eastAsia"), east_asia)
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set(qn("w:val"), str(int(size_pt * 2)))
    color = rPr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rPr.append(color)
    color.set(qn("w:val"), "000000")


def make_run(text: str):
    r = OxmlElement("w:r")
    set_run_font_elem(r)
    t = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace() or "\t" in text:
        t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def insert_paragraph_after(para: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    para._p.addnext(new_p)
    new_para = Paragraph(new_p, para._parent)
    if text:
        new_para.add_run(text)
    return new_para


def add_page_break(para: Paragraph) -> None:
    run = para.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def clear_bookmarks_prefix(p: Paragraph, prefix: str) -> None:
    for child in list(p._p):
        if child.tag == qn("w:bookmarkStart") and (child.get(qn("w:name")) or "").startswith(prefix):
            bid = child.get(qn("w:id"))
            p._p.remove(child)
            for child2 in list(p._p):
                if child2.tag == qn("w:bookmarkEnd") and child2.get(qn("w:id")) == bid:
                    p._p.remove(child2)


def add_bookmark(p: Paragraph, name: str, bid: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    el = p._p
    if len(el) and el[0].tag == qn("w:pPr"):
        el.insert(1, start)
    else:
        el.insert(0, start)
    el.append(end)


def style_caption(para: Paragraph, text: str) -> None:
    for child in list(para._p):
        if child.tag != qn("w:pPr"):
            para._p.remove(child)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(0)
    run = para.add_run(text)
    set_run_font(run, east_asia="宋体", size_pt=10.5, bold=False)


def style_title(para: Paragraph) -> None:
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(12)
    pf.space_after = Pt(12)
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    if not para.runs:
        para.add_run(para.text or "表目录")
    for run in para.runs:
        set_run_font(run, east_asia="黑体", size_pt=16, bold=True)


def style_entry_pPr(p: Paragraph, tab_pos: float) -> None:
    pf = p.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.left_indent = Cm(0)
    pPr = p._p.get_or_add_pPr()
    for tabs in pPr.findall(qn("w:tabs")):
        pPr.remove(tabs)
    pf.tab_stops.add_tab_stop(Cm(tab_pos), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def write_linked_entry(p: Paragraph, title: str, page: int, anchor: str, tab_pos: float) -> None:
    style_entry_pPr(p, tab_pos)
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    hyper = OxmlElement("w:hyperlink")
    hyper.set(qn("w:anchor"), anchor)
    hyper.set(qn("w:history"), "1")
    hyper.append(make_run(title))
    hyper.append(make_run("\t"))
    hyper.append(make_run(str(page)))
    p._p.append(hyper)


def has_drawing(p: Paragraph) -> bool:
    return bool(p._p.findall(".//" + qn("w:drawing"))) or bool(p._p.findall(".//" + qn("w:pict")))


def has_page_br(p: Paragraph) -> bool:
    return any(br.get(qn("w:type")) == "page" for br in p._p.findall(".//" + qn("w:br")))


def drawing_height_emu(p: Paragraph, usable_h: int) -> int:
    total = 0
    for ext in p._p.findall(".//" + qn("wp:extent")):
        cy = ext.get("cy")
        if cy:
            total += int(cy)
    if total <= 0:
        return int(3.0 * 914400)
    return min(total, int(usable_h * 0.85))


def vis_len(s: str) -> float:
    return sum(0.55 if ord(ch) < 128 else 1.0 for ch in s)


def para_height(p: Paragraph, usable_w: int, usable_h: int) -> int:
    text = p.text or ""
    pf = p.paragraph_format
    sb = min(pf.space_before.emu if pf.space_before else 0, int(Pt(24).emu))
    sa = min(pf.space_after.emu if pf.space_after else 0, int(Pt(24).emu))
    h = sb + sa
    size = Pt(12).emu
    if p.runs:
        for r in p.runs:
            if r.font.size:
                size = r.font.size.emu
                break
    ls = 1.5
    if pf.line_spacing and isinstance(pf.line_spacing, (int, float)) and pf.line_spacing >= 1:
        ls = float(pf.line_spacing)
    line_h = int(size * ls)
    if has_drawing(p):
        h += drawing_height_emu(p, usable_h)
    if not text.strip():
        return h if has_drawing(p) else max(h, line_h // 3)
    chars_per_line = max(int(usable_w / size), 18)
    lines = max(1, int((vis_len(text) + chars_per_line - 0.01) // chars_per_line))
    return h + lines * line_h


def is_real_caption(text: str) -> bool:
    t = text.strip()
    m = TABLE_CAPTION_SHORT.match(t)
    if not m:
        return False
    title = (m.group(3) or "").strip()
    if not title or len(t) > 60:
        return False
    if "来自对" in t or "实验关闭" in t or "事实一致性最高" in t:
        return False
    return True


def remove_existing_table_toc(doc: Document) -> None:
    seen = False
    to_remove = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "表目录":
            seen = True
            to_remove.append(p._p)
            continue
        if not seen:
            continue
        if CHAPTER_RE.match(t):
            break
        to_remove.append(p._p)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def set_para_text(para: Paragraph, text: str) -> None:
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def main() -> None:
    backup = ROOT / f"华东师范大学硕士论文.backup.{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    shutil.copy2(SRC, backup)
    doc = Document(str(SRC))

    # 修正误把分析段当成表7-3题注的情况
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("表 7-3 来自对") or t.startswith("表7-3 来自对"):
            set_para_text(p, re.sub(r"^表\s*7-3\s*", "表7-3的结果", t, count=1))
            break

    body = doc.element.body
    tables = [c for c in body.iterchildren() if c.tag == qn("w:tbl")]
    desired = {
        0: "表7-1  核心依赖包版本",
        1: "表7-2  消融实验基线设置",
        2: "表7-3  快速消融实验结果",
    }
    for ti, tbl in enumerate(tables):
        want = desired[ti]
        prev = tbl.getprevious()
        prev_p = None
        prev_text = ""
        while prev is not None:
            if prev.tag == qn("w:p"):
                prev_text = "".join(prev.itertext()).strip()
                prev_p = prev
                if prev_text:
                    break
            prev = prev.getprevious()
        if is_real_caption(prev_text):
            m = TABLE_CAPTION_SHORT.match(prev_text)
            if m and m.group(1) == "7" and m.group(2) == str(ti + 1):
                style_caption(Paragraph(prev_p, doc), want)
                continue
        if ti == 2 or not is_real_caption(prev_text):
            new_p = OxmlElement("w:p")
            tbl.addprevious(new_p)
            style_caption(Paragraph(new_p, doc), want)

    remove_existing_table_toc(doc)

    captions: list[tuple[Paragraph, str, str, str]] = []
    body_started = False
    gate = False
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t in ("图目录", "表目录") or t.startswith("致谢"):
            gate = True
        if gate and CHAPTER_RE.match(t):
            body_started = True
        if not body_started:
            continue
        if is_real_caption(t):
            m = TABLE_CAPTION_SHORT.match(t)
            captions.append((p, m.group(1), m.group(2), m.group(3).strip()))

    seen: set[tuple[str, str]] = set()
    uniq: list[tuple[Paragraph, str, str, str]] = []
    for item in captions:
        key = (item[1], item[2])
        if key in seen:
            continue
        seen.add(key)
        uniq.append(item)
    captions = uniq
    if not captions:
        raise SystemExit("未找到表题")

    for i, (p, _, _, _) in enumerate(captions):
        clear_bookmarks_prefix(p, "_TocTbl")
        add_bookmark(p, f"_TocTbl{i:03d}", 6000 + i)

    sec = doc.sections[0]
    usable_h = sec.page_height - sec.top_margin - sec.bottom_margin
    usable_w = sec.page_width - sec.left_margin - sec.right_margin
    body_start = None
    seen_ack = in_front = False
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if t == "目录":
            in_front = True
        if in_front and t.startswith("致谢"):
            seen_ack = True
        if seen_ack and t.startswith("第一章") and "绪论" in t and i > 100:
            body_start = i
            break
    if body_start is None:
        raise SystemExit("未找到正文第一章")

    page, y = 1, 0
    tbl_pages: dict[tuple[str, str], int] = {}
    for i, p in enumerate(doc.paragraphs):
        if i < body_start:
            continue
        if i > body_start and has_page_br(doc.paragraphs[i - 1]):
            page, y = page + 1, 0
        ph = para_height(p, usable_w, usable_h)
        nxt = p._p.getnext()
        if nxt is not None and nxt.tag == qn("w:tbl"):
            ph += int(2.2 * 914400)
        if y + ph > usable_h and y > 0:
            page, y = page + 1, 0
        t = (p.text or "").strip()
        if is_real_caption(t):
            m = TABLE_CAPTION_SHORT.match(t)
            tbl_pages[(m.group(1), m.group(2))] = page
        y += ph

    seen_fig = False
    fig_end = None
    body_ch1 = None
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "图目录":
            seen_fig = True
            fig_end = p
            continue
        if not seen_fig:
            continue
        if CHAPTER_RE.match(t):
            body_ch1 = p
            break
        fig_end = p
    if fig_end is None or body_ch1 is None:
        raise SystemExit("无法定位图目录结束或正文第一章")

    to_remove = []
    crossed = False
    for p in doc.paragraphs:
        if p._p is fig_end._p:
            crossed = True
            continue
        if p._p is body_ch1._p:
            break
        if crossed:
            to_remove.append(p._p)
    for el in to_remove:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)

    tab_pos = max((sec.page_width - sec.left_margin - sec.right_margin) / 360000 - 0.1, 12.0)
    cursor = fig_end
    pb1 = insert_paragraph_after(cursor, "")
    add_page_break(pb1)
    cursor = pb1

    title_p = insert_paragraph_after(cursor, "表目录")
    style_title(title_p)
    cursor = title_p

    prev_ch = None
    for i, (_, ch, seq, title) in enumerate(captions):
        if prev_ch is not None and ch != prev_ch:
            cursor = insert_paragraph_after(cursor, "")
        entry = insert_paragraph_after(cursor, "")
        title_norm = WS_RE.sub(" ", title)
        label = f"表 {ch}.{seq}  {title_norm}"
        pg = tbl_pages.get((ch, seq), 1)
        write_linked_entry(entry, label, pg, f"_TocTbl{i:03d}", tab_pos)
        print(f"  {label} -> {pg}")
        cursor = entry
        prev_ch = ch

    pb2 = insert_paragraph_after(cursor, "")
    add_page_break(pb2)
    doc.save(str(SRC))
    print(f"backup: {backup.name}")
    print(f"tables: {len(captions)}")
    print("请先关闭已打开的论文再重新打开，避免旧内容覆盖。")


if __name__ == "__main__":
    main()
