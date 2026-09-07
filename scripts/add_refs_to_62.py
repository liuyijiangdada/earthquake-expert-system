# -*- coding: utf-8 -*-
"""补充参考文献至62条，并在第2章相关节补充引用述评。"""
import copy
from docx import Document
from docx.oxml.ns import qn

THESIS = "华东师范大学硕士论文.docx"

# 6 条新文献（[57]-[62]）
NEW_REFS = [
    "[57] 程麟淇, 安璐, 彭泽, 等. 基于双层知识图谱与检索增强生成的突发事件应急问答研究[J]. 情报理论与实践, 2026, 49(6): 171-180.",
    "[58] ZHOU W, HUANG M, LIU S, et al. Research on the construction and application of earthquake emergency information knowledge graph based on large language models[J]. IEEE Access, 2025, 13: 127742-127757.",
    "[59] 姚立伟, 任福, 王双燕, 等. 地震灾害应急预案的生成式编制方法[J]. 武汉大学学报(信息科学版), 2025, 50(6): 1159-1174.",
    "[60] 邵舒羽, 张扬, 刘艳. 基于KGCN的地质地震灾害事件演化结果预测[J]. 中国安全科学学报, 2025, 35(2): 212-219.",
    "[61] 段乙好, 李晓丽, 徐志双, 等. 中国地震局应急指挥中心地震应急快速响应技术系统设计与实现[J]. 中国地震, 2024, 40(3): 709-717.",
    "[62] 马秀丹, 陈雅慧, 李华玥, 等. 2024年中国大陆地震灾害损失述评[J]. 中国地震, 2025, 41(1): 173-180.",
]

# 第2章各节追加的引用述评（附加到对应段落末尾）
# 键=段落索引，值=追加文本
APPEND_TO_PARA = {
    # 2.1 末段（应急垂直领域 KG 实践），追加国内 KG+LLM 应急应用述评
    188: "近年国内研究亦将知识图谱与大语言模型引入地震应急场景：周文涛等[58]基于大语言模型构建地震应急信息知识图谱并以图数据库为底座支撑问答检索；姚立伟等[59]以大模型、智能体与知识图谱协同驱动地震应急预案的生成式编制；邵舒羽等[60]利用知识图谱卷积网络预测地质地震灾害事件演化结果。上述工作证实了结构化知识在应急决策中的价值，但多以预案编制或灾情预测为目标，尚未形成面向分阶段问答的动静态协同机制。",
    # 2.2 末段（RAG 与图谱互补），追加 KG+RAG 应急问答对比述评
    191: "程麟淇等[57]提出基于双层知识图谱与检索增强生成的突发事件应急问答方法，以事件层—规则层并行结构与混合检索验证了“图谱+检索”在应急问答中的可行性，但其面向通用突发事件、未引入震前/震中/震后的阶段先验与动态震情调度，这正是本文方法在地震应急场景下的差异化设计所在。",
    # 2.3 末段（动态数据接入），追加应急响应系统与灾情数据述评
    194: "在工程实践层面，段乙好等[61]设计并实现了中国地震局应急指挥中心的地震应急快速响应技术系统，集成触发判定、震害快速评估与产品服务于一体；马秀丹等[62]对2024年中国大陆地震灾害损失的系统述评则进一步印证了多源灾情数据规范化整合对应急决策的支撑作用。",
}


def main():
    doc = Document(THESIS)
    paras = doc.paragraphs

    # 1) 追加参考文献 [57]-[62]：定位 [56] 所在段落，其后逐条插入
    ref56_idx = None
    for i, p in enumerate(paras):
        if p.text.strip().startswith("[56]"):
            ref56_idx = i
            break
    if ref56_idx is None:
        raise RuntimeError("未找到 [56] 参考文献段落")

    # 以 [56] 段落为模板克隆格式
    template_para = paras[ref56_idx]
    template_pPr = template_para._p.find(qn('w:pPr'))
    template_run_rPr = None
    if template_para.runs:
        template_run_rPr = template_para.runs[0]._r.find(qn('w:rPr'))

    anchor_elem = template_para._p
    for ref_text in NEW_REFS:
        new_p = copy.deepcopy(template_para._p)
        # 清空内容
        for child in list(new_p):
            if child.tag == qn('w:r'):
                new_p.remove(child)
        # 新增 run
        from docx.oxml import OxmlElement
        r = OxmlElement('w:r')
        if template_run_rPr is not None:
            r.append(copy.deepcopy(template_run_rPr))
        t = OxmlElement('w:t')
        t.text = ref_text
        t.set(qn('xml:space'), 'preserve')
        r.append(t)
        new_p.append(r)
        anchor_elem.addnext(new_p)
        anchor_elem = new_p  # 下一条插在这条之后

    # 2) 第2章各段末尾追加引用述评
    # 重新读取段落（结构已变），按文本定位
    paras = doc.paragraphs
    targets = {
        "本章仅梳理通用构建范式与领域适配要点。": APPEND_TO_PARA[188],
        "本文第三章将据此设计主题级向量索引及其与图谱的联合注入顺序。": APPEND_TO_PARA[191],
        "本章仅给出动态接入与实时处理的技术基础。": APPEND_TO_PARA[194],
    }
    appended = 0
    for p in paras:
        txt = p.text.strip()
        for anchor, extra in targets.items():
            if txt.endswith(anchor):
                # 追加一个 run，复用末尾 run 格式
                src_rPr = None
                if p.runs:
                    src_rPr = p.runs[-1]._r.find(qn('w:rPr'))
                from docx.oxml import OxmlElement
                r = OxmlElement('w:r')
                if src_rPr is not None:
                    r.append(copy.deepcopy(src_rPr))
                t = OxmlElement('w:t')
                t.text = extra
                t.set(qn('xml:space'), 'preserve')
                r.append(t)
                p._p.append(r)
                appended += 1
                break

    if appended != 3:
        print(f"警告：预期追加3处述评，实际追加 {appended} 处")

    doc.save(THESIS)
    print("完成：已追加6条参考文献与3处第2章引用述评")


if __name__ == "__main__":
    main()
