#!/usr/bin/env python3
"""将第二章 2.5 系统架构迁入第五章 5.2，并精简第二章。

- 删除 2.5.1–2.5.3 详述，改为 2.5 本章小结（指向第五章）
- 把原图2-1/2-2/2-3 迁到 5.2，改号为图5-1/5-2/5-3
- 原界面图5-2～5-8 → 图5-4～5-10；实验图5-9/5-10 → 图5-11/5-12
"""
from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"
BACKUP = ROOT / "华东师范大学硕士论文.docx.bak-before-move-arch-2-5"


def _set_run_font(run, *, east_asia="宋体", size_pt=12.0, bold=False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), east_asia)


def set_para_text(para: Paragraph, text: str) -> None:
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def format_para(para: Paragraph, level: str) -> None:
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if level == "h2":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=14, bold=True)
    elif level == "h3":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=12, bold=True)
    elif level == "figure_caption":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
        for run in para.runs:
            _set_run_font(run, east_asia="宋体", size_pt=10.5, bold=False)
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.74)
        for run in para.runs:
            _set_run_font(run, east_asia="宋体", size_pt=12, bold=False)


def delete_paragraph(para: Paragraph) -> None:
    el = para._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def insert_before(anchor: Paragraph, text: str, level: str) -> Paragraph:
    new_p = deepcopy(anchor._p)
    for child in list(new_p):
        local = child.tag.split("}")[-1]
        if local in ("r", "drawing", "pict", "hyperlink"):
            new_p.remove(child)
    anchor._p.addprevious(new_p)
    para = Paragraph(new_p, anchor._parent)
    set_para_text(para, text)
    format_para(para, level)
    return para


def insert_drawing_before(anchor: Paragraph, drawing_el) -> Paragraph:
    new_p = deepcopy(drawing_el)
    anchor._p.addprevious(new_p)
    return Paragraph(new_p, anchor._parent)


def find_idx(doc: Document, pred) -> int:
    for i, p in enumerate(doc.paragraphs):
        if pred(p.text.strip()):
            return i
    raise KeyError("not found")


def has_drawing(p: Paragraph) -> bool:
    return bool(p._element.xpath('.//*[local-name()="drawing"]'))


CH2_SUMMARY = (
    "本章梳理知识图谱构建、向量检索与 RAG、动态数据接入以及问答评价指标等理论基础，"
    "为第三、四章方法设计与第五章系统实现提供术语与概念准备。"
    "系统总体逻辑架构、请求数据流与技术栈映射属于实现层内容，详见第五章 5.2；"
    "本章不再展开系统说明书式架构叙述，以免与实现章重复。"
)

ARCH_BLOCKS: list[tuple[str, str]] = [
    ("h3", "5.2.2  系统总体逻辑架构"),
    (
        "body",
        "系统按展示、应用服务、协同调度、知识数据与模型五层组织，另设配置模块管理各功能开关，"
        "部署和消融实验共用同一套启停方式。",
    ),
    (
        "body",
        "展示层提供 Web 与 Android 两种入口。Web 端含三阶段快捷提问、阶段标识、多模态资源和可靠性提示，"
        "元信息字段默认不对普通用户展示；Android 端调用同一后端完成问答与资源获取。",
    ),
    (
        "body",
        "应用服务层对外提供文本问答、图文问答、数据更新、评测问集、阶段分类和地图代理等接口，"
        "负责参数校验与结果返回。文本问答时先组装图谱、检索与动态震情上下文，再调用生成模块。",
    ),
    (
        "body",
        "协同调度层包含阶段分类、规则阈值调度、动态震情检索和第三层输出增强。"
        "它根据静态置信度、动态可用性与紧急程度，决定是否启用动态源，并写入提示约束与可靠性提示；"
        "图谱与检索是否进入上下文，同时受全局配置约束。",
    ),
    (
        "body",
        "知识数据层提供结构化与非结构化知识支撑，包括 Neo4j 知识图谱、向量检索库（Milvus 或内存索引）、"
        "避难所与多模态资源数据，以及轻量 CEIC 抓取与 USGS 合并后的动态震情数据。"
        "模型层部署常驻的 Qwen1.5-1.8B（LoRA）文本生成模型，并按需加载多模态模型。",
    ),
    (
        "body",
        "配置模块集中存放功能开关和外部服务参数。图谱、检索、阶段分类、动态检索和输出增强均可按实验需要打开或关闭，"
        "便于对照复现。系统总体逻辑架构见图5-1。",
    ),
    ("drawing", "arch"),
    ("figure_caption", "图5-1  系统总体逻辑架构图"),
    ("h3", "5.2.3  数据流与请求处理流程"),
    (
        "body",
        "（1）文本问答流：前端提交用户问句及可选多轮历史（最近 3 轮）。"
        "服务端依次完成阶段分类、知识信号探测、规则调度与上下文检索组装，"
        "写入【知识图谱】【参考资料】【动态信息】【阶段提示】等段落，再套入对话模板，"
        "经左侧截断保尾后由 LoRA 模型生成；必要时对异常或低置信度回答回退为检索摘要。"
        "响应可附带阶段判定、调度理由、可靠性提示与可选多媒体资源元数据。",
    ),
    (
        "body",
        "（2）图谱直连流：按查询类型调用全部列表、按地区、按震级或按深度等结构化检索，"
        "无需经过大模型，适用于列表筛选与低延迟查询场景。",
    ),
    (
        "body",
        "（3）数据更新流：通过数据更新接口触发图谱增量更新，经统一动态震情源拉取轻量 CEIC 抓取结果与 USGS 目录合并后写入 Neo4j；"
        "动态问答路径由动态检索模块带缓存读取同一数据源（单源失败时可回退至另一源或返回空结果并提示）。"
        "问答请求数据流见图5-2。",
    ),
    ("drawing", "flow"),
    ("figure_caption", "图5-2  问答请求数据流图"),
    ("h3", "5.2.4  技术选型与模块映射"),
    (
        "body",
        "一次文本问答大致经过：前端提交、服务路由、上下文组装、知识获取、质量兜底与第三层增强，最后由生成模型输出。"
        "前端为 Vue 3 Web 与 Android 客户端；服务端汇总图谱、向量检索和动态震情后，"
        "再交给 Transformers + PEFT/LoRA 文本模型，或按需调用多模态模型。",
    ),
    (
        "body",
        "文本生成采用 Qwen1.5-1.8B 基座，并加载地震应急领域 LoRA 适配器（秩 r=8，α=16，dropout=0.1）；"
        "嵌入模型为 BAAI/bge-small-zh-v1.5。图存储采用 Neo4j；向量检索支持主题级索引，并可选用 Milvus 或内存检索模式。"
        "后端以 Flask 提供服务。知识侧分别对应 Neo4j 图谱模块、应急知识向量检索模块，"
        "以及 CEIC 轻量抓取与 USGS 合并的动态震情模块。"
        "基础设施方面，可通过 Docker Compose 部署 Neo4j 与 Milvus；本地实验亦可启用内存向量检索模式。"
        "设备优先级为 MPS→CUDA→CPU。技术栈与模块映射关系见图5-3。",
    ),
    ("drawing", "stack"),
    ("figure_caption", "图5-3  技术栈与模块映射关系"),
]


def replace_figure_numbers(doc: Document) -> None:
    """先升后降，避免连锁替换冲突。界面原5-2..5-8→5-4..5-10；实验5-9/5-10→5-11/5-12。"""
    # Only touch chapter 5 UI/experiment captions and mentions; architecture already set as 5-1..5-3.
    mapping_pairs = [
        ("图5-10", "图5-12"),
        ("图5-9", "图5-11"),
        ("图5-8", "图5-10"),
        ("图5-7", "图5-9"),
        ("图5-6", "图5-8"),
        ("图5-5", "图5-7"),
        ("图5-4", "图5-6"),
        ("图5-3", "图5-5"),
        ("图5-2", "图5-4"),
    ]
    # Find start of 5.3 so we don't rewrite the new 5.2 architecture captions we just inserted.
    start_53 = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("5.3") and "详细设计" in p.text:
            start_53 = i
            break
    if start_53 is None:
        raise RuntimeError("找不到 5.3")

    for i, p in enumerate(doc.paragraphs):
        if i < start_53:
            continue
        t = p.text
        if not t:
            continue
        new = t
        for old, nw in mapping_pairs:
            new = new.replace(old, nw)
        if new != t:
            set_para_text(p, new)
            # keep caption centering if it looks like a figure caption
            if new.startswith("图5-"):
                format_para(p, "figure_caption")


def patch_toc_figures(doc: Document) -> None:
    repl = {
        "图 2.1 系统总体逻辑架构图": "图 5.1 系统总体逻辑架构图",
        "图 2.2 问答请求数据流图": "图 5.2 问答请求数据流图",
        "图 2.3 技术栈与模块映射关系": "图 5.3 技术栈与模块映射关系",
    }
    for p in doc.paragraphs:
        t = p.text
        for old, new in repl.items():
            if old in t:
                set_para_text(p, t.replace(old, new))
                break


def main() -> None:
    shutil.copy2(THESIS, BACKUP)
    print(f"备份 -> {BACKUP.name}")

    doc = Document(str(THESIS))

    # --- capture drawings from 2.5 ---
    i25 = find_idx(doc, lambda t: t.startswith("2.5") and "架构" in t.replace(" ", ""))
    i3 = find_idx(doc, lambda t: t.startswith("第三章") and "静态" in t)
    drawings = {}
    for i in range(i25, i3):
        p = doc.paragraphs[i]
        t = p.text.strip()
        if has_drawing(p):
            # associate by following caption
            cap = ""
            for j in range(i + 1, min(i + 4, i3)):
                ct = doc.paragraphs[j].text.strip()
                if ct.startswith("图2-"):
                    cap = ct
                    break
            if "图2-1" in cap:
                drawings["arch"] = deepcopy(p._element)
            elif "图2-2" in cap:
                drawings["flow"] = deepcopy(p._element)
            elif "图2-3" in cap:
                drawings["stack"] = deepcopy(p._element)
    if set(drawings) != {"arch", "flow", "stack"}:
        raise RuntimeError(f"未完整捕获架构图 drawing: {set(drawings)}")
    print("已捕获图2-1/2-2/2-3 drawing")

    # --- replace 2.5 block with summary ---
    # delete i25 .. i3-1
    to_del = [doc.paragraphs[i] for i in range(i25, i3)]
    for p in to_del:
        delete_paragraph(p)
    # re-find chapter 3 and insert summary before it
    i3 = find_idx(doc, lambda t: t.startswith("第三章") and "静态" in t)
    ch3 = doc.paragraphs[i3]
    insert_before(ch3, CH2_SUMMARY, "body")
    insert_before(ch3, "2.5  本章小结", "h2")
    print("已将 2.5 改为本章小结")

    # --- rebuild 5.2.2 onward until 5.3 ---
    i522 = find_idx(doc, lambda t: t.startswith("5.2.2"))
    i53 = find_idx(doc, lambda t: t.startswith("5.3") and "详细设计" in t)
    to_del = [doc.paragraphs[i] for i in range(i522, i53)]
    for p in to_del:
        delete_paragraph(p)

    i53 = find_idx(doc, lambda t: t.startswith("5.3") and "详细设计" in t)
    anchor = doc.paragraphs[i53]
    for level, text in ARCH_BLOCKS:
        if level == "drawing":
            insert_drawing_before(anchor, drawings[text])
        else:
            insert_before(anchor, text, level)
    print("已写入 5.2.2–5.2.4 架构内容与图5-1～5-3")

    # --- renumber UI / experiment figures in 5.3+ ---
    replace_figure_numbers(doc)
    patch_toc_figures(doc)
    print("已重编号界面/实验图题")

    doc.save(str(THESIS))

    # verify
    doc2 = Document(str(THESIS))
    blob = "\n".join(p.text for p in doc2.paragraphs)
    checks = {
        "no_25_arch_detail": "2.5.1" not in blob and "图2-1" not in blob,
        "has_25_summary": "2.5  本章小结" in blob or "2.5 本章小结" in blob,
        "has_fig51_arch": "图5-1  系统总体逻辑架构图" in blob,
        "has_fig52_flow": "图5-2  问答请求数据流图" in blob,
        "has_fig53_stack": "图5-3  技术栈与模块映射关系" in blob,
        "has_ui_fig54": "图5-4  Web端问答主界面" in blob,
        "has_exp_fig511": "图5-11" in blob and "图5-12" in blob,
        "order": True,
    }
    # order: 2.5 summary before ch3; 5.2.2 before 5.3; fig5-1 before fig5-4
    texts = [p.text.strip() for p in doc2.paragraphs]

    def first_startswith(prefix: str) -> int:
        for i, t in enumerate(texts):
            if t.startswith(prefix):
                return i
        return -1

    i25s = first_startswith("2.5")
    i3 = first_startswith("第三章")
    i522 = first_startswith("5.2.2")
    i53 = first_startswith("5.3")
    i51 = first_startswith("图5-1")
    i54 = first_startswith("图5-4")
    checks["order"] = (
        0 <= i25s < i3
        and 0 <= i522 < i53
        and 0 <= i51 < i54
    )
    # drawings still present
    n_draw = sum(1 for p in doc2.paragraphs if has_drawing(p))
    checks["drawings_kept"] = n_draw >= 8  # 3 arch + UI/exp images

    for k, v in checks.items():
        print(f"  {'OK' if v else 'FAIL'}: {k}")
    if not all(checks.values()):
        raise SystemExit(1)
    print(f"完成。inline drawings 约 {n_draw} 处。请在 Word 中更新目录域。")


if __name__ == "__main__":
    main()
