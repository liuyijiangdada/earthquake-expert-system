#!/usr/bin/env python3
"""将《华东师范大学硕士论文.docx》第五、六章按方案 A 合并，第八章改号为第六章。

用法:
  python scripts/merge_thesis_ch5_ch6.py
  python scripts/merge_thesis_ch5_ch6.py --no-backup
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"
SUMMARY = ROOT / "data" / "eval" / "table_6_3_summary_60.json"
BACKUP = ROOT / "华东师范大学硕士论文.docx.bak-before-ch5-merge"

sys.path.insert(0, str(ROOT))
from scripts.thesis_ch5_merge_content import expand_blocks  # noqa: E402


def _set_run_font(
    run,
    *,
    east_asia: str = "宋体",
    western: str = "Times New Roman",
    size_pt: float = 12,
    bold: bool = False,
) -> None:
    run.font.name = western
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), east_asia)


def _clear_runs(para: Paragraph) -> None:
    for run in para.runs:
        run.text = ""


def set_para_text(para: Paragraph, text: str) -> None:
    _clear_runs(para)
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def _format_para(para: Paragraph, level: str) -> None:
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if level == "h1":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(12)
        pf.space_after = Pt(12)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=16, bold=True)
    elif level == "h2":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=14, bold=True)
    elif level == "h3":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(3)
        pf.space_after = Pt(3)
        for run in para.runs:
            _set_run_font(run, east_asia="黑体", size_pt=12, bold=True)
    elif level in ("caption", "figure_caption", "table_title"):
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
        size = 12 if level == "table_title" else 10.5
        ea = "黑体" if level == "table_title" else "宋体"
        for run in para.runs:
            _set_run_font(run, east_asia=ea, size_pt=size, bold=(level == "table_title"))
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.74)
        for run in para.runs:
            _set_run_font(run, east_asia="宋体", size_pt=12, bold=False)


def _insert_para_before(anchor: Paragraph, text: str, level: str) -> Paragraph:
    new_p = deepcopy(anchor._p)
    for child in list(new_p):
        if child.tag == qn("w:r") or child.tag.endswith("}r"):
            new_p.remove(child)
        # remove drawings from template clone
        if child.tag.endswith("}drawing") or child.tag.endswith("}pict"):
            new_p.remove(child)
    # also clear nested drawings in runs already removed; ensure empty
    for child in list(new_p):
        local = child.tag.split("}")[-1]
        if local in ("drawing", "pict", "r"):
            new_p.remove(child)
    anchor._p.addprevious(new_p)
    para = Paragraph(new_p, anchor._parent)
    set_para_text(para, text)
    _format_para(para, level)
    return para


def _delete_paragraph(para: Paragraph) -> None:
    el = para._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def _find_body_chapter(doc: Document, title_prefix: str, contain: str) -> int:
    """找正文中的章标题（排除目录：含制表符的行）。取最后一个匹配。"""
    idxs = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if "\t" in p.text:
            continue
        if t.startswith(title_prefix) and contain in t:
            idxs.append(i)
    if not idxs:
        raise KeyError(f"找不到章节: {title_prefix} {contain}")
    return idxs[-1]


def _find_para_index(doc: Document, predicate) -> int:
    for i, p in enumerate(doc.paragraphs):
        if predicate(p.text.strip()):
            return i
    raise KeyError("paragraph not found")


def build_eval_mapping(summary_path: Path) -> dict[str, str]:
    data = json.loads(summary_path.read_text(encoding="utf-8"))
    overall = data.get("overall") or data
    by_phase = data.get("by_phase") or {}
    m: dict[str, str] = {}
    for bid in ("B0", "B1", "B2", "B3"):
        s = overall.get(bid, {})
        m[f"FACT_{bid}"] = str(s.get("factual_accuracy_pct", "—"))
        m[f"COMP_{bid}"] = str(s.get("completeness_mean", "—"))
        m[f"FMT_{bid}"] = str(s.get("format_compliance_pct", "—"))
        phase_map = by_phase.get(bid, {})
        for phase, key in (("震前", "PRE"), ("震中", "DUR"), ("震后", "POST")):
            ps = phase_map.get(phase, {})
            m[f"FACT_{bid}_{key}"] = str(ps.get("factual_accuracy_pct", "—"))
            m[f"COMP_{bid}_{key}"] = str(ps.get("completeness_mean", "—"))
            m[f"FMT_{bid}_{key}"] = str(ps.get("format_compliance_pct", "—"))
    return m


def fill_placeholders(text: str, mapping: dict[str, str]) -> str:
    def repl(m: re.Match) -> str:
        return mapping.get(m.group(1), m.group(0))

    return re.sub(r"\{\{([A-Z0-9_]+)\}\}", repl, text)


def _add_picture(para: Paragraph, image_path: Path, *, width_in: float) -> None:
    run = para.add_run()
    run.add_picture(str(image_path), width=Inches(width_in))
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)


def _image_width(path: Path) -> float:
    # landscape screenshots wider; portrait mobile narrower
    name = path.name.lower()
    if name in {"image3.png", "image4.png", "image5.png", "image6.png", "image7.png"}:
        return 2.6
    if "fig-6" in name or "factual" in name or "completeness" in name:
        return 5.2
    return 5.0


def patch_toc(doc: Document) -> None:
    """尽量更新目录区章标题文字（页码仍需 Word 更新域）。"""
    replacements = {
        "第六章  协同推理与系统实现": "第六章  总结与展望",
        "第六章 协同推理与系统实现": "第六章 总结与展望",
        "第八章  总结与展望": "第六章  总结与展望",
        "第八章 总结与展望": "第六章 总结与展望",
    }
    for p in doc.paragraphs:
        t = p.text
        for old, new in replacements.items():
            if old in t:
                set_para_text(p, t.replace(old, new))
                break
        # 目录行：第六章\t...
        if p.text.strip().startswith("第六章") and "\t" in p.text:
            set_para_text(p, "第六章 总结与展望\t")
        if p.text.strip().startswith("第八章") and "\t" in p.text:
            set_para_text(p, "第六章 总结与展望\t")


def merge(doc: Document, mapping: dict[str, str]) -> None:
    ch5 = _find_body_chapter(doc, "第五章", "系统")
    ref_i = _find_para_index(
        doc,
        lambda t: t.startswith("参考文献") and "\t" not in t,
    )
    # 删除第五章到参考文献之前
    paras = list(doc.paragraphs)
    to_delete = [paras[i] for i in range(ch5, ref_i)]
    for p in to_delete:
        _delete_paragraph(p)

    # 重新定位参考文献
    ref = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("参考文献") and "\t" not in p.text:
            ref = p
            break
    if ref is None:
        raise RuntimeError("参考文献段落丢失")

    # 在参考文献前插入占位空段作为锚点
    spacer = OxmlElement("w:p")
    ref._p.addprevious(spacer)
    anchor = Paragraph(spacer, ref._parent)

    blocks = expand_blocks()
    # 正序插入：每次插到参考文献正前方，后插入的块会落在更靠近参考文献处，整体顺序正确。
    for level, text in blocks:
        text = fill_placeholders(text, mapping)
        if level == "image":
            path = ROOT / text
            new_p = deepcopy(ref._p)
            for child in list(new_p):
                local = child.tag.split("}")[-1]
                if local in ("r", "drawing", "pict", "hyperlink"):
                    new_p.remove(child)
            ref._p.addprevious(new_p)
            para = Paragraph(new_p, ref._parent)
            set_para_text(para, "")
            if path.exists():
                _add_picture(para, path, width_in=_image_width(path))
            else:
                set_para_text(para, f"【插图占位：文件缺失 {text}】")
                _format_para(para, "body")
        else:
            _insert_para_before(ref, text, level)

    # 删除临时 spacer（若仍为空）
    for p in list(doc.paragraphs):
        if p._element is spacer:
            _delete_paragraph(p)
            break

    patch_toc(doc)


def verify(doc: Document) -> dict[str, bool]:
    texts = [p.text.strip() for p in doc.paragraphs]
    blob = "\n".join(texts)

    def idx_start(prefix: str) -> int:
        for i, t in enumerate(texts):
            if "\t" in t:
                continue
            if t.startswith(prefix):
                return i
        return -1

    i5 = idx_start("第五章")
    i51 = idx_start("5.1  ")
    if i51 < 0:
        i51 = idx_start("5.1")
    i531 = idx_start("5.3.1")
    i54 = idx_start("5.4")
    i6 = idx_start("第六章")
    i61 = idx_start("6.1")
    iref = idx_start("参考文献")

    checks = {
        "has_ch5": i5 >= 0 and "系统实现" in texts[i5],
        "has_ch6_summary": i6 >= 0 and "总结" in texts[i6],
        "no_old_ch6_impl": "协同推理与系统实现" not in blob,
        "no_ch8": "第八章" not in blob,
        "has_ui_section": i531 >= 0 and "前端页面与交互" in texts[i531],
        "has_ui_figs": "图5-2" in blob and "图5-8" in blob,
        "has_fact_825": "82.5" in blob,
        "has_exp_section": i54 >= 0,
        "no_ch7_ref": "见第 7 章" not in blob and "见第7章" not in blob and "表 7-3" not in blob,
        "order_ok": all(x >= 0 for x in (i5, i51, i531, i54, i6, i61, iref))
        and i5 < i51 < i531 < i54 < i6 < i61 < iref,
    }
    return checks


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-backup", action="store_true")
    parser.add_argument("--thesis", type=Path, default=THESIS)
    args = parser.parse_args()

    thesis: Path = args.thesis
    if not thesis.exists():
        raise SystemExit(f"找不到论文: {thesis}")

    if not args.no_backup:
        shutil.copy2(thesis, BACKUP)
        print(f"已备份 -> {BACKUP.name}")

    mapping = build_eval_mapping(SUMMARY) if SUMMARY.exists() else {}
    if mapping:
        print(f"已加载评测数字: B3事实={mapping.get('FACT_B3')}%")
    else:
        print("警告: 未找到 summary JSON，占位符可能残留")

    doc = Document(str(thesis))
    merge(doc, mapping)
    doc.save(str(thesis))
    print(f"已保存 -> {thesis.name}")

    doc2 = Document(str(thesis))
    checks = verify(doc2)
    for k, v in checks.items():
        print(f"  {'OK' if v else 'FAIL'}: {k}")
    if not all(checks.values()):
        raise SystemExit(1)
    print("验收通过")


if __name__ == "__main__":
    main()
