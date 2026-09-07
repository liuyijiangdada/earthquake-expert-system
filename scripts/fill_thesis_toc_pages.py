#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""为目录写入页码，并将超链接样式固定为黑字无下划线（对齐戚媛媛目录观感）。

用法:
  .venv/bin/python scripts/fill_thesis_toc_pages.py
"""
from __future__ import annotations

import argparse
import math
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"
BACKUP = ROOT / "华东师范大学硕士论文.docx.bak-before-toc-pages"
BOOKMARK_PREFIX = "_TocThesis"
CHARS_PER_LINE = 32
LINES_PER_PAGE = 28


def ensure_black_hyperlink_styles(doc: Document) -> None:
    styles_el = doc.styles.element
    for style_id in ("17", "18", "Hyperlink", "FollowedHyperlink"):
        for st in styles_el.findall(qn("w:style")):
            if st.get(qn("w:styleId")) != style_id:
                continue
            rPr = st.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                st.append(rPr)
            for tag in ("w:color", "w:u"):
                node = rPr.find(qn(tag))
                if node is not None:
                    rPr.remove(node)
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "000000")
            rPr.append(color)
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "none")
            rPr.append(u)

    for p in doc.paragraphs:
        if "HYPERLINK" not in p._p.xml and "PAGEREF" not in p._p.xml and "w:hyperlink" not in p._p.xml:
            continue
        for r in p._p.findall(qn("w:r")):
            rPr = r.find(qn("w:rPr"))
            if rPr is None:
                continue
            for rs in list(rPr.findall(qn("w:rStyle"))):
                rPr.remove(rs)
            for tag in ("w:color", "w:u"):
                node = rPr.find(qn(tag))
                if node is not None:
                    rPr.remove(node)
            c = OxmlElement("w:color")
            c.set(qn("w:val"), "000000")
            rPr.append(c)
            uu = OxmlElement("w:u")
            uu.set(qn("w:val"), "none")
            rPr.append(uu)


def find_body_chapter1_element(doc: Document):
    el = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("第一章") and "绪论" in t and "\t" not in p.text:
            if "HYPERLINK" not in p._p.xml and "PAGEREF" not in p._p.xml:
                el = p._p
    return el


def estimate_bookmark_pages(doc: Document) -> tuple[dict[str, int], int]:
    body_el = find_body_chapter1_element(doc)
    if body_el is None:
        raise RuntimeError("找不到正文第一章")

    para_by_el = {p._p: p for p in doc.paragraphs}
    tbl_by_el = {t._tbl: t for t in doc.tables}
    page = 1
    line_acc = 0.0
    bookmark_page: dict[str, int] = {}
    passed_body = False

    def flush_lines(n: float) -> None:
        nonlocal page, line_acc
        line_acc += n
        while line_acc >= LINES_PER_PAGE:
            line_acc -= LINES_PER_PAGE
            page += 1

    for child in list(doc.element.body):
        tag = child.tag.split("}")[-1]
        if tag == "p":
            if child is body_el:
                passed_body = True
            if not passed_body:
                continue
            for bm in child.findall(qn("w:bookmarkStart")):
                name = bm.get(qn("w:name")) or ""
                if name.startswith(BOOKMARK_PREFIX):
                    bookmark_page[name] = page
            p = para_by_el.get(child)
            t = p.text.strip() if p is not None else ""
            if child.xpath('.//*[local-name()="drawing"]'):
                flush_lines(10)
            if not t:
                flush_lines(0.3)
                continue
            if re.match(r"^第[一二三四五六]章", t) or t in ("参考文献", "致谢"):
                flush_lines(2.0)
            elif re.match(r"^\d+\.\d+", t) and len(t) < 80:
                flush_lines(1.0)
            flush_lines(max(1.0, float(math.ceil(len(t) / CHARS_PER_LINE))))
        elif tag == "tbl":
            if not passed_body:
                continue
            tbl = tbl_by_el.get(child)
            rows = len(tbl.rows) if tbl is not None else 3
            flush_lines(max(3.0, rows * 1.2 + 1.5))
    return bookmark_page, page


def fill_pageref_results(doc: Document, bookmark_page: dict[str, int]) -> int:
    filled = 0
    for p in doc.paragraphs:
        if "PAGEREF" not in p._p.xml:
            continue
        runs = list(p._p.findall(qn("w:r")))
        i = 0
        while i < len(runs):
            fld = runs[i].find(qn("w:fldChar"))
            if fld is not None and fld.get(qn("w:fldCharType")) == "begin":
                instr = ""
                j = i + 1
                while j < len(runs):
                    instr_el = runs[j].find(qn("w:instrText"))
                    if instr_el is not None:
                        instr += instr_el.text or ""
                        j += 1
                        continue
                    fld2 = runs[j].find(qn("w:fldChar"))
                    if fld2 is not None and fld2.get(qn("w:fldCharType")) == "separate":
                        j += 1
                        break
                    j += 1
                if j < len(runs) and "PAGEREF" in instr:
                    m = re.search(r"PAGEREF\s+(\S+)", instr)
                    if m and m.group(1) in bookmark_page:
                        t_el = runs[j].find(qn("w:t"))
                        if t_el is not None:
                            t_el.text = str(bookmark_page[m.group(1)])
                            filled += 1
                i = j
            i += 1
    return filled


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-backup", action="store_true")
    args = ap.parse_args()
    if not args.no_backup:
        shutil.copy2(THESIS, BACKUP)
        print(f"backup -> {BACKUP.name}")

    doc = Document(str(THESIS))
    ensure_black_hyperlink_styles(doc)
    bookmark_page, total = estimate_bookmark_pages(doc)
    filled = fill_pageref_results(doc, bookmark_page)
    doc.save(str(THESIS))
    print(f"bookmarks={len(bookmark_page)} pages≈{total} filled={filled}")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "目录":
            print([doc.paragraphs[i + j].text for j in range(1, 6)])
            break


if __name__ == "__main__":
    main()
