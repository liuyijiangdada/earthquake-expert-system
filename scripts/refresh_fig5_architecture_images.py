#!/usr/bin/env python3
"""重新生成图5-1/5-2/5-3（改序号+超清），并替换论文 docx 中对应插图。"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
THESIS = ROOT / "华东师范大学硕士论文.docx"
BACKUP = ROOT / "华东师范大学硕士论文.docx.bak-before-refresh-fig5-1-3"
ARCH = ROOT / "docs" / "superpowers" / "architecture"

sys.path.insert(0, str(ROOT / "scripts"))
import generate_query_flow as qf  # noqa: E402
import generate_system_architecture as sa  # noqa: E402
import generate_tech_stack_table as ts  # noqa: E402


REPLACEMENTS = [
    ("图5-1", ARCH / "fig-5-1-system-architecture.png", 5.8),
    ("图5-2", ARCH / "fig-5-2-query-flow.png", 4.6),
    ("图5-3", ARCH / "fig-5-3-tech-stack.png", 5.8),
]


def has_drawing(p: Paragraph) -> bool:
    return bool(p._element.xpath('.//*[local-name()="drawing"]'))


def clear_runs(para: Paragraph) -> None:
    for run in para.runs:
        run.text = ""
    # remove drawings from paragraph
    for child in list(para._element):
        local = child.tag.split("}")[-1]
        if local in ("drawing", "pict"):
            para._element.remove(child)
        if local == "r":
            for sub in list(child):
                sl = sub.tag.split("}")[-1]
                if sl in ("drawing", "pict"):
                    child.remove(sub)


def replace_image_before_caption(doc: Document, caption_prefix: str, image: Path, width_in: float) -> None:
    paras = list(doc.paragraphs)
    cap_i = None
    for i, p in enumerate(paras):
        t = p.text.strip()
        if t.startswith(caption_prefix) and ("系统" in t or "问答" in t or "技术栈" in t or "数据流" in t or "架构" in t):
            cap_i = i
            break
    if cap_i is None:
        # softer match
        for i, p in enumerate(paras):
            if p.text.strip().startswith(caption_prefix):
                cap_i = i
                break
    if cap_i is None:
        raise KeyError(f"找不到题注: {caption_prefix}")

    # Prefer previous drawing paragraph; else empty/near previous
    target = None
    for j in range(cap_i - 1, max(-1, cap_i - 5), -1):
        p = paras[j]
        if has_drawing(p) or not p.text.strip():
            target = p
            if has_drawing(p):
                break
    if target is None:
        target = paras[cap_i - 1]

    clear_runs(target)
    # wipe leftover empty runs with drawings
    for child in list(target._element):
        local = child.tag.split("}")[-1]
        if local in ("drawing", "pict"):
            target._element.remove(child)

    run = target.add_run()
    run.add_picture(str(image), width=Inches(width_in))
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = target.paragraph_format
    pf.first_line_indent = Cm(0)
    print(f"  已替换 {caption_prefix} <- {image.name} ({image.stat().st_size // 1024} KB)")


def main() -> None:
    shutil.copy2(THESIS, BACKUP)
    print(f"备份 -> {BACKUP.name}")

    print("生成超清图…")
    p1 = sa.generate_png()
    sa.generate_svg()
    p2 = qf.generate_png()
    qf.generate_svg()
    p3 = ts.generate_png()
    ts.generate_svg()
    for p in (p1, p2, p3):
        from PIL import Image

        im = Image.open(p)
        print(f"  {p.name}: {im.size[0]}×{im.size[1]}px")

    doc = Document(str(THESIS))
    print("写入 docx…")
    for prefix, path, width in REPLACEMENTS:
        if not path.exists():
            raise FileNotFoundError(path)
        replace_image_before_caption(doc, prefix, path, width)
    doc.save(str(THESIS))
    print(f"已保存 -> {THESIS.name}")


if __name__ == "__main__":
    main()
