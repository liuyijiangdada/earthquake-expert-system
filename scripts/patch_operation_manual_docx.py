#!/usr/bin/env python3
"""根据操作手册截图补齐对应操作说明，重新生成 操作手册.docx。"""
from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "操作手册.docx"
BACKUP_PATH = ROOT / "操作手册.backup.docx"
MEDIA_DIR = ROOT / "data" / "manual_media"

SECTIONS = [
    {
        "title": "一、环境准备与服务启动",
        "steps": [
            "启动 Neo4j 知识图谱：在项目根目录执行 docker start biyelunwen-neo4j（首次部署可用 docker compose up -d neo4j）。",
            "（可选）启动 Milvus 向量库以启用完整 RAG：docker compose up -d。",
            "启动后端服务：python app.py，默认监听 http://127.0.0.1:8000。",
            "浏览器访问 http://127.0.0.1:8000 进入 Web 问答界面；开发调试时也可执行 cd frontend && npm run dev，访问 http://localhost:5173。",
            "Android 移动端：使用 Android Studio 打开 disaster_app 目录，编译安装到手机或模拟器。",
        ],
        "image": None,
    },
    {
        "title": "二、Web 端主界面",
        "steps": [
            "页面顶部显示系统名称「地震应急智能问答」及「服务已连接」状态；右侧为知识库统计卡片（地震条数、最大震级、地点种类）。",
            "左侧为「应急问答对话」区：上方能力标签（KG / RAG / USGS / 高德 / 多模态）标识当前系统能力。",
            "「场景快捷提问」按震前、震中、震后三阶段分类，提供家庭应急包、室内避险、最新震情等一键提问按钮。",
            "快捷按钮操作：单击将问题填入输入框；双击直接发送；⌘/Ctrl + 单击亦可直接发送。",
            "右侧「知识图谱震情洞察」展示震级分布环形图及浅源/中源/深源统计，便于了解知识库覆盖情况。",
        ],
        "image": "image1.png",
    },
    {
        "title": "三、Web 端文字问答",
        "steps": [
            "在底部输入框输入地震应急相关问题，或点击上方快捷按钮填入示例问题。",
            "点击「发送」或按 Enter 提交问题；Shift + Enter 可换行。",
            "系统返回带阶段标签（如「震前 · 科普 · 预防」）的结构化回答，并可能附带相关示意图或地图链接。",
            "回答下方可点击「有用 / 需改进」提交反馈，帮助优化回答质量。",
            "底部工具栏：「上传图片」支持多模态图片问答；「清空对话」清除当前会话；「刷新地震数据」从数据源更新知识图谱。",
        ],
        "image": "image2.png",
    },
    {
        "title": "四、移动端 · 智能问答",
        "steps": [
            "打开 App 后默认进入底部导航「问答」页。",
            "顶部显示系统标题与能力标签（KG、RAG、USGS、高德、多模态）。",
            "「场景快捷提问」横向滑动，可点选「应急包」「室内避险」「避难所」「最新震情」等常用问题。",
            "在输入框输入问题，点击右侧发送按钮提交；左侧图片图标可选择本地图片进行多模态问答。",
            "对话区展示助手欢迎语及问答历史，支持震前准备、震中避险、震后恢复等全周期咨询。",
        ],
        "image": "image3.png",
    },
    {
        "title": "五、移动端 · 应急安全指引",
        "steps": [
            "点击底部导航「指引」进入本页。",
            "阅读「震前 / 震中 / 震后」三阶段要点卡片，了解防御准备、避险动作与恢复注意事项。",
            "中部「安全状态」卡片显示当前状态（未知 / 安全 / 注意 / 危险）。",
            "在「上报我的安全状态」区域选择安全等级（安全、注意、危险），填写「当前位置」。",
            "点击「更新状态」提交个人安全状态，便于家人或救援方了解你的情况。",
        ],
        "image": "image4.png",
    },
    {
        "title": "六、移动端 · 震后紧急求助",
        "steps": [
            "点击底部导航「求助」进入震后紧急求助页。",
            "点击右上角「+ 发起求助」打开求助表单。",
            "选择「求助类型」（救援、医疗、物资、避难所等）。",
            "填写「位置」与「描述情况」，必要时开启「紧急求助」开关。",
            "点击「提交」发送求助信息；列表区将展示已提交的求助记录。",
        ],
        "image": "image5.png",
    },
    {
        "title": "七、移动端 · 灾情上报",
        "steps": [
            "点击底部导航「上报」进入灾情上报页。",
            "点击右上角「+ 上报」打开上报表单。",
            "在「灾情类型」下拉框中选择类型（地震、火灾、洪水、滑坡、其他）。",
            "填写「灾害位置」与「详细描述」，尽量提供可核实的位置与现场情况。",
            "点击「提交」完成上报；下方列表可查看历史上报及处理状态（待处理 / 处理中 / 已解决）。",
        ],
        "image": "image6.png",
    },
    {
        "title": "八、移动端 · 志愿者招募",
        "steps": [
            "点击底部导航「志愿」进入志愿者招募页。",
            "点击右上角「+ 报名」打开注册表单。",
            "填写「姓名」「电话」「所在地区」及「专业技能」（多个技能用逗号分隔，如：医疗、搜救、心理辅导）。",
            "点击「报名」提交申请；报名成功后可在列表中查看志愿者信息。",
            "本功能用于灾后志愿力量登记，与 Web 端应急问答形成互补。",
        ],
        "image": "image7.png",
    },
]


def _set_run_font(run, east_asia: str = "宋体", size_pt: float = 12) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _add_heading(doc: Document, text: str, level: int) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        _set_run_font(run, east_asia="黑体", size_pt=16 if level == 1 else 14)


def _add_step(doc: Document, index: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run(f"{index}. {text}")
    _set_run_font(run)


def _extract_images_from_docx(docx_path: Path, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(docx_path) as zf:
        for name in zf.namelist():
            if name.startswith("word/media/") and not name.endswith("/"):
                target = out_dir / Path(name).name
                target.write_bytes(zf.read(name))


def build_manual() -> None:
    if not MEDIA_DIR.exists() or not any(MEDIA_DIR.glob("image*.png")):
        if DOCX_PATH.exists():
            _extract_images_from_docx(DOCX_PATH, MEDIA_DIR)
        else:
            raise FileNotFoundError(f"未找到截图资源，请确保 {DOCX_PATH} 存在")

    if DOCX_PATH.exists() and not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)

    doc = Document()
    title = doc.add_heading("地震应急智能问答系统 操作手册", level=0)
    for run in title.runs:
        _set_run_font(run, east_asia="黑体", size_pt=22)

    intro = doc.add_paragraph(
        "本文档说明 Web 端与 Android 移动端的主要功能与操作步骤，"
        "配图与系统实际界面一致，便于答辩演示与日常使用。"
    )
    intro.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for run in intro.runs:
        _set_run_font(run)

    for section in SECTIONS:
        _add_heading(doc, section["title"], level=1)
        for i, step in enumerate(section["steps"], start=1):
            _add_step(doc, i, step)

        image_name = section.get("image")
        if image_name:
            image_path = MEDIA_DIR / image_name
            if image_path.exists():
                doc.add_paragraph()
                doc.add_picture(str(image_path), width=Inches(5.8))
                cap = doc.add_paragraph(f"图：{section['title'].split('、', 1)[-1]}界面截图")
                cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                for run in cap.runs:
                    _set_run_font(run, size_pt=10.5)
                    run.italic = True

        doc.add_paragraph()

    doc.save(str(DOCX_PATH))
    print(f"已更新: {DOCX_PATH}")
    if BACKUP_PATH.exists():
        print(f"原文件备份: {BACKUP_PATH}")


if __name__ == "__main__":
    build_manual()
